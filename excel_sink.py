from __future__ import annotations

import json
import os
import time
from contextlib import contextmanager
from typing import Any, Dict, Iterable, Optional

from openpyxl import Workbook, load_workbook


DEFAULT_SHEETS = {
    "Calls": ["session_id", "customer_id", "start_ts", "end_ts", "disposition", "ptp_date", "callback_time"],
    "Actions": ["ts", "session_id", "action_name", "payload_json", "result_json", "ok"],
    "Tickets": ["ts", "session_id", "customer_id", "category", "reason", "summary", "status"],
    "Messages": ["ts", "session_id", "role", "content_redacted"],
    "Outbox": ["ts", "session_id", "channel", "to", "body", "link"],
}


@contextmanager
def _file_lock(path: str):
    """Advisory lock (best-effort) so we can safely append rows in demos.

    On macOS/Linux, fcntl is available. If locking fails, we proceed without it.
    """
    lock_path = f"{path}.lock"
    fd = None
    try:
        fd = os.open(lock_path, os.O_CREAT | os.O_RDWR, 0o644)
        try:
            import fcntl  # type: ignore

            fcntl.flock(fd, fcntl.LOCK_EX)
        except Exception:
            pass
        yield
    finally:
        if fd is not None:
            try:
                import fcntl  # type: ignore

                fcntl.flock(fd, fcntl.LOCK_UN)
            except Exception:
                pass
            try:
                os.close(fd)
            except Exception:
                pass


class ExcelOutcomeSink:
    def __init__(self, path: str) -> None:
        self.path = path

    def ensure_workbook(self) -> None:
        if os.path.exists(self.path):
            return
        os.makedirs(os.path.dirname(self.path) or ".", exist_ok=True)
        wb = Workbook()
        # Remove default sheet.
        wb.remove(wb.active)
        for name, headers in DEFAULT_SHEETS.items():
            ws = wb.create_sheet(name)
            ws.append(list(headers))
        wb.save(self.path)
        wb.close()

    def append_row(self, sheet: str, row: Dict[str, Any]) -> None:
        self.ensure_workbook()
        with _file_lock(self.path):
            wb = load_workbook(self.path)
            try:
                if sheet not in wb.sheetnames:
                    ws = wb.create_sheet(sheet)
                    ws.append(DEFAULT_SHEETS.get(sheet, list(row.keys())))
                ws = wb[sheet]
                headers = [c.value for c in ws[1]]
                out = []
                for h in headers:
                    out.append(row.get(h))
                # Add extra keys (if any) as appended columns (once).
                extras = [k for k in row.keys() if k not in headers]
                if extras:
                    for k in extras:
                        headers.append(k)
                        ws.cell(row=1, column=len(headers)).value = k
                        out.append(row.get(k))
                ws.append(out)
                wb.save(self.path)
            finally:
                wb.close()

    def log_message(self, *, session_id: str, role: str, content_redacted: str) -> None:
        self.append_row(
            "Messages",
            {
                "ts": round(time.time(), 3),
                "session_id": session_id,
                "role": role,
                "content_redacted": content_redacted,
            },
        )

    def log_action(self, *, session_id: str, action_name: str, payload: Dict[str, Any], result: Dict[str, Any], ok: bool) -> None:
        self.append_row(
            "Actions",
            {
                "ts": round(time.time(), 3),
                "session_id": session_id,
                "action_name": action_name,
                "payload_json": json.dumps(payload, ensure_ascii=False),
                "result_json": json.dumps(result, ensure_ascii=False),
                "ok": bool(ok),
            },
        )

    def log_outbox(self, *, session_id: str, channel: str, to: str, body: str, link: str) -> None:
        self.append_row(
            "Outbox",
            {
                "ts": round(time.time(), 3),
                "session_id": session_id,
                "channel": channel,
                "to": to,
                "body": body,
                "link": link,
            },
        )

    def upsert_call(self, *, session_id: str, customer_id: Optional[str], start_ts: Optional[float] = None, end_ts: Optional[float] = None, disposition: Optional[str] = None, ptp_date: Optional[str] = None, callback_time: Optional[str] = None) -> None:
        # For the demo we append snapshots rather than true upserts (simpler & auditable).
        self.append_row(
            "Calls",
            {
                "session_id": session_id,
                "customer_id": customer_id,
                "start_ts": start_ts,
                "end_ts": end_ts,
                "disposition": disposition,
                "ptp_date": ptp_date,
                "callback_time": callback_time,
            },
        )

