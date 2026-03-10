import hashlib
import hmac
import json
import os
import tempfile
import time
import unittest
from contextlib import suppress
from datetime import datetime, timedelta

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import web_app


class V2PlatformApiTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db_path = os.path.join(self.tmp.name, "demo.db")
        self.output_path = os.path.join(self.tmp.name, "demo_output.xlsx")

    def tearDown(self):
        self._reset_app_state()
        self.tmp.cleanup()

    def _reset_app_state(self):
        with suppress(Exception):
            task = getattr(web_app, "_report_scheduler_task", None)
            if task and not task.done():
                task.cancel()
        web_app._report_scheduler_task = None
        web_app._demo_state.clear()

    def _client(self, *, webhook_secret: str | None = None) -> TestClient:
        os.environ["DEMO_DB_PATH"] = self.db_path
        os.environ["DEMO_OUTPUT_XLSX_PATH"] = self.output_path
        os.environ["APP_ENV"] = "demo"
        os.environ["DEMO_MODE"] = "1"
        os.environ["PILOT_MODE"] = "1"
        os.environ["JWT_SECRET"] = "test-secret"
        os.environ["COOKIE_SECURE"] = "0"
        if webhook_secret:
            os.environ["V2_WEBHOOK_SECRET"] = webhook_secret
        else:
            os.environ.pop("V2_WEBHOOK_SECRET", None)
        self._reset_app_state()
        return TestClient(web_app.app)

    @staticmethod
    def _login(client: TestClient, username: str, password: str) -> str:
        res = client.post("/api/auth/login", json={"username": username, "password": password})
        assert res.status_code == 200, res.text
        token = str(res.json().get("access_token") or "")
        assert token
        return token

    @staticmethod
    def _headers(token: str, tenant_id: str) -> dict:
        return {
            "Authorization": f"Bearer {token}",
            "x-tenant-id": tenant_id,
            "x-request-id": "req-test-v2",
        }

    @staticmethod
    def _webhook_headers(*, secret: str, body_text: str, timestamp: int | None = None) -> dict:
        ts = int(time.time()) if timestamp is None else int(timestamp)
        signed = f"{ts}.{body_text}".encode("utf-8")
        sig = hmac.new(secret.encode("utf-8"), signed, hashlib.sha256).hexdigest()
        return {
            "x-webhook-timestamp": str(ts),
            "x-webhook-signature": sig,
            "content-type": "application/json",
        }

    def test_v2_tenant_journey_channel_payment_and_settlement_flow(self):
        with self._client() as client:
            admin_token = self._login(client, "admin", "admin123")

            create_tenant = client.post(
                "/api/v2/tenants",
                json={"tenant_id": "alpha", "name": "Alpha NBFC", "settings": {"timezone": "Asia/Kolkata"}},
                headers={"Authorization": f"Bearer {admin_token}"},
            )
            self.assertEqual(create_tenant.status_code, 200, create_tenant.text)
            self.assertEqual(create_tenant.json()["tenant"]["tenant_id"], "alpha")

            patch_settings = client.patch(
                "/api/v2/tenants/alpha/settings",
                json={"settings": {"channels": {"voice": True, "whatsapp": True, "sms": True, "email": True}}},
                headers={"Authorization": f"Bearer {admin_token}"},
            )
            self.assertEqual(patch_settings.status_code, 200, patch_settings.text)

            headers = self._headers(admin_token, "alpha")
            imported = client.post(
                "/api/v2/customers/import",
                json={
                    "rows": [
                        {
                            "customer_id": "CUST-ALPHA-1",
                            "full_name": "Asha Rao",
                            "phone": "+919999999901",
                            "email": "asha@example.com",
                            "dpd": 54,
                            "risk_band": "MEDIUM",
                        }
                    ]
                },
                headers=headers,
            )
            self.assertEqual(imported.status_code, 200, imported.text)
            self.assertEqual(imported.json()["result"]["inserted"], 1)

            loan = client.post(
                "/api/v2/loan_accounts/upsert",
                json={
                    "loan_account_id": "LN-ALPHA-1",
                    "customer_id": "CUST-ALPHA-1",
                    "principal_outstanding": 25000,
                    "emi_amount": 3200,
                    "dpd": 54,
                    "status": "ACTIVE",
                    "product_type": "PL",
                },
                headers=headers,
            )
            self.assertEqual(loan.status_code, 200, loan.text)
            self.assertEqual(loan.json()["loan_account"]["loan_account_id"], "LN-ALPHA-1")

            start = client.post(
                "/api/v2/journeys/start",
                json={
                    "customer_id": "CUST-ALPHA-1",
                    "loan_account_id": "LN-ALPHA-1",
                    "channel_hint": "whatsapp",
                    "context": {"bucket": "31-60"},
                },
                headers=headers,
            )
            self.assertEqual(start.status_code, 200, start.text)
            journey = start.json()["journey"]
            journey_id = journey["id"]

            sent = client.post(
                "/api/v2/channels/send",
                json={
                    "journey_id": journey_id,
                    "customer_id": "CUST-ALPHA-1",
                    "loan_account_id": "LN-ALPHA-1",
                    "channel": "whatsapp",
                    "content": "Your EMI is overdue. Can we help you close today?",
                },
                headers=headers,
            )
            self.assertEqual(sent.status_code, 200, sent.text)
            self.assertEqual(sent.json()["message"]["direction"], "outbound")

            listed = client.get("/api/v2/channels/messages?customer_id=CUST-ALPHA-1", headers=headers)
            self.assertEqual(listed.status_code, 200, listed.text)
            self.assertGreaterEqual(listed.json()["total"], 1)

            intent_resp = client.post(
                "/api/v2/payments/intents",
                json={
                    "customer_id": "CUST-ALPHA-1",
                    "loan_account_id": "LN-ALPHA-1",
                    "journey_id": journey_id,
                    "amount": 3200,
                    "currency": "INR",
                    "rail": "upi",
                },
                headers=headers,
            )
            self.assertEqual(intent_resp.status_code, 200, intent_resp.text)
            intent_id = intent_resp.json()["payment_intent"]["id"]

            webhook = client.post(
                "/api/v2/payments/webhooks/upi",
                json={
                    "tenant_id": "alpha",
                    "provider_event_id": "evt-pay-1",
                    "payment_intent_id": intent_id,
                    "status": "success",
                    "event_type": "payment.captured",
                },
            )
            self.assertEqual(webhook.status_code, 200, webhook.text)
            self.assertEqual(webhook.json()["event"]["status"], "SUCCEEDED")

            offer_resp = client.post(
                "/api/v2/settlements/offers",
                json={
                    "customer_id": "CUST-ALPHA-1",
                    "loan_account_id": "LN-ALPHA-1",
                    "journey_id": journey_id,
                    "offered_amount": 21000,
                    "original_due_amount": 25000,
                    "terms": {"installments": 2},
                },
                headers=headers,
            )
            self.assertEqual(offer_resp.status_code, 200, offer_resp.text)
            offer_id = offer_resp.json()["offer"]["id"]

            accept_resp = client.post(
                f"/api/v2/settlements/{offer_id}/accept",
                json={"accepted_amount": 21000, "acceptance_channel": "voice"},
                headers=headers,
            )
            self.assertEqual(accept_resp.status_code, 200, accept_resp.text)
            self.assertEqual(accept_resp.json()["offer"]["status"], "ACCEPTED")

            c360 = client.get("/api/v2/customers/CUST-ALPHA-1/360", headers=headers)
            self.assertEqual(c360.status_code, 200, c360.text)
            payload = c360.json()
            self.assertEqual(payload["customer_id"], "CUST-ALPHA-1")
            self.assertGreaterEqual(len(payload["loan_accounts"]), 1)
            self.assertGreaterEqual(len(payload["channel_messages"]), 1)
            self.assertGreaterEqual(len(payload["payment_intents"]), 1)
            self.assertGreaterEqual(len(payload["settlement_offers"]), 1)

            wrong_tenant = client.get(
                f"/api/v2/journeys/{journey_id}",
                headers=self._headers(admin_token, "default"),
            )
            self.assertEqual(wrong_tenant.status_code, 404)

    def test_v2_recovery_experiments_and_approval_queue(self):
        with self._client() as client:
            admin_token = self._login(client, "admin", "admin123")
            create_tenant = client.post(
                "/api/v2/tenants",
                json={"tenant_id": "beta", "name": "Beta Finance"},
                headers={"Authorization": f"Bearer {admin_token}"},
            )
            self.assertEqual(create_tenant.status_code, 200, create_tenant.text)

            headers = self._headers(admin_token, "beta")

            exp_resp = client.post(
                "/api/v2/recovery/experiments",
                json={
                    "name": "Settlement Offer Message Test",
                    "objective": "Improve commitment conversion",
                    "control": "control",
                    "treatment": "new_offer_copy",
                    "split_ratio": 0.5,
                    "unit_id": "CUST-BETA-1",
                },
                headers=headers,
            )
            self.assertEqual(exp_resp.status_code, 200, exp_resp.text)
            experiment_id = exp_resp.json()["experiment"]["id"]
            self.assertIsNotNone(exp_resp.json()["assignment"])

            decide = client.post(
                "/api/v2/recovery/nba/decide",
                json={
                    "customer_id": "CUST-BETA-1",
                    "loan_account_id": "LN-BETA-1",
                    "journey_id": "JNY-BETA-1",
                    "experiment_id": experiment_id,
                    "amount_due": 50000,
                    "features": {
                        "dpd": 132,
                        "amount_due": 50000,
                        "hardship": True,
                        "last_outcome": "no_answer",
                    },
                },
                headers=headers,
            )
            self.assertEqual(decide.status_code, 200, decide.text)
            decision = decide.json()["decision"]
            self.assertTrue(decision["approval_required"])
            approval_id = decision.get("approval_id")
            self.assertTrue(approval_id)

            queue = client.get("/api/v2/approvals/queue?status=PENDING", headers=headers)
            self.assertEqual(queue.status_code, 200, queue.text)
            self.assertGreaterEqual(queue.json()["total"], 1)

            approve = client.post(
                f"/api/v2/approvals/{approval_id}/approve",
                json={"note": "Risk reviewed by supervisor"},
                headers=headers,
            )
            self.assertEqual(approve.status_code, 200, approve.text)
            self.assertEqual(approve.json()["approval"]["status"], "APPROVED")

            uplift = client.get(f"/api/v2/recovery/uplift?experiment_id={experiment_id}", headers=headers)
            self.assertEqual(uplift.status_code, 200, uplift.text)
            self.assertEqual(uplift.json()["tenant_id"], "beta")
            self.assertGreaterEqual(len(uplift.json()["snapshots"]), 1)

    def test_webhook_signature_validation_and_replay_protection(self):
        secret = "test-hook-secret"
        with self._client(webhook_secret=secret) as client:
            admin_token = self._login(client, "admin", "admin123")
            create_tenant = client.post(
                "/api/v2/tenants",
                json={"tenant_id": "alpha", "name": "Alpha NBFC"},
                headers={"Authorization": f"Bearer {admin_token}"},
            )
            self.assertEqual(create_tenant.status_code, 200, create_tenant.text)
            headers = self._headers(admin_token, "alpha")
            intent_resp = client.post(
                "/api/v2/payments/intents",
                json={
                    "customer_id": "CUST-ALPHA-1",
                    "loan_account_id": "LN-ALPHA-1",
                    "journey_id": "JNY-ALPHA-1",
                    "amount": 3200,
                    "currency": "INR",
                    "rail": "upi",
                },
                headers=headers,
            )
            self.assertEqual(intent_resp.status_code, 200, intent_resp.text)
            intent_id = intent_resp.json()["payment_intent"]["id"]

            payload = {
                "tenant_id": "alpha",
                "provider_event_id": "evt-pay-42",
                "payment_intent_id": intent_id,
                "status": "success",
                "event_type": "payment.captured",
            }
            body_text = json.dumps(payload, separators=(",", ":"))

            bad_sig = client.post(
                "/api/v2/payments/webhooks/upi",
                data=body_text,
                headers={
                    "x-webhook-timestamp": str(int(time.time())),
                    "x-webhook-signature": "deadbeef",
                    "content-type": "application/json",
                },
            )
            self.assertEqual(bad_sig.status_code, 401, bad_sig.text)

            stale = client.post(
                "/api/v2/payments/webhooks/upi",
                data=body_text,
                headers=self._webhook_headers(secret=secret, body_text=body_text, timestamp=int(time.time()) - 301),
            )
            self.assertEqual(stale.status_code, 401, stale.text)

            ok = client.post(
                "/api/v2/payments/webhooks/upi",
                data=body_text,
                headers=self._webhook_headers(secret=secret, body_text=body_text),
            )
            self.assertEqual(ok.status_code, 200, ok.text)

            replay = client.post(
                "/api/v2/payments/webhooks/upi",
                data=body_text,
                headers=self._webhook_headers(secret=secret, body_text=body_text),
            )
            self.assertEqual(replay.status_code, 200, replay.text)
            self.assertTrue(bool(replay.json().get("deduplicated")))

    def test_metrics_read_only_and_ops_tick_mutates(self):
        with self._client() as client:
            admin_token = self._login(client, "admin", "admin123")
            auth = {"Authorization": f"Bearer {admin_token}"}
            demo = web_app._get_demo_singletons()
            followups = demo["followups"]
            followups.schedule_ptp_followups(
                session_id="sess-metrics-1",
                customer_id="CUST-METRICS-1",
                ptp_date="2000-01-01",
                phone="+919999999999",
                channel="whatsapp",
            )
            due_before = followups.due_followups(limit=20)
            self.assertGreaterEqual(len(due_before), 1)

            metrics = client.get("/api/metrics", headers=auth)
            self.assertEqual(metrics.status_code, 200, metrics.text)
            due_after_metrics = followups.due_followups(limit=20)
            self.assertEqual(len(due_after_metrics), len(due_before))

            tick = client.post("/api/ops/tick", headers=auth)
            self.assertEqual(tick.status_code, 200, tick.text)
            due_after_tick = followups.due_followups(limit=20)
            self.assertEqual(len(due_after_tick), 0)

    def test_task_ptp_update_schedules_followups_and_updates_metrics(self):
        with self._client() as client:
            admin_token = self._login(client, "admin", "admin123")
            auth = {"Authorization": f"Bearer {admin_token}"}
            demo = web_app._get_demo_singletons()
            campaign = demo["campaign_service"]
            workbench = demo["workbench"]

            created = campaign.create_campaign(name="Buyer Demo", customer_ids=["BUY-1"], max_attempts=2, retry_delay_minutes=15, batch_size=20)
            campaign_id = created["campaign_id"]
            workbench.seed_tasks(
                campaign_id=campaign_id,
                portfolio_id="pfl-buyer",
                rows=[
                    {
                        "customer_id": "BUY-1",
                        "phone": "+919999999981",
                        "amount_due": 1750,
                        "dpd": 18,
                        "customer_name": "Buyer Demo Customer",
                    }
                ],
                actor="seed",
            )
            task = workbench.list_tasks(campaign_id=campaign_id, page=1, page_size=10)["rows"][0]

            update = client.post(
                f"/api/tasks/{task['id']}/update",
                json={
                    "state": "PTP",
                    "disposition": "ptp_captured",
                    "ptp_date": "2026-03-12",
                    "notes": "manual demo commitment",
                },
                headers=auth,
            )
            self.assertEqual(update.status_code, 200, update.text)
            payload = update.json()
            self.assertTrue(payload["ok"])
            self.assertTrue(payload.get("session_id"))
            self.assertEqual(payload["task"]["ptp_date"], "2026-03-12")
            self.assertEqual(len(payload.get("followups") or []), 3)

            metrics = client.get(f"/api/metrics?campaign_id={campaign_id}", headers=auth)
            self.assertEqual(metrics.status_code, 200, metrics.text)
            metrics_payload = metrics.json()
            self.assertEqual(metrics_payload["accounts_assigned"], 1)
            self.assertEqual(metrics_payload["accounts_contacted"], 1)
            self.assertEqual(metrics_payload["expected_recovery_amount"], 1750.0)
            self.assertEqual(metrics_payload["followups_scheduled_total"], 3)
            self.assertEqual(metrics_payload["queue_snapshot"]["PTP"], 1)

            wb = load_workbook(self.output_path)
            try:
                calls = wb["Calls"]
                headers = [cell.value for cell in calls[1]]
                call_rows = [dict(zip(headers, row)) for row in calls.iter_rows(min_row=2, values_only=True)]
            finally:
                wb.close()
            matching = [row for row in call_rows if row.get("session_id") == payload["session_id"]]
            self.assertGreaterEqual(len(matching), 1)
            self.assertEqual(matching[-1]["ptp_date"], "2026-03-12")
            self.assertEqual(matching[-1]["disposition"], "ptp_captured")

    def test_portfolio_to_agent_to_dashboard_and_ops_copilot_flow(self):
        with self._client() as client:
            mgr_token = self._login(client, "mgr", "mgr123")
            agent_token = self._login(client, "agent", "agent123")
            mgr_headers = self._headers(mgr_token, "default")
            agent_headers = self._headers(agent_token, "default")

            csv_text = "\n".join(
                [
                    "customer_id,customer_name,phone,dpd,initial_amount,remaining_amount,language,due_date",
                    "FLOW-1,Asha Rao,9876543210,42,100000,22000,en,2026-03-14",
                    "FLOW-2,Rohan Verma,9876543211,68,135000,41000,hi,2026-03-16",
                ]
            )
            upload = client.post(
                "/api/portfolio/upload",
                files={"file": ("portfolio.csv", csv_text.encode("utf-8"), "text/csv")},
                headers=mgr_headers,
            )
            self.assertEqual(upload.status_code, 200, upload.text)
            upload_payload = upload.json()
            upload_id = upload_payload["upload_id"]

            mapped = client.post(
                f"/api/portfolio/{upload_id}/map",
                json={
                    "portfolio_name": "Ops Proof Portfolio",
                    "mappings": {
                        "customer_id": "customer_id",
                        "customer_name": "customer_name",
                        "phone": "phone",
                        "dpd": "dpd",
                        "initial_amount": "initial_amount",
                        "remaining_amount": "remaining_amount",
                        "language": "language",
                        "due_date": "due_date",
                    },
                },
                headers=mgr_headers,
            )
            self.assertEqual(mapped.status_code, 200, mapped.text)
            portfolio_id = mapped.json()["portfolio_id"]

            validate = client.post(f"/api/portfolio/{portfolio_id}/validate", headers=mgr_headers)
            self.assertEqual(validate.status_code, 200, validate.text)
            validate_payload = validate.json()
            self.assertEqual(validate_payload["valid_rows"], 2)

            preview = client.get(f"/api/portfolio/{portfolio_id}/preview?limit=5", headers=mgr_headers)
            self.assertEqual(preview.status_code, 200, preview.text)
            preview_rows = preview.json()["rows"]
            self.assertEqual(preview_rows[0]["initial_amount"], 100000.0)
            self.assertEqual(preview_rows[0]["remaining_amount"], 22000.0)
            self.assertEqual(preview_rows[0]["amount_due"], 22000.0)

            launch = client.post(
                f"/api/portfolio/{portfolio_id}/launch",
                json={"campaign_name": "Ops Proof Campaign", "retry_policy": {"max_attempts": 2, "retry_delay_minutes": 15}, "throttle": 25},
                headers=mgr_headers,
            )
            self.assertEqual(launch.status_code, 200, launch.text)
            launch_payload = launch.json()
            campaign_id = launch_payload["campaign_id"]
            self.assertEqual(launch_payload["seeded_accounts"], 2)
            self.assertEqual(launch_payload["tasks_created"], 2)

            demo = web_app._get_demo_singletons()
            workbench = demo["workbench"]
            audit = demo["audit"]

            tasks = workbench.list_tasks(campaign_id=campaign_id, page=1, page_size=10)["rows"]
            task_by_customer = {row["customer_id"]: row for row in tasks}
            callback_at = (datetime.now() + timedelta(hours=2)).isoformat()

            callback_call = client.post(
                "/api/telephony/agent_call",
                json={
                    "phone": task_by_customer["FLOW-1"]["phone"],
                    "customer_name": task_by_customer["FLOW-1"]["customer_name"],
                    "customer_id": "FLOW-1",
                    "campaign_id": campaign_id,
                    "amount_due": str(task_by_customer["FLOW-1"]["amount_due"]),
                    "language": "en",
                    "tts_speaker": "meera",
                },
                headers=agent_headers,
            )
            self.assertEqual(callback_call.status_code, 200, callback_call.text)
            callback_call_payload = callback_call.json()
            self.assertTrue(callback_call_payload["ok"])
            self.assertTrue(callback_call_payload.get("session_id"))

            ptp_call = client.post(
                "/api/telephony/agent_call",
                json={
                    "phone": task_by_customer["FLOW-2"]["phone"],
                    "customer_name": task_by_customer["FLOW-2"]["customer_name"],
                    "customer_id": "FLOW-2",
                    "campaign_id": campaign_id,
                    "amount_due": str(task_by_customer["FLOW-2"]["amount_due"]),
                    "language": "hi",
                    "tts_speaker": "shubh",
                },
                headers=agent_headers,
            )
            self.assertEqual(ptp_call.status_code, 200, ptp_call.text)
            ptp_call_payload = ptp_call.json()
            self.assertTrue(ptp_call_payload["ok"])
            self.assertTrue(ptp_call_payload.get("session_id"))

            callback_update = client.post(
                f"/api/tasks/{task_by_customer['FLOW-1']['id']}/update",
                json={
                    "state": "CALLBACK",
                    "disposition": "callback_scheduled",
                    "callback_at": callback_at,
                    "session_id": callback_call_payload["session_id"],
                    "notes": "Customer requested afternoon callback",
                },
                headers=agent_headers,
            )
            self.assertEqual(callback_update.status_code, 200, callback_update.text)
            callback_payload = callback_update.json()
            self.assertTrue(callback_payload["ok"])
            self.assertEqual(len(callback_payload.get("followups") or []), 1)

            ptp_date = (datetime.now() + timedelta(days=2)).date().isoformat()
            ptp_update = client.post(
                f"/api/tasks/{task_by_customer['FLOW-2']['id']}/update",
                json={
                    "state": "PTP",
                    "disposition": "ptp_captured",
                    "ptp_date": ptp_date,
                    "session_id": ptp_call_payload["session_id"],
                    "notes": "Customer committed after salary credit",
                },
                headers=agent_headers,
            )
            self.assertEqual(ptp_update.status_code, 200, ptp_update.text)
            ptp_payload = ptp_update.json()
            self.assertTrue(ptp_payload["ok"])
            self.assertEqual(len(ptp_payload.get("followups") or []), 3)

            agent_queue = client.get(f"/api/tasks?campaign_id={campaign_id}", headers=agent_headers)
            self.assertEqual(agent_queue.status_code, 200, agent_queue.text)
            agent_rows = agent_queue.json()["rows"]
            self.assertEqual(len(agent_rows), 2)
            self.assertTrue(all(str(row.get("owner") or "") == "agent" for row in agent_rows))

            audit.record_violation(
                session_id=callback_call_payload["session_id"],
                kind="profanity",
                detail="Customer used abusive language before accepting a callback",
            )

            metrics = client.get(f"/api/metrics?campaign_id={campaign_id}", headers=mgr_headers)
            self.assertEqual(metrics.status_code, 200, metrics.text)
            metrics_payload = metrics.json()
            self.assertEqual(metrics_payload["accounts_assigned"], 2)
            self.assertEqual(metrics_payload["accounts_contacted"], 2)
            self.assertEqual(metrics_payload["callback_count"], 1)
            self.assertEqual(metrics_payload["ptp_count"], 1)
            self.assertEqual(metrics_payload["expected_recovery_amount"], 41000.0)
            self.assertEqual(metrics_payload["followups_scheduled_total"], 4)
            self.assertEqual(metrics_payload["queue_snapshot"]["CALLBACK"], 1)
            self.assertEqual(metrics_payload["queue_snapshot"]["PTP"], 1)
            self.assertEqual(metrics_payload["profanity_incidents"], 1)
            self.assertGreaterEqual(metrics_payload["ended_sessions"], 2)

            roll = client.get(f"/api/metrics/roll-forward?campaign_id={campaign_id}&days=30", headers=mgr_headers)
            self.assertEqual(roll.status_code, 200, roll.text)
            roll_payload = roll.json()
            self.assertEqual(roll_payload["total_transitions"], 2)
            self.assertEqual(roll_payload["cure_count"], 0)
            self.assertEqual(roll_payload["roll_forward_count"], 0)
            self.assertEqual(roll_payload["matrix"]["31-60"]["31-60"], 1)
            self.assertEqual(roll_payload["matrix"]["61-90"]["61-90"], 1)

            wb = load_workbook(self.output_path)
            try:
                calls = wb["Calls"]
                headers = [cell.value for cell in calls[1]]
                call_rows = [dict(zip(headers, row)) for row in calls.iter_rows(min_row=2, values_only=True)]
            finally:
                wb.close()
            callback_rows = [row for row in call_rows if row.get("session_id") == callback_call_payload["session_id"]]
            ptp_rows = [row for row in call_rows if row.get("session_id") == ptp_call_payload["session_id"]]
            self.assertGreaterEqual(len(callback_rows), 2)
            self.assertGreaterEqual(len(ptp_rows), 2)
            self.assertTrue(any(row.get("callback_time") for row in callback_rows))
            self.assertTrue(any(row.get("ptp_date") == ptp_date for row in ptp_rows))

            copilot = client.post(
                "/api/control-layer/chat",
                json={"campaign_id": campaign_id, "message": "Which bucket is underperforming, why, and what should I change?"},
                headers=mgr_headers,
            )
            self.assertEqual(copilot.status_code, 200, copilot.text)
            copilot_payload = copilot.json()
            self.assertEqual(copilot_payload["scope"]["campaign_id"], campaign_id)
            self.assertGreaterEqual(len(copilot_payload["evidence"]), 2)
            self.assertGreaterEqual(len(copilot_payload["recommendations"]), 1)

            first_rec = copilot_payload["recommendations"][0]
            exp_action = [row for row in first_rec["actions"] if row["action_type"] == "launch_experiment"][0]
            exp_response = client.post(
                "/api/control-layer/actions",
                json={"action_type": exp_action["action_type"], "payload": exp_action["payload"]},
                headers=mgr_headers,
            )
            self.assertEqual(exp_response.status_code, 200, exp_response.text)
            self.assertEqual(exp_response.json()["result_type"], "experiment_created")

            strategy_action = [row for row in first_rec["actions"] if row["action_type"] == "request_strategy_change"][0]
            strategy_response = client.post(
                "/api/control-layer/actions",
                json={"action_type": strategy_action["action_type"], "payload": strategy_action["payload"]},
                headers=mgr_headers,
            )
            self.assertEqual(strategy_response.status_code, 200, strategy_response.text)
            self.assertEqual(strategy_response.json()["result_type"], "approval_queued")


if __name__ == "__main__":
    unittest.main()
