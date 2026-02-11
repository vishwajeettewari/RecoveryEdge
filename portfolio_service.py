from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import sqlite3
import time
import uuid
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import load_workbook


REQUIRED_FIELDS = ["customer_id", "phone", "amount_due", "dpd"]
OPTIONAL_FIELDS = ["due_date", "language", "customer_name"]
ALLOWED_FIELDS = set(REQUIRED_FIELDS + OPTIONAL_FIELDS)


class PortfolioService:
    def __init__(self, db_path: str, data_dir: str) -> None:
        self.db_path = db_path
        self.base_dir = os.path.join(data_dir, "portfolio_uploads")
        os.makedirs(self.base_dir, exist_ok=True)
        self._init_db()

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.db_path, check_same_thread=False)
        conn.row_factory = sqlite3.Row
        return conn

    def _init_db(self) -> None:
        conn = self._connect()
        try:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS portfolios (
                    id TEXT PRIMARY KEY,
                    name TEXT NOT NULL,
                    created_at REAL NOT NULL,
                    source_type TEXT NOT NULL,
                    status TEXT NOT NULL,
                    row_count INTEGER NOT NULL DEFAULT 0,
                    schema_hash TEXT,
                    validation_summary_json TEXT,
                    data_quality_score INTEGER DEFAULT 0,
                    error_csv_path TEXT
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS portfolio_uploads (
                    id TEXT PRIMARY KEY,
                    portfolio_id TEXT NOT NULL,
                    original_filename TEXT NOT NULL,
                    uploaded_at REAL NOT NULL,
                    raw_path TEXT NOT NULL,
                    parsed_path TEXT NOT NULL,
                    validation_summary_json TEXT
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS portfolio_columns (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    portfolio_id TEXT NOT NULL,
                    source_col TEXT NOT NULL,
                    mapped_field TEXT,
                    required INTEGER NOT NULL DEFAULT 0,
                    inferred_type TEXT,
                    sample_values_json TEXT
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS portfolio_rows (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    portfolio_id TEXT NOT NULL,
                    row_index INTEGER NOT NULL,
                    customer_id TEXT,
                    phone TEXT,
                    amount_due REAL,
                    dpd INTEGER,
                    due_date TEXT,
                    language TEXT,
                    customer_name TEXT,
                    extra_json TEXT,
                    is_valid INTEGER,
                    is_duplicate INTEGER DEFAULT 0,
                    normalized_phone TEXT,
                    error_json TEXT
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS campaign_launches (
                    id TEXT PRIMARY KEY,
                    portfolio_id TEXT NOT NULL,
                    campaign_id TEXT NOT NULL,
                    launched_at REAL NOT NULL,
                    launch_config_json TEXT,
                    status TEXT NOT NULL
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS portfolio_exclusions (
                    id TEXT PRIMARY KEY,
                    portfolio_id TEXT NOT NULL,
                    uploaded_at REAL NOT NULL,
                    path TEXT NOT NULL,
                    row_count INTEGER NOT NULL DEFAULT 0
                )
                """
            )
            conn.commit()
        finally:
            conn.close()

    def upload(self, *, filename: str, content: bytes) -> Dict[str, Any]:
        source_type = self._source_type(filename)
        if source_type not in {"csv", "xlsx"}:
            raise ValueError("Only .csv and .xlsx uploads are supported")

        upload_id = f"upl-{uuid.uuid4().hex[:12]}"
        portfolio_id = f"pfl-{uuid.uuid4().hex[:12]}"
        ts = time.time()
        folder = os.path.join(self.base_dir, upload_id)
        os.makedirs(folder, exist_ok=True)
        raw_path = os.path.join(folder, f"raw-{self._safe_name(filename)}")
        parsed_path = os.path.join(folder, "parsed.jsonl")

        with open(raw_path, "wb") as f:
            f.write(content)

        rows = self._parse_file(raw_path, source_type)
        self._write_jsonl(parsed_path, rows)

        columns = self._infer_columns(rows)
        schema_hash = self._schema_hash([c["source_col"] for c in columns])

        conn = self._connect()
        try:
            conn.execute(
                """
                INSERT INTO portfolios (id, name, created_at, source_type, status, row_count, schema_hash)
                VALUES (?, ?, ?, ?, 'uploaded', ?, ?)
                """,
                (portfolio_id, portfolio_id, ts, source_type, len(rows), schema_hash),
            )
            conn.execute(
                """
                INSERT INTO portfolio_uploads (id, portfolio_id, original_filename, uploaded_at, raw_path, parsed_path)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (upload_id, portfolio_id, filename, ts, raw_path, parsed_path),
            )
            for c in columns:
                conn.execute(
                    """
                    INSERT INTO portfolio_columns (portfolio_id, source_col, mapped_field, required, inferred_type, sample_values_json)
                    VALUES (?, ?, NULL, 0, ?, ?)
                    """,
                    (portfolio_id, c["source_col"], c["inferred_type"], json.dumps(c["sample_values"], ensure_ascii=False)),
                )
            conn.commit()
        finally:
            conn.close()

        return {
            "upload_id": upload_id,
            "portfolio_id": portfolio_id,
            "source_type": source_type,
            "columns": columns,
            "sample_preview": rows[:20],
            "row_count": len(rows),
        }

    def map_columns(self, *, upload_id: str, mappings: Dict[str, str], portfolio_name: str) -> Dict[str, Any]:
        upload = self._get_upload(upload_id)
        if not upload:
            raise ValueError("upload_id_not_found")
        portfolio_id = str(upload["portfolio_id"])

        mapped_values = {v for v in mappings.values() if v}
        missing = [f for f in REQUIRED_FIELDS if f not in mapped_values]
        if missing:
            raise ValueError(f"required_mappings_missing:{','.join(missing)}")

        rows = self._read_jsonl(str(upload["parsed_path"]))

        conn = self._connect()
        mapped_count = 0
        try:
            conn.execute("DELETE FROM portfolio_rows WHERE portfolio_id = ?", (portfolio_id,))
            conn.execute("DELETE FROM portfolio_columns WHERE portfolio_id = ?", (portfolio_id,))

            source_cols = list(rows[0].keys()) if rows else list(mappings.keys())
            for source_col in source_cols:
                mapped_field = mappings.get(source_col)
                required = 1 if mapped_field in REQUIRED_FIELDS else 0
                inferred_type = self._infer_value_type([r.get(source_col) for r in rows[:50]])
                sample_values = [str(r.get(source_col) or "") for r in rows[:5]]
                conn.execute(
                    """
                    INSERT INTO portfolio_columns (portfolio_id, source_col, mapped_field, required, inferred_type, sample_values_json)
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (portfolio_id, source_col, mapped_field, required, inferred_type, json.dumps(sample_values, ensure_ascii=False)),
                )

            for idx, row in enumerate(rows, start=1):
                mapped = {f: None for f in REQUIRED_FIELDS + OPTIONAL_FIELDS}
                for source_col, target in mappings.items():
                    if not target or target not in ALLOWED_FIELDS:
                        continue
                    mapped[target] = row.get(source_col)
                extra = {k: v for k, v in row.items() if k not in mappings or mappings.get(k) not in ALLOWED_FIELDS}
                conn.execute(
                    """
                    INSERT INTO portfolio_rows (portfolio_id, row_index, customer_id, phone, amount_due, dpd, due_date, language, customer_name, extra_json, is_valid, is_duplicate, normalized_phone, error_json)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, 0, NULL, NULL)
                    """,
                    (
                        portfolio_id,
                        idx,
                        self._to_text(mapped.get("customer_id")),
                        self._to_text(mapped.get("phone")),
                        self._to_float(mapped.get("amount_due")),
                        self._to_int(mapped.get("dpd")),
                        self._to_text(mapped.get("due_date")),
                        self._to_text(mapped.get("language")),
                        self._to_text(mapped.get("customer_name")),
                        json.dumps(extra, ensure_ascii=False),
                    ),
                )
                mapped_count += 1

            conn.execute(
                """
                UPDATE portfolios
                SET name = ?, status = 'mapped', row_count = ?
                WHERE id = ?
                """,
                (portfolio_name or portfolio_id, mapped_count, portfolio_id),
            )
            conn.commit()
        finally:
            conn.close()

        return {"portfolio_id": portfolio_id, "mapped_rows": mapped_count}

    def validate(self, *, portfolio_id: str) -> Dict[str, Any]:
        conn = self._connect()
        try:
            rows = conn.execute(
                """
                SELECT id, row_index, customer_id, phone, amount_due, dpd, due_date, language, customer_name, extra_json
                FROM portfolio_rows
                WHERE portfolio_id = ?
                ORDER BY id ASC
                """,
                (portfolio_id,),
            ).fetchall()

            latest_by_customer: Dict[str, int] = {}
            for r in rows:
                cid = (r["customer_id"] or "").strip()
                if cid:
                    latest_by_customer[cid] = int(r["id"])

            issues = {
                "missing_required_fields": 0,
                "invalid_phones": 0,
                "invalid_dpd": 0,
                "duplicates": 0,
                "missing_amount_due": 0,
            }
            bucket_counts = {"1-30": 0, "31-60": 0, "61-90": 0, "90+": 0, "0": 0}
            lang_dist: Dict[str, int] = {}
            valid_count = 0
            invalid_count = 0
            dup_count = 0
            error_rows: List[Dict[str, Any]] = []

            for r in rows:
                rid = int(r["id"])
                cid = (r["customer_id"] or "").strip()
                phone = (r["phone"] or "").strip()
                dpd_raw = r["dpd"]
                amount_raw = r["amount_due"]
                errs: List[str] = []
                is_duplicate = 0

                if cid and latest_by_customer.get(cid) != rid:
                    is_duplicate = 1
                    dup_count += 1
                    issues["duplicates"] += 1
                    errs.append("duplicate_customer_id_dropped")

                missing_required = False
                if not cid:
                    missing_required = True
                if not phone:
                    missing_required = True
                if dpd_raw is None:
                    missing_required = True
                if amount_raw is None:
                    missing_required = True
                if missing_required:
                    issues["missing_required_fields"] += 1
                    errs.append("missing_required_fields")

                normalized_phone = self._normalize_phone(phone)
                if phone and not normalized_phone:
                    issues["invalid_phones"] += 1
                    errs.append("invalid_phone")

                dpd_val = self._parse_int(dpd_raw)
                if dpd_raw is not None and (dpd_val is None or dpd_val < 0):
                    issues["invalid_dpd"] += 1
                    errs.append("invalid_dpd")

                amount_val = self._parse_float(amount_raw)
                if amount_raw is None:
                    issues["missing_amount_due"] += 1
                    errs.append("missing_amount_due")
                elif amount_val is None or amount_val < 0:
                    issues["missing_amount_due"] += 1
                    errs.append("invalid_amount_due")

                is_valid = 1 if (not errs) else 0
                if is_valid:
                    valid_count += 1
                    if dpd_val is None:
                        bucket_counts["0"] += 1
                    elif dpd_val <= 0:
                        bucket_counts["0"] += 1
                    elif dpd_val <= 30:
                        bucket_counts["1-30"] += 1
                    elif dpd_val <= 60:
                        bucket_counts["31-60"] += 1
                    elif dpd_val <= 90:
                        bucket_counts["61-90"] += 1
                    else:
                        bucket_counts["90+"] += 1
                    lang = (r["language"] or "unknown").strip() or "unknown"
                    lang_dist[lang] = int(lang_dist.get(lang, 0)) + 1
                else:
                    invalid_count += 1
                    error_rows.append(
                        {
                            "row_index": r["row_index"],
                            "customer_id": cid,
                            "phone": phone,
                            "dpd": dpd_raw,
                            "amount_due": amount_raw,
                            "errors": "|".join(errs),
                        }
                    )

                conn.execute(
                    """
                    UPDATE portfolio_rows
                    SET is_valid = ?,
                        is_duplicate = ?,
                        normalized_phone = ?,
                        error_json = ?,
                        phone = COALESCE(?, phone),
                        dpd = COALESCE(?, dpd),
                        amount_due = COALESCE(?, amount_due)
                    WHERE id = ?
                    """,
                    (
                        is_valid,
                        is_duplicate,
                        normalized_phone,
                        json.dumps(errs, ensure_ascii=False),
                        normalized_phone,
                        dpd_val,
                        amount_val,
                        rid,
                    ),
                )

            total = len(rows)
            dq = 100
            if issues["missing_required_fields"] > 0:
                dq -= 30
            if issues["invalid_phones"] > 0:
                dq -= 20
            if issues["invalid_dpd"] > 0:
                dq -= 20
            if issues["duplicates"] > 0:
                dq -= 10
            if issues["missing_amount_due"] > 0:
                dq -= 20
            if total > 0 and (invalid_count / float(total)) > 0.05:
                dq -= 20
            dq = max(0, dq)

            top_issues = sorted(
                [{"issue": k, "count": int(v)} for k, v in issues.items() if int(v) > 0],
                key=lambda x: x["count"],
                reverse=True,
            )[:5]

            errors_csv_path = self._write_error_csv(portfolio_id, error_rows)
            summary = {
                "portfolio_id": portfolio_id,
                "total_rows": total,
                "valid_rows": valid_count,
                "invalid_rows": invalid_count,
                "duplicates_dropped": dup_count,
                "bucket_counts": bucket_counts,
                "language_distribution": lang_dist,
                "issues": issues,
                "data_quality_score": dq,
                "top_issues": top_issues,
                "errors_csv_path": errors_csv_path,
            }

            conn.execute(
                """
                UPDATE portfolios
                SET status = 'validated', validation_summary_json = ?, data_quality_score = ?, error_csv_path = ?
                WHERE id = ?
                """,
                (json.dumps(summary, ensure_ascii=False), dq, errors_csv_path, portfolio_id),
            )
            conn.execute(
                """
                UPDATE portfolio_uploads
                SET validation_summary_json = ?
                WHERE portfolio_id = ?
                """,
                (json.dumps(summary, ensure_ascii=False), portfolio_id),
            )
            conn.commit()
            return summary
        finally:
            conn.close()

    def preview(self, *, portfolio_id: str, limit: int = 20) -> List[Dict[str, Any]]:
        conn = self._connect()
        try:
            rows = conn.execute(
                """
                SELECT customer_id, normalized_phone AS phone, amount_due, dpd, due_date, language, customer_name
                FROM portfolio_rows
                WHERE portfolio_id = ?
                  AND is_valid = 1
                  AND is_duplicate = 0
                ORDER BY row_index ASC
                LIMIT ?
                """,
                (portfolio_id, max(1, int(limit))),
            ).fetchall()
            return [dict(r) for r in rows]
        finally:
            conn.close()

    def add_exclusions(self, *, portfolio_id: str, filename: str, content: bytes) -> Dict[str, Any]:
        exclusion_id = f"pex-{uuid.uuid4().hex[:12]}"
        folder = os.path.join(self.base_dir, portfolio_id)
        os.makedirs(folder, exist_ok=True)
        path = os.path.join(folder, f"exclusions-{self._safe_name(filename)}")
        with open(path, "wb") as f:
            f.write(content)

        ids = self._read_exclusion_ids(path)
        conn = self._connect()
        try:
            conn.execute(
                """
                INSERT INTO portfolio_exclusions (id, portfolio_id, uploaded_at, path, row_count)
                VALUES (?, ?, ?, ?, ?)
                """,
                (exclusion_id, portfolio_id, time.time(), path, len(ids)),
            )
            conn.commit()
        finally:
            conn.close()
        return {"exclusion_id": exclusion_id, "row_count": len(ids)}

    def launch_candidates(
        self,
        *,
        portfolio_id: str,
        exclusion_id: Optional[str],
        exclude_predicate: Optional[str],
    ) -> Dict[str, Any]:
        conn = self._connect()
        try:
            rows = conn.execute(
                """
                SELECT * FROM portfolio_rows
                WHERE portfolio_id = ?
                  AND is_valid = 1
                  AND is_duplicate = 0
                ORDER BY row_index ASC
                """,
                (portfolio_id,),
            ).fetchall()

            excluded_customers = set()
            if exclusion_id:
                ex = conn.execute(
                    "SELECT path FROM portfolio_exclusions WHERE id = ? AND portfolio_id = ?",
                    (exclusion_id, portfolio_id),
                ).fetchone()
                if ex:
                    excluded_customers = set(self._read_exclusion_ids(str(ex["path"])))

            predicate = self._compile_predicate(exclude_predicate)

            selected: List[Dict[str, Any]] = []
            excluded: List[Dict[str, Any]] = []
            for row in rows:
                rec = dict(row)
                cid = str(rec.get("customer_id") or "").strip()
                drop = False
                if cid and cid in excluded_customers:
                    drop = True
                if not drop and predicate and predicate(rec):
                    drop = True
                out = {
                    "customer_id": cid,
                    "phone": rec.get("normalized_phone") or rec.get("phone"),
                    "amount_due": rec.get("amount_due"),
                    "dpd": rec.get("dpd"),
                    "due_date": rec.get("due_date"),
                    "language": rec.get("language"),
                    "customer_name": rec.get("customer_name"),
                }
                if drop:
                    excluded.append(out)
                else:
                    selected.append(out)

            return {
                "selected_rows": selected,
                "excluded_rows": excluded,
                "selected_count": len(selected),
                "excluded_count": len(excluded),
            }
        finally:
            conn.close()

    def record_launch(self, *, portfolio_id: str, campaign_id: str, launch_config: Dict[str, Any], status: str = "launched") -> str:
        launch_id = f"lch-{uuid.uuid4().hex[:12]}"
        conn = self._connect()
        try:
            conn.execute(
                """
                INSERT INTO campaign_launches (id, portfolio_id, campaign_id, launched_at, launch_config_json, status)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (launch_id, portfolio_id, campaign_id, time.time(), json.dumps(launch_config, ensure_ascii=False), status),
            )
            conn.commit()
        finally:
            conn.close()
        return launch_id

    def get_portfolio(self, portfolio_id: str) -> Optional[Dict[str, Any]]:
        conn = self._connect()
        try:
            row = conn.execute("SELECT * FROM portfolios WHERE id = ?", (portfolio_id,)).fetchone()
            return dict(row) if row else None
        finally:
            conn.close()

    def list_portfolios(self, limit: int = 50, offset: int = 0) -> Dict[str, Any]:
        conn = self._connect()
        try:
            total = conn.execute("SELECT COUNT(*) AS n FROM portfolios").fetchone()["n"]
            rows = conn.execute(
                "SELECT * FROM portfolios ORDER BY created_at DESC LIMIT ? OFFSET ?",
                (max(1, int(limit)), max(0, int(offset))),
            ).fetchall()
            return {"total": int(total or 0), "rows": [dict(r) for r in rows]}
        finally:
            conn.close()

    def get_errors_csv_path(self, portfolio_id: str) -> Optional[str]:
        conn = self._connect()
        try:
            row = conn.execute("SELECT error_csv_path FROM portfolios WHERE id = ?", (portfolio_id,)).fetchone()
            if not row:
                return None
            return row["error_csv_path"]
        finally:
            conn.close()

    def _get_upload(self, upload_id: str) -> Optional[Dict[str, Any]]:
        conn = self._connect()
        try:
            row = conn.execute("SELECT * FROM portfolio_uploads WHERE id = ?", (upload_id,)).fetchone()
            return dict(row) if row else None
        finally:
            conn.close()

    @staticmethod
    def _safe_name(filename: str) -> str:
        out = re.sub(r"[^A-Za-z0-9._-]+", "_", filename or "upload")
        return out[:120]

    @staticmethod
    def _source_type(filename: str) -> str:
        name = (filename or "").lower().strip()
        if name.endswith(".csv"):
            return "csv"
        if name.endswith(".xlsx"):
            return "xlsx"
        return ""

    def _parse_file(self, path: str, source_type: str) -> List[Dict[str, Any]]:
        if source_type == "csv":
            return self._parse_csv(path)
        return self._parse_xlsx(path)

    @staticmethod
    def _parse_csv(path: str) -> List[Dict[str, Any]]:
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            out = []
            for row in reader:
                if row is None:
                    continue
                cleaned = {str(k).strip(): (v.strip() if isinstance(v, str) else v) for k, v in row.items() if k is not None}
                out.append(cleaned)
            return out

    @staticmethod
    def _parse_xlsx(path: str) -> List[Dict[str, Any]]:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return []
            headers = [str(h or "").strip() for h in rows[0]]
            out: List[Dict[str, Any]] = []
            for r in rows[1:]:
                rec: Dict[str, Any] = {}
                for i, h in enumerate(headers):
                    if not h:
                        continue
                    val = r[i] if i < len(r) else None
                    if isinstance(val, str):
                        val = val.strip()
                    rec[h] = val
                if rec:
                    out.append(rec)
            return out
        finally:
            wb.close()

    @staticmethod
    def _write_jsonl(path: str, rows: List[Dict[str, Any]]) -> None:
        with open(path, "w", encoding="utf-8") as f:
            for row in rows:
                f.write(json.dumps(row, ensure_ascii=False) + "\n")

    @staticmethod
    def _read_jsonl(path: str) -> List[Dict[str, Any]]:
        out: List[Dict[str, Any]] = []
        with open(path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                out.append(json.loads(line))
        return out

    def _infer_columns(self, rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        cols: Dict[str, List[Any]] = {}
        for row in rows[:200]:
            for k, v in row.items():
                cols.setdefault(str(k).strip(), []).append(v)
        out: List[Dict[str, Any]] = []
        for col, vals in cols.items():
            sample = [str(v) for v in vals[:5] if v not in (None, "")]
            out.append(
                {
                    "source_col": col,
                    "inferred_type": self._infer_value_type(vals),
                    "sample_values": sample,
                }
            )
        return out

    @staticmethod
    def _infer_value_type(values: Iterable[Any]) -> str:
        vals = [v for v in values if v not in (None, "")]
        if not vals:
            return "string"
        numeric = 0
        date_like = 0
        for v in vals:
            s = str(v).strip()
            if re.fullmatch(r"[-+]?\d+(?:\.\d+)?", s):
                numeric += 1
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s) or re.fullmatch(r"\d{1,2}/\d{1,2}/\d{2,4}", s):
                date_like += 1
        if numeric == len(vals):
            return "number"
        if date_like >= max(1, int(len(vals) * 0.6)):
            return "date"
        return "string"

    @staticmethod
    def _schema_hash(cols: List[str]) -> str:
        raw = "|".join(sorted([str(c) for c in cols]))
        return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]

    @staticmethod
    def _to_text(v: Any) -> Optional[str]:
        if v is None:
            return None
        s = str(v).strip()
        return s if s else None

    @staticmethod
    def _to_float(v: Any) -> Optional[float]:
        if v in (None, ""):
            return None
        try:
            return float(str(v).replace(",", "").strip())
        except Exception:
            return None

    @staticmethod
    def _to_int(v: Any) -> Optional[int]:
        if v in (None, ""):
            return None
        try:
            return int(float(str(v).strip()))
        except Exception:
            return None

    @staticmethod
    def _parse_float(v: Any) -> Optional[float]:
        try:
            if v in (None, ""):
                return None
            return float(v)
        except Exception:
            return None

    @staticmethod
    def _parse_int(v: Any) -> Optional[int]:
        try:
            if v in (None, ""):
                return None
            return int(v)
        except Exception:
            return None

    @staticmethod
    def _normalize_phone(phone: str) -> Optional[str]:
        p = re.sub(r"[^0-9+]", "", (phone or "").strip())
        if re.fullmatch(r"\+91\d{10}", p):
            return p
        digits = re.sub(r"\D", "", p)
        if len(digits) == 12 and digits.startswith("91"):
            return "+" + digits
        if len(digits) == 10 and digits[0] in {"6", "7", "8", "9"}:
            return "+91" + digits
        return None

    def _write_error_csv(self, portfolio_id: str, rows: List[Dict[str, Any]]) -> str:
        folder = os.path.join(self.base_dir, portfolio_id)
        os.makedirs(folder, exist_ok=True)
        path = os.path.join(folder, "errors.csv")
        with open(path, "w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=["row_index", "customer_id", "phone", "dpd", "amount_due", "errors"])
            writer.writeheader()
            for r in rows:
                writer.writerow(r)
        return path

    @staticmethod
    def _read_exclusion_ids(path: str) -> List[str]:
        out: List[str] = []
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            for row in reader:
                if not row:
                    continue
                cid = (row[0] or "").strip()
                if cid and cid.lower() not in {"customer_id", "id"}:
                    out.append(cid)
        return out

    @staticmethod
    def _compile_predicate(expr: Optional[str]):
        if not expr:
            return None
        tokens = [t.strip() for t in str(expr).split(",") if t.strip()]
        parsed: List[Tuple[str, str, float]] = []
        for t in tokens:
            m = re.fullmatch(r"([a-zA-Z_][a-zA-Z0-9_]*)\s*(<=|>=|==|!=|<|>)\s*([-+]?\d+(?:\.\d+)?)", t)
            if not m:
                continue
            field = m.group(1)
            op = m.group(2)
            val = float(m.group(3))
            if field not in {"dpd", "amount_due"}:
                continue
            parsed.append((field, op, val))
        if not parsed:
            return None

        def _fn(rec: Dict[str, Any]) -> bool:
            for field, op, val in parsed:
                cur = rec.get(field)
                try:
                    curf = float(cur)
                except Exception:
                    continue
                if op == "<" and curf < val:
                    return True
                if op == ">" and curf > val:
                    return True
                if op == "<=" and curf <= val:
                    return True
                if op == ">=" and curf >= val:
                    return True
                if op == "==" and curf == val:
                    return True
                if op == "!=" and curf != val:
                    return True
            return False

        return _fn
