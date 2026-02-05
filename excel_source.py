from __future__ import annotations

import os
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook


REQUIRED_COLUMNS = [
    "customer_id",
    "customer_name",
    "phone",
    "language_preference",
    "overdue_amount",
    "due_date",
    "dpd",
    "risk_band",
]


class ExcelCustomerSource:
    def __init__(self, path: str, sheet: str = "Customers") -> None:
        self.path = path
        self.sheet = sheet

    def _ensure_exists(self) -> None:
        if not os.path.exists(self.path):
            raise FileNotFoundError(self.path)

    def list_customers(self) -> List[Dict[str, Any]]:
        self._ensure_exists()
        wb = load_workbook(self.path, read_only=True, data_only=True)
        try:
            ws = wb[self.sheet]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return []
            headers = [str(h or "").strip() for h in rows[0]]
            out: List[Dict[str, Any]] = []
            for r in rows[1:]:
                if not r:
                    continue
                rec = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
                cid = (rec.get("customer_id") or "").strip() if isinstance(rec.get("customer_id"), str) else rec.get("customer_id")
                if not cid:
                    continue
                out.append(self._normalize_record(rec))
            return out
        finally:
            wb.close()

    def get_customer(self, customer_id: str) -> Optional[Dict[str, Any]]:
        customer_id = str(customer_id or "").strip()
        if not customer_id:
            return None
        for rec in self.list_customers():
            if str(rec.get("customer_id")) == customer_id:
                return rec
        return None

    def _normalize_record(self, rec: Dict[str, Any]) -> Dict[str, Any]:
        # Keep keys stable; strip strings.
        out: Dict[str, Any] = {}
        for k in REQUIRED_COLUMNS:
            v = rec.get(k)
            if isinstance(v, str):
                v = v.strip()
            out[k] = v
        # Pass through any extra fields (but prefer REQUIRED_COLUMNS first).
        for k, v in rec.items():
            if k in out:
                continue
            out[k] = v.strip() if isinstance(v, str) else v
        return out

