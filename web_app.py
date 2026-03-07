import asyncio
import base64
import contextlib
import csv
import json
import logging
import os
import time
import hashlib
import hmac
import io
import secrets
import struct
import uuid
from typing import Dict, Optional, List, Any, Tuple

from fastapi import FastAPI, Request, WebSocket, WebSocketDisconnect, UploadFile, File
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, ValidationError
from starlette.websockets import WebSocketState

from config import get_env, get_env_bool, load_dotenv
from logging_utils import configure_logging, log_event, mask_secret
from sarvam_llm_service import SarvamLLMService
from sarvam_stt_service import SaarikaSTTService
from web_session import WebCallSession
from audit_store import SQLiteAuditStore
from excel_source import ExcelCustomerSource
from excel_sink import ExcelOutcomeSink
from actions import ActionRouter
from knowledge_store import SQLiteFTSKnowledgeStore
from strategy_engine import StrategyEngine
from compliance_engine import ComplianceEngine
from campaign_service import CampaignService
from followup_service import FollowupService
from integrations.crm_adapter import CRMAdapter
from integrations.zoho_crm import ZohoCRM
from portfolio_service import PortfolioService
from workbench_service import WorkbenchService
from alerts_service import AlertsService
from reports_service import ReportsService
from sync_service import SyncService
from approval_service import ApprovalService
from channel_hub_service import ChannelHubService
from customer_360_service import Customer360Service
from event_bus_service import EventBusService
from experiment_service import ExperimentService
from journey_orchestrator_service import JourneyOrchestratorService
from model_router_service import ModelRouterService
from payment_orchestration_service import PaymentOrchestrationService
from recovery_brain_service import RecoveryBrainService
from settlement_service import SettlementService
from tenant_service import TenantService
from auth_rbac import (
    ALERTS_MUTATE,
    ALERTS_VIEW,
    ALERT_RULES_EDIT,
    INTEGRATIONS_RESOLVE_CONFLICTS,
    INTEGRATIONS_VIEW,
    PORTFOLIO_MANAGE,
    REPORTS_SCHEDULE,
    REPORTS_VIEW,
    SYSTEM_VIEW_BUILD_INFO,
    USERS_MANAGE,
    VIEW_DASHBOARD,
    WORKBENCH_MUTATE,
    WORKBENCH_VIEW,
    ROLES,
    decode_token,
    decode_access_token,
    has_permissions,
    hash_password,
    issue_refresh_token,
    issue_ws_token,
    issue_access_token,
    password_policy_errors,
    role_permissions,
    verify_password,
)

logger = logging.getLogger(__name__)

BASE_DIR = os.path.dirname(__file__)
FRONTEND_DIR = os.path.join(BASE_DIR, "frontend")
FRONTEND_REACT_DIST_DIR = os.path.join(BASE_DIR, "frontend-react", "dist")
DATA_DIR = os.path.join(BASE_DIR, "data")
KNOWLEDGE_DIR_DEFAULT = os.path.join(BASE_DIR, "knowledge")
ACCESS_COOKIE_NAME = "ca_access_token"
REFRESH_COOKIE_NAME = "ca_refresh_token"
WS_CONNECTION_LIMIT_PER_USER = int(get_env("WS_CONNECTION_LIMIT_PER_USER", "3") or "3")
WEBHOOK_MAX_SKEW_SECONDS = int(get_env("WEBHOOK_MAX_SKEW_SECONDS", "300") or "300")
LOCKOUT_THRESHOLD = int(get_env("AUTH_LOCKOUT_THRESHOLD", "5") or "5")
LOCKOUT_DURATION_SECONDS = int(get_env("AUTH_LOCKOUT_SECONDS", "900") or "900")
PASSWORD_RESET_TOKEN_TTL_SECONDS = int(get_env("PASSWORD_RESET_TOKEN_TTL_SECONDS", "1800") or "1800")
IDEMPOTENCY_WINDOW_SECONDS = int(get_env("IDEMPOTENCY_WINDOW_SECONDS", "86400") or "86400")


def _app_env() -> str:
    raw = (get_env("APP_ENV", "") or "").strip().lower()
    if raw in {"prod", "production"}:
        return "prod"
    if raw in {"demo"}:
        return "demo"
    if raw in {"dev", "development", "local"}:
        return "dev"
    if get_env_bool("DEMO_MODE", False):
        return "demo"
    return "dev"


def _is_prod() -> bool:
    return _app_env() == "prod"


def _is_demo() -> bool:
    return _app_env() == "demo"


def _cookie_secure() -> bool:
    if _is_prod():
        return True
    return get_env_bool("COOKIE_SECURE", False)


def _access_token_ttl_seconds() -> int:
    raw = int(get_env("JWT_EXP_SECONDS", "28800") or "28800")
    if _is_prod():
        return max(60, min(900, raw))
    return max(60, raw)


def _refresh_token_ttl_seconds() -> int:
    raw = int(get_env("JWT_REFRESH_EXP_SECONDS", str(30 * 24 * 3600)) or str(30 * 24 * 3600))
    return max(300, raw)


def _jwt_secret() -> str:
    secret = str(get_env("JWT_SECRET", "dev-jwt-secret") or "dev-jwt-secret")
    if _is_prod() and len(secret.strip()) < 32:
        raise RuntimeError("weak_or_missing_jwt_secret")
    return secret


def _require_demo_mode() -> Optional[JSONResponse]:
    if _is_demo():
        return None
    return JSONResponse({"error": "demo_mode_required"}, status_code=403)


def _simulation_blocked_in_prod(request: Request) -> Optional[JSONResponse]:
    if not _is_prod():
        return None
    return _error_response(
        request,
        status_code=503,
        code="simulation_disabled_in_prod",
        message="Simulation-only path is disabled in production",
    )


def _error_response(
    request: Request,
    *,
    status_code: int,
    code: str,
    message: str,
    details: Optional[Dict[str, Any]] = None,
) -> JSONResponse:
    return JSONResponse(
        {
            "error": {
                "code": code,
                "message": message,
                "details": details or {},
            },
            "request_id": _request_id(request),
        },
        status_code=status_code,
    )


def _json_body_hash(body: Any) -> str:
    blob = json.dumps(body if isinstance(body, dict) else {}, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode("utf-8")
    return hashlib.sha256(blob).hexdigest()


def _idempotency_preflight(
    *,
    request: Request,
    audit: SQLiteAuditStore,
    route_key: str,
    body: Dict[str, Any],
) -> Tuple[str, str, Optional[JSONResponse]]:
    idem_key = (request.headers.get("Idempotency-Key") or "").strip()
    if not idem_key:
        return "", "", None
    body_hash = _json_body_hash(body)
    row = audit.get_idempotency(route_key=route_key, idempotency_key=idem_key)
    if not row:
        return idem_key, body_hash, None
    existing_hash = str(row.get("body_hash") or "")
    if existing_hash and existing_hash != body_hash:
        return idem_key, body_hash, _error_response(
            request,
            status_code=409,
            code="idempotency_conflict",
            message="Idempotency key was already used with a different request body",
            details={"idempotency_key": idem_key},
        )
    payload = row.get("response_json")
    if not isinstance(payload, dict):
        payload = {}
    return idem_key, body_hash, JSONResponse(payload, status_code=int(row.get("status_code") or 200))


def _idempotency_store(
    *,
    audit: SQLiteAuditStore,
    route_key: str,
    idempotency_key: str,
    body_hash: str,
    response_payload: Dict[str, Any],
    status_code: int = 200,
) -> None:
    if not idempotency_key:
        return
    audit.store_idempotency(
        route_key=route_key,
        idempotency_key=idempotency_key,
        body_hash=body_hash,
        response_json=response_payload,
        status_code=status_code,
        window_seconds=IDEMPOTENCY_WINDOW_SECONDS,
    )


def _verify_webhook_signature(request: Request, raw_body: bytes) -> Optional[JSONResponse]:
    secret = (get_env("V2_WEBHOOK_SECRET", "") or "").strip()
    if _is_prod() and not secret:
        return _error_response(
            request,
            status_code=500,
            code="webhook_secret_missing",
            message="Webhook secret is not configured",
        )
    if not secret:
        return None
    timestamp_raw = (request.headers.get("x-webhook-timestamp") or "").strip()
    signature_raw = (request.headers.get("x-webhook-signature") or "").strip()
    if not timestamp_raw or not signature_raw:
        return _error_response(
            request,
            status_code=401,
            code="missing_webhook_signature",
            message="Missing webhook signature headers",
        )
    try:
        timestamp = int(timestamp_raw)
    except Exception:
        return _error_response(
            request,
            status_code=401,
            code="invalid_webhook_timestamp",
            message="Invalid webhook timestamp",
        )
    if abs(int(time.time()) - timestamp) > WEBHOOK_MAX_SKEW_SECONDS:
        return _error_response(
            request,
            status_code=401,
            code="stale_webhook_timestamp",
            message="Webhook timestamp is outside allowed skew",
        )
    payload_to_sign = f"{timestamp_raw}.{raw_body.decode('utf-8', errors='replace')}".encode("utf-8")
    expected = hmac.new(secret.encode("utf-8"), payload_to_sign, hashlib.sha256).hexdigest()
    supplied = signature_raw.lower()
    if supplied.startswith("sha256="):
        supplied = supplied.split("=", 1)[1]
    if not hmac.compare_digest(supplied, expected):
        return _error_response(
            request,
            status_code=401,
            code="invalid_webhook_signature",
            message="Webhook signature verification failed",
        )
    return None


def _extract_webhook_event_id(body: Dict[str, Any]) -> str:
    for key in ("provider_event_id", "external_event_id", "event_id", "id"):
        val = body.get(key)
        if val is None:
            continue
        text = str(val).strip()
        if text:
            return text
    return ""

app = FastAPI()


@app.middleware("http")
async def https_redirect(request: Request, call_next):
    """Redirect HTTP to HTTPS on Azure App Service (or any reverse proxy that sets X-Forwarded-Proto)."""
    proto = request.headers.get("x-forwarded-proto", "")
    host = request.headers.get("host", "")
    if proto == "http" and host and "localhost" not in host and "127.0.0.1" not in host:
        url = request.url.replace(scheme="https")
        return RedirectResponse(url, status_code=301)
    return await call_next(request)


@app.middleware("http")
async def request_context(request: Request, call_next):
    req_id = (request.headers.get("x-request-id") or "").strip()
    if not req_id:
        req_id = f"req-{uuid.uuid4().hex[:12]}"
    request.state.request_id = req_id
    response = await call_next(request)
    response.headers["x-request-id"] = req_id
    return response


app.mount("/static", StaticFiles(directory=FRONTEND_DIR), name="static")

# Global singletons for the demo server process.
_demo_state: Dict[str, object] = {}
_report_scheduler_task: Optional[asyncio.Task] = None
_ws_user_connections: Dict[str, int] = {}


def _validate_runtime_config() -> None:
    # Strict production guardrails.
    _ = _jwt_secret()
    if _is_prod():
        if get_env_bool("DEMO_MODE", False):
            raise RuntimeError("DEMO_MODE must be disabled when APP_ENV=prod")
        if get_env("DEMO_PAY_BASE_URL", "https://pay.example/demo").find("pay.example") >= 0:
            raise RuntimeError("DEMO_PAY_BASE_URL uses placeholder domain in production")
        if not (get_env("V2_WEBHOOK_SECRET", "") or "").strip():
            raise RuntimeError("V2_WEBHOOK_SECRET is required in production")
        if get_env("ADMIN_TOKEN"):
            log_event(logger, "prod_config_warning", warning="ADMIN_TOKEN is ignored in prod")


@app.on_event("startup")
async def _startup_validate_config() -> None:
    configure_logging()
    _validate_runtime_config()
    log_event(logger, "app_startup", app_env=_app_env(), demo_mode=_is_demo(), pilot_mode=get_env_bool("PILOT_MODE", False))


def _file_token(path: str) -> str:
    try:
        st = os.stat(path)
        raw = f"{path}:{st.st_mtime_ns}:{st.st_size}".encode("utf-8")
        return hashlib.sha1(raw).hexdigest()[:10]
    except Exception:
        return "dev"


def _build_info_payload() -> Dict[str, Any]:
    admin_js = os.path.join(FRONTEND_DIR, "admin.js")
    app_js = os.path.join(FRONTEND_DIR, "app.js")
    styles_css = os.path.join(FRONTEND_DIR, "styles.css")
    react_main = os.path.join(BASE_DIR, "frontend-react", "src", "main.tsx")
    react_index = os.path.join(FRONTEND_REACT_DIST_DIR, "index.html")
    t1 = _file_token(admin_js)
    t2 = _file_token(app_js)
    t3 = _file_token(styles_css)
    t4 = _file_token(react_main)
    t5 = _file_token(react_index)
    static_token = hashlib.sha1(f"{t1}|{t2}|{t3}|{t4}|{t5}".encode("utf-8")).hexdigest()[:10]
    return {
        "app": "collections-agent",
        "static_token": static_token,
        "admin_js_token": t1,
        "app_js_token": t2,
        "styles_token": t3,
        "react_main_token": t4,
        "react_dist_token": t5,
        "demo_mode": _is_demo(),
        "app_env": _app_env(),
        "pilot_mode": get_env_bool("PILOT_MODE", False),
        "git_sha": get_env("GIT_SHA", "") or None,
        "timestamp": int(time.time()),
    }


def _render_html(path: str, *, no_store: bool = False) -> HTMLResponse:
    with open(path, "r", encoding="utf-8") as handle:
        html = handle.read()
    build = _build_info_payload()
    token = str(build.get("static_token") or "dev")
    html = html.replace("__STATIC_VER__", token)
    html = html.replace("__BUILD_TOKEN__", token)
    headers = {}
    if no_store:
        headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        headers["Pragma"] = "no-cache"
    return HTMLResponse(html, headers=headers)


def _render_react_html(*, no_store: bool = False) -> HTMLResponse:
    index_path = os.path.join(FRONTEND_REACT_DIST_DIR, "index.html")
    if not os.path.exists(index_path):
        html = (
            "<!doctype html><html><head><meta charset='utf-8' />"
            "<meta name='viewport' content='width=device-width, initial-scale=1' />"
            "<title>Collections Agent App</title></head><body>"
            "<div style='font-family:system-ui;padding:24px'>"
            "<h2>React app not built yet</h2>"
            "<p>Run <code>npm --prefix frontend-react install</code> and "
            "<code>npm --prefix frontend-react run build</code>, or open Vite dev server at "
            "<a href='http://127.0.0.1:5173/app'>http://127.0.0.1:5173/app</a>.</p>"
            "</div></body></html>"
        )
        headers = {}
        if no_store:
            headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
            headers["Pragma"] = "no-cache"
        return HTMLResponse(html, headers=headers)
    return _render_html(index_path, no_store=no_store)


def _role_tokens() -> Dict[str, str]:
    if not _is_demo():
        return {}
    out: Dict[str, str] = {}
    raw = get_env("ROLE_TOKENS_JSON")
    if raw:
        try:
            data = json.loads(raw)
            if isinstance(data, dict):
                for token, role in data.items():
                    if token and role:
                        out[str(token)] = str(role).upper()
        except Exception:
            pass
    admin_token = get_env("ADMIN_TOKEN")
    if admin_token and admin_token not in out:
        out[admin_token] = "ADMIN"
    return out


def _get_role(request: Request) -> str:
    ctx = _get_auth_context(request)
    return str(ctx.get("role") or "ANON")

def _extract_access_token(request: Request) -> str:
    auth = (request.headers.get("authorization") or "").strip()
    if auth.lower().startswith("bearer "):
        return auth[7:].strip()
    if _is_demo():
        hdr = (request.headers.get("x-admin-token") or "").strip()
        if hdr:
            return hdr
    cookie_token = (request.cookies.get(ACCESS_COOKIE_NAME) or "").strip()
    return cookie_token


def _to_public_user(row: Optional[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    if not row:
        return None
    return {
        "id": row.get("id"),
        "username": row.get("username"),
        "full_name": row.get("full_name"),
        "email": row.get("email"),
        "role": row.get("role"),
        "is_active": bool(int(row.get("is_active") or 0)),
        "default_tenant_id": row.get("default_tenant_id"),
        "must_change_password": bool(int(row.get("must_change_password") or 0)),
        "password_changed_at": row.get("password_changed_at"),
        "created_at": row.get("created_at"),
        "last_login_at": row.get("last_login_at"),
    }


def _get_auth_context(request: Request) -> Dict[str, Any]:
    token = _extract_access_token(request)
    if not token:
        return {}
    try:
        secret = _jwt_secret()
    except Exception:
        return {}
    payload = decode_access_token(token, secret)
    if payload and payload.get("sub"):
        demo = _get_demo_singletons()
        audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
        user = audit.get_user_by_id(str(payload.get("sub")))
        if user and int(user.get("is_active") or 0) == 1:
            sid = str(payload.get("sid") or "").strip() or None
            if sid:
                sess = audit.get_user_session(session_id=sid)
                if not sess or sess.get("revoked_at") or float(sess.get("expires_at") or 0) < time.time():
                    return {}
            role = str(user.get("role") or "").upper()
            tenants = audit.list_user_tenants(user_id=str(user.get("id") or ""))
            default_tenant = str(user.get("default_tenant_id") or (tenants[0] if tenants else "default"))
            return {
                "token": token,
                "role": role,
                "permissions": sorted(role_permissions(role)),
                "user": _to_public_user(user),
                "tenant_ids": tenants or ["default"],
                "default_tenant_id": default_tenant or "default",
                "session_id": sid,
                "jti": str(payload.get("jti") or "").strip() or None,
                "source": "jwt",
            }
    legacy_role = _role_tokens().get(token)
    if legacy_role:
        role = str(legacy_role).upper()
        return {
            "token": token,
            "role": role,
            "permissions": sorted(role_permissions(role)),
            "user": {"id": f"legacy-{role.lower()}", "username": f"{role.lower()}_legacy", "role": role, "is_active": True},
            "tenant_ids": ["default"],
            "default_tenant_id": "default",
            "source": "legacy",
        }
    return {}


def _require_permissions(request: Request, required: Tuple[str, ...]) -> Optional[JSONResponse]:
    ctx = _get_auth_context(request)
    if not ctx:
        return _error_response(
            request,
            status_code=401,
            code="unauthorized",
            message="Authentication required",
            details={"required_permissions": list(required)},
        )
    role = str(ctx.get("role") or "")
    if has_permissions(role, required):
        tenant_deny = _tenant_access_denied(request, ctx)
        if tenant_deny:
            return tenant_deny
        return None
    return _error_response(
        request,
        status_code=403,
        code="forbidden",
        message="Insufficient permissions",
        details={
            "required_permissions": list(required),
            "role": role,
            "permissions": ctx.get("permissions") or [],
        },
    )


def _require_role(request: Request, allowed: Tuple[str, ...]) -> Optional[JSONResponse]:
    ctx = _get_auth_context(request)
    if not ctx:
        return _error_response(
            request,
            status_code=401,
            code="unauthorized",
            message="Authentication required",
            details={"required_roles": list(allowed)},
        )
    role = str(ctx.get("role") or "")
    if role in allowed:
        tenant_deny = _tenant_access_denied(request, ctx)
        if tenant_deny:
            return tenant_deny
        return None
    return _error_response(
        request,
        status_code=403,
        code="forbidden",
        message="Role not allowed",
        details={"required_roles": list(allowed), "current_role": role},
    )


def _require_pilot_mode() -> Optional[JSONResponse]:
    if get_env_bool("PILOT_MODE", False):
        return None
    return JSONResponse({"error": "pilot_mode_disabled"}, status_code=403)


def _actor(request: Request) -> str:
    actor = (request.headers.get("x-user") or "").strip()
    if actor:
        return actor
    ctx = _get_auth_context(request)
    user = ctx.get("user") if isinstance(ctx.get("user"), dict) else {}
    if isinstance(user, dict) and user.get("username"):
        return str(user["username"])
    role = str(ctx.get("role") or "system").lower()
    return f"{role}_user"


def _request_id(request: Request) -> str:
    req_id = getattr(getattr(request, "state", object()), "request_id", None)
    if isinstance(req_id, str) and req_id.strip():
        return req_id.strip()
    hdr = (request.headers.get("x-request-id") or "").strip()
    if hdr:
        return hdr
    return f"req-{uuid.uuid4().hex[:12]}"


def _requested_tenant(request: Request) -> str:
    return (request.headers.get("x-tenant-id") or request.query_params.get("tenant_id") or "").strip().lower()


def _tenant_access_denied(request: Request, ctx: Dict[str, Any]) -> Optional[JSONResponse]:
    requested = _requested_tenant(request)
    if not requested:
        return None
    allowed = {str(t).strip().lower() for t in (ctx.get("tenant_ids") or []) if str(t).strip()}
    if not allowed:
        return None
    if requested in allowed:
        return None
    return _error_response(
        request,
        status_code=403,
        code="forbidden_tenant",
        message="Tenant access denied",
        details={"tenant_id": requested},
    )


def _tenant_id(request: Request) -> str:
    ctx = _get_auth_context(request)
    requested = _requested_tenant(request)
    default_tid = str(ctx.get("default_tenant_id") or "default").strip().lower() if ctx else "default"
    tid = requested or default_tid or "default"
    demo = _get_demo_singletons()
    tenants: TenantService = demo["tenant_service"]  # type: ignore[assignment]
    return tenants.ensure_tenant(tid)


def _v2_context(request: Request) -> Dict[str, str]:
    return {
        "tenant_id": _tenant_id(request),
        "actor": _actor(request),
        "request_id": _request_id(request),
    }


def _pagination(request: Request) -> Tuple[int, int]:
    try:
        page = int(request.query_params.get("page") or 1)
    except Exception:
        page = 1
    try:
        page_size = int(request.query_params.get("page_size") or 25)
    except Exception:
        page_size = 25
    return max(1, page), max(1, min(100, page_size))


def _seed_default_users(audit: SQLiteAuditStore) -> None:
    if not _is_demo():
        return
    if audit.has_users():
        return
    defaults = [
        ("admin", "admin123", "ADMIN", "Platform Admin"),
        ("ceo", "ceo123", "CEO", "Chief Executive Officer"),
        ("cfo", "cfo123", "CFO", "Chief Financial Officer"),
        ("mgr", "mgr123", "COLLECTIONS_MANAGER", "Collections Manager"),
        ("agent", "agent123", "CALLING_AGENT", "Calling Agent"),
        ("comp", "comp123", "COMPLIANCE_OFFICER", "Compliance Officer"),
        ("viewer", "viewer123", "VIEWER", "Read Only Viewer"),
    ]
    for username, password, role, full_name in defaults:
        audit.create_user(
            user_id=f"usr-{uuid.uuid4().hex[:12]}",
            username=username,
            password_hash=hash_password(password),
            full_name=full_name,
            email=f"{username}@demo.local",
            role=role,
            is_active=True,
            default_tenant_id="default",
            actor="system",
        )


@app.get("/")
async def index() -> HTMLResponse:
    return RedirectResponse("/app", status_code=302)


@app.get("/console")
async def console() -> HTMLResponse:
    index_path = os.path.join(FRONTEND_DIR, "index.html")
    return _render_html(index_path, no_store=True)


@app.get("/login")
async def login_page() -> HTMLResponse:
    return _render_react_html(no_store=True)


@app.get("/app")
async def app_index() -> HTMLResponse:
    return _render_react_html(no_store=True)


@app.get("/app/{full_path:path}")
async def app_spa(full_path: str):
    if os.path.exists(FRONTEND_REACT_DIST_DIR):
        candidate = os.path.abspath(os.path.join(FRONTEND_REACT_DIST_DIR, full_path))
        dist_abs = os.path.abspath(FRONTEND_REACT_DIST_DIR)
        if candidate.startswith(dist_abs) and os.path.isfile(candidate):
            return FileResponse(candidate)
    return _render_react_html(no_store=True)

@app.get("/admin")
async def admin(request: Request) -> HTMLResponse:
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return RedirectResponse("/login?next=/admin", status_code=302)
    admin_path = os.path.join(FRONTEND_DIR, "admin.html")
    return _render_html(admin_path, no_store=True)


@app.get("/api/system/build_info")
async def api_build_info(request: Request):
    deny = _require_permissions(request, (SYSTEM_VIEW_BUILD_INFO,))
    if deny:
        return deny
    return _build_info_payload()


@app.get("/api/public/build_info")
async def api_public_build_info():
    return _build_info_payload()


def _auth_me_payload(ctx: Dict[str, Any]) -> Dict[str, Any]:
    user = ctx.get("user") if isinstance(ctx.get("user"), dict) else {}
    return {
        "user": user,
        "role": ctx.get("role"),
        "permissions": ctx.get("permissions") or [],
        "tenant_ids": ctx.get("tenant_ids") or [],
        "default_tenant_id": ctx.get("default_tenant_id") or "default",
        "session_id": ctx.get("session_id"),
        "must_change_password": bool((user or {}).get("must_change_password")),
    }


class LoginRequest(BaseModel):
    username: str = Field(min_length=1)
    password: str = Field(min_length=1)
    tenant_id: Optional[str] = None


class RefreshRequest(BaseModel):
    refresh_token: Optional[str] = None


class ChangePasswordRequest(BaseModel):
    current_password: str = Field(min_length=1)
    new_password: str = Field(min_length=1)


class WSTokenRequest(BaseModel):
    tenant_id: Optional[str] = None


class ForgotPasswordRequest(BaseModel):
    username_or_email: str = Field(min_length=1)


class ResetPasswordRequest(BaseModel):
    token: str = Field(min_length=1)
    new_password: str = Field(min_length=1)


class TenantSettingsPatchRequest(BaseModel):
    settings: Dict[str, Any]


class TelephonyTestCallRequest(BaseModel):
    phone: str = Field(min_length=8)
    customer_name: Optional[str] = None
    amount_due: Optional[str] = None
    timeout_seconds: Optional[int] = Field(default=25, ge=5, le=60)


class TelephonyAgentCallRequest(BaseModel):
    phone: str = Field(min_length=8)
    customer_name: Optional[str] = None
    customer_id: Optional[str] = None
    campaign_id: Optional[str] = None
    amount_due: Optional[str] = None
    language: Optional[str] = None
    tts_speaker: Optional[str] = None
    timeout_seconds: Optional[int] = Field(default=25, ge=5, le=60)


class V2ChannelSendRequest(BaseModel):
    customer_id: Optional[str] = None
    journey_id: Optional[str] = None
    loan_account_id: Optional[str] = None
    channel: str = Field(min_length=1)
    content: str = Field(min_length=1)
    provider: Optional[str] = None
    provider_message_id: Optional[str] = None
    message_type: Optional[str] = None
    metadata: Optional[Dict[str, Any]] = None


class V2PaymentIntentRequest(BaseModel):
    customer_id: Optional[str] = None
    loan_account_id: Optional[str] = None
    journey_id: Optional[str] = None
    amount: float
    currency: Optional[str] = "INR"
    rail: str = Field(min_length=1)
    provider_ref: Optional[str] = None
    payment_link: Optional[str] = None
    metadata: Optional[Dict[str, Any]] = None


class V2SettlementOfferRequest(BaseModel):
    customer_id: str = Field(min_length=1)
    loan_account_id: str = Field(min_length=1)
    journey_id: Optional[str] = None
    offered_amount: float
    original_due_amount: float
    terms: Optional[Dict[str, Any]] = None


class V2SettlementAcceptRequest(BaseModel):
    accepted_amount: Optional[float] = None
    acceptance_channel: Optional[str] = None
    note: Optional[str] = None


def _resolve_user_tenant(
    *,
    audit: SQLiteAuditStore,
    user_id: str,
    requested_tenant: Optional[str],
) -> Tuple[str, List[str]]:
    tenant_ids = audit.list_user_tenants(user_id=user_id)
    if not tenant_ids:
        tenant_ids = ["default"]
        audit.add_user_tenant_membership(user_id=user_id, tenant_id="default", is_default=True, actor="system")
    requested = (requested_tenant or "").strip().lower()
    if requested:
        if requested not in tenant_ids:
            raise ValueError("forbidden_tenant")
        return requested, tenant_ids
    default_tid = (audit.get_user_default_tenant(user_id=user_id) or "default").strip().lower()
    if default_tid not in tenant_ids:
        default_tid = tenant_ids[0]
    return default_tid, tenant_ids


def _set_auth_cookies(resp: JSONResponse, *, access_token: str, refresh_token: str) -> None:
    access_ttl = _access_token_ttl_seconds()
    refresh_ttl = _refresh_token_ttl_seconds()
    resp.set_cookie(
        ACCESS_COOKIE_NAME,
        access_token,
        max_age=access_ttl,
        httponly=True,
        samesite="lax",
        secure=_cookie_secure(),
        path="/",
    )
    resp.set_cookie(
        REFRESH_COOKIE_NAME,
        refresh_token,
        max_age=refresh_ttl,
        httponly=True,
        samesite="lax",
        secure=_cookie_secure(),
        path="/",
    )


def _clear_auth_cookies(resp: JSONResponse) -> None:
    resp.delete_cookie(ACCESS_COOKIE_NAME, path="/")
    resp.delete_cookie(REFRESH_COOKIE_NAME, path="/")


def _issue_session_tokens(
    *,
    audit: SQLiteAuditStore,
    user_row: Dict[str, Any],
    tenant_id: str,
    request: Request,
    session_id: Optional[str] = None,
) -> Dict[str, Any]:
    user_id = str(user_row["id"])
    username = str(user_row["username"])
    role = str(user_row.get("role") or "").upper()
    sid = session_id or f"ses-{uuid.uuid4().hex[:16]}"
    access_jti = f"aj-{uuid.uuid4().hex[:16]}"
    refresh_jti = f"rj-{uuid.uuid4().hex[:16]}"
    secret = _jwt_secret()
    access_token = issue_access_token(
        user_id=user_id,
        username=username,
        role=role,
        tenant_id=tenant_id,
        session_id=sid,
        jti=access_jti,
        secret=secret,
        expires_in_seconds=_access_token_ttl_seconds(),
    )
    refresh_token = issue_refresh_token(
        user_id=user_id,
        username=username,
        role=role,
        tenant_id=tenant_id,
        session_id=sid,
        jti=refresh_jti,
        secret=secret,
        expires_in_seconds=_refresh_token_ttl_seconds(),
    )
    refresh_hash = audit.hash_token(refresh_token)
    expires_at = time.time() + _refresh_token_ttl_seconds()
    if session_id:
        audit.rotate_user_session(
            session_id=sid,
            refresh_token_hash=refresh_hash,
            access_jti=access_jti,
            expires_at=expires_at,
        )
    else:
        audit.create_user_session(
            session_id=sid,
            user_id=user_id,
            tenant_id=tenant_id,
            refresh_token_hash=refresh_hash,
            expires_at=expires_at,
            access_jti=access_jti,
            user_agent=(request.headers.get("user-agent") or "")[:512] or None,
            ip_addr=(request.client.host if request.client else None),
        )
    return {
        "access_token": access_token,
        "refresh_token": refresh_token,
        "session_id": sid,
    }


def _generate_temporary_password() -> str:
    # Enforce policy shape: upper + lower + digit + special, minimum 12 chars.
    return f"Tmp{secrets.token_urlsafe(10)}!9a"


@app.post("/api/auth/login")
async def api_auth_login(request: Request):
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    try:
        body = LoginRequest.model_validate(await request.json())
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid login payload",
            details={"errors": exc.errors()},
        )
    except Exception:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid JSON body",
        )
    username = str(body.username or "").strip().lower()
    password = str(body.password or "")
    row = audit.get_user_by_username(username)
    if not row or int(row.get("is_active") or 0) != 1:
        return _error_response(request, status_code=401, code="invalid_credentials", message="Invalid credentials")
    if audit.is_user_locked(user_id=str(row["id"])):
        return _error_response(request, status_code=423, code="account_locked", message="Account temporarily locked")
    if not verify_password(password, str(row.get("password_hash") or "")):
        audit.mark_login_failure(
            user_id=str(row["id"]),
            threshold=LOCKOUT_THRESHOLD,
            lock_seconds=LOCKOUT_DURATION_SECONDS,
        )
        return _error_response(request, status_code=401, code="invalid_credentials", message="Invalid credentials")
    role = str(row.get("role") or "").upper()
    try:
        tenant_id, tenant_ids = _resolve_user_tenant(
            audit=audit,
            user_id=str(row["id"]),
            requested_tenant=body.tenant_id,
        )
    except ValueError:
        return _error_response(request, status_code=403, code="forbidden_tenant", message="Tenant access denied")
    issued = _issue_session_tokens(audit=audit, user_row=row, tenant_id=tenant_id, request=request)
    audit.mark_user_login(user_id=str(row["id"]))
    payload = {
        "access_token": issued["access_token"],
        "token_type": "bearer",
        "user": _to_public_user(row),
        "permissions": sorted(role_permissions(role)),
        "tenant_ids": tenant_ids,
        "default_tenant_id": tenant_id,
        "session_id": issued["session_id"],
        "must_change_password": bool(int(row.get("must_change_password") or 0)),
    }
    if _is_demo():
        payload["refresh_token"] = issued["refresh_token"]
    resp = JSONResponse(payload)
    _set_auth_cookies(resp, access_token=issued["access_token"], refresh_token=issued["refresh_token"])
    return resp


@app.post("/api/auth/refresh")
async def api_auth_refresh(request: Request):
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    try:
        body_raw = await request.json()
    except Exception:
        body_raw = {}
    try:
        body = RefreshRequest.model_validate(body_raw if isinstance(body_raw, dict) else {})
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid refresh payload",
            details={"errors": exc.errors()},
        )
    refresh_token = (
        (body.refresh_token or "").strip()
        or (request.cookies.get(REFRESH_COOKIE_NAME) or "").strip()
    )
    if not refresh_token:
        return _error_response(request, status_code=401, code="missing_refresh_token", message="Refresh token required")
    payload = decode_token(refresh_token, _jwt_secret(), expected_token_type="refresh", expected_aud="refresh")
    if not payload:
        return _error_response(request, status_code=401, code="invalid_refresh_token", message="Invalid refresh token")
    sid = str(payload.get("sid") or "").strip()
    sub = str(payload.get("sub") or "").strip()
    if not sid or not sub:
        return _error_response(request, status_code=401, code="invalid_refresh_token", message="Invalid refresh token")
    row = audit.get_user_session(session_id=sid)
    if not row or str(row.get("user_id") or "") != sub:
        return _error_response(request, status_code=401, code="invalid_refresh_token", message="Session not found")
    if row.get("revoked_at") or float(row.get("expires_at") or 0) < time.time():
        return _error_response(request, status_code=401, code="session_expired", message="Session expired")
    if str(row.get("refresh_token_hash") or "") != audit.hash_token(refresh_token):
        return _error_response(request, status_code=401, code="invalid_refresh_token", message="Refresh token mismatch")
    user = audit.get_user_by_id(sub)
    if not user or int(user.get("is_active") or 0) != 1:
        return _error_response(request, status_code=401, code="invalid_session_user", message="Session user inactive")
    tenant_id = str(payload.get("tenant_id") or row.get("tenant_id") or audit.get_user_default_tenant(user_id=sub) or "default")
    if not audit.user_has_tenant(user_id=sub, tenant_id=tenant_id):
        return _error_response(request, status_code=403, code="forbidden_tenant", message="Tenant access denied")
    issued = _issue_session_tokens(audit=audit, user_row=user, tenant_id=tenant_id, request=request, session_id=sid)
    resp = JSONResponse(
        {
            "access_token": issued["access_token"],
            "token_type": "bearer",
            "session_id": sid,
            "tenant_id": tenant_id,
        }
    )
    _set_auth_cookies(resp, access_token=issued["access_token"], refresh_token=issued["refresh_token"])
    return resp


@app.post("/api/auth/logout")
async def api_auth_logout(request: Request):
    ctx = _get_auth_context(request)
    if ctx.get("session_id"):
        demo = _get_demo_singletons()
        audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
        with contextlib.suppress(Exception):
            audit.revoke_user_session(session_id=str(ctx["session_id"]))
    resp = JSONResponse({"ok": True})
    _clear_auth_cookies(resp)
    return resp


@app.get("/api/auth/me")
async def api_auth_me(request: Request):
    ctx = _get_auth_context(request)
    if not ctx:
        return _error_response(request, status_code=401, code="unauthorized", message="Authentication required")
    return _auth_me_payload(ctx)


@app.post("/api/auth/ws-token")
async def api_auth_ws_token(request: Request):
    deny = _require_permissions(request, (WORKBENCH_VIEW,))
    if deny:
        return deny
    ctx = _get_auth_context(request)
    user = ctx.get("user") if isinstance(ctx.get("user"), dict) else {}
    if not user:
        return _error_response(request, status_code=401, code="unauthorized", message="Authentication required")
    try:
        body = WSTokenRequest.model_validate(await request.json())
    except Exception:
        body = WSTokenRequest()
    tenant_id = (body.tenant_id or _tenant_id(request)).strip().lower()
    if tenant_id not in {str(t) for t in (ctx.get("tenant_ids") or ["default"])}:
        return _error_response(request, status_code=403, code="forbidden_tenant", message="Tenant access denied")
    token = issue_ws_token(
        user_id=str(user.get("id") or ""),
        username=str(user.get("username") or ""),
        role=str(ctx.get("role") or ""),
        tenant_id=tenant_id,
        secret=_jwt_secret(),
        expires_in_seconds=60,
        session_id=str(ctx.get("session_id") or "") or None,
        jti=f"ws-{uuid.uuid4().hex[:16]}",
    )
    return {"ws_token": token, "expires_in_seconds": 60}


@app.get("/api/auth/sessions")
async def api_auth_sessions(request: Request):
    ctx = _get_auth_context(request)
    if not ctx:
        return _error_response(request, status_code=401, code="unauthorized", message="Authentication required")
    user = ctx.get("user") if isinstance(ctx.get("user"), dict) else {}
    actor_user_id = str((user or {}).get("id") or "")
    if not actor_user_id:
        return _error_response(request, status_code=401, code="unauthorized", message="Authentication required")
    target_user_id = actor_user_id
    requested_user_id = (request.query_params.get("user_id") or "").strip()
    if requested_user_id and has_permissions(str(ctx.get("role") or ""), (USERS_MANAGE,)):
        target_user_id = requested_user_id
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    return {"rows": audit.list_user_sessions(user_id=target_user_id)}


@app.post("/api/auth/sessions/{session_id}/revoke")
async def api_auth_session_revoke(session_id: str, request: Request):
    ctx = _get_auth_context(request)
    if not ctx:
        return _error_response(request, status_code=401, code="unauthorized", message="Authentication required")
    user = ctx.get("user") if isinstance(ctx.get("user"), dict) else {}
    actor_user_id = str((user or {}).get("id") or "")
    if not actor_user_id:
        return _error_response(request, status_code=401, code="unauthorized", message="Authentication required")
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    sess = audit.get_user_session(session_id=session_id)
    if not sess:
        return _error_response(request, status_code=404, code="session_not_found", message="Session not found")
    session_user_id = str(sess.get("user_id") or "")
    is_admin = has_permissions(str(ctx.get("role") or ""), (USERS_MANAGE,))
    if session_user_id != actor_user_id and not is_admin:
        return _error_response(request, status_code=403, code="forbidden", message="Cannot revoke this session")
    return {"ok": audit.revoke_user_session(session_id=session_id)}


@app.post("/api/auth/forgot-password")
async def api_auth_forgot_password(request: Request):
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    try:
        body = ForgotPasswordRequest.model_validate(await request.json())
    except Exception:
        body = ForgotPasswordRequest(username_or_email="")
    try:
        body = ForgotPasswordRequest.model_validate({"username_or_email": body.username_or_email})
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid forgot-password payload",
            details={"errors": exc.errors()},
        )
    needle = (body.username_or_email or "").strip().lower()
    user = audit.get_user_by_username(needle)
    if not user:
        for row in audit.list_users():
            if str(row.get("email") or "").strip().lower() == needle:
                user = audit.get_user_by_id(str(row.get("id")))
                break
    token_preview = None
    if user:
        token_preview = audit.create_password_reset_token(user_id=str(user["id"]), ttl_seconds=PASSWORD_RESET_TOKEN_TTL_SECONDS)
    result: Dict[str, Any] = {"ok": True, "message": "If the account exists, reset instructions were issued."}
    if _is_demo() and token_preview:
        result["reset_token_preview"] = token_preview
    return result


@app.post("/api/auth/reset-password")
async def api_auth_reset_password(request: Request):
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    try:
        body = ResetPasswordRequest.model_validate(await request.json())
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid reset-password payload",
            details={"errors": exc.errors()},
        )
    errs = password_policy_errors(body.new_password)
    if errs:
        return _error_response(
            request,
            status_code=400,
            code="weak_password",
            message="Password does not meet policy",
            details={"violations": errs},
        )
    token_row = audit.consume_password_reset_token(token=body.token)
    if not token_row:
        return _error_response(request, status_code=400, code="invalid_reset_token", message="Invalid or expired reset token")
    user_id = str(token_row["user_id"])
    audit.set_user_password(user_id=user_id, password_hash=hash_password(body.new_password))
    with contextlib.suppress(Exception):
        audit.revoke_user_sessions(user_id=user_id)
    return {"ok": True}


@app.post("/api/auth/change_password")
async def api_auth_change_password(request: Request):
    ctx = _get_auth_context(request)
    if not ctx:
        return _error_response(request, status_code=401, code="unauthorized", message="Authentication required")
    user = ctx.get("user") if isinstance(ctx.get("user"), dict) else {}
    user_id = str((user or {}).get("id") or "")
    if not user_id:
        return _error_response(request, status_code=401, code="unauthorized", message="Authentication required")
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    try:
        body = ChangePasswordRequest.model_validate(await request.json())
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid change-password payload",
            details={"errors": exc.errors()},
        )
    db_user = audit.get_user_by_id(user_id)
    if not db_user:
        return _error_response(request, status_code=404, code="user_not_found", message="User not found")
    if not verify_password(body.current_password, str(db_user.get("password_hash") or "")):
        return _error_response(request, status_code=400, code="invalid_current_password", message="Invalid current password")
    errs = password_policy_errors(body.new_password)
    if errs:
        return _error_response(
            request,
            status_code=400,
            code="weak_password",
            message="Password does not meet policy",
            details={"violations": errs},
        )
    audit.set_user_password(user_id=user_id, password_hash=hash_password(body.new_password))
    with contextlib.suppress(Exception):
        audit.revoke_user_sessions(user_id=user_id)
    return {"ok": True}


@app.get("/api/settings/tenant")
async def api_tenant_settings_get(request: Request):
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    tenants: TenantService = demo["tenant_service"]  # type: ignore[assignment]
    tenant_id = _tenant_id(request)
    row = tenants.get_tenant(tenant_id)
    if not row:
        return _error_response(request, status_code=404, code="tenant_not_found", message="Tenant not found")
    return {"tenant": row}


@app.patch("/api/settings/tenant")
async def api_tenant_settings_patch(request: Request):
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return deny
    try:
        body = TenantSettingsPatchRequest.model_validate(await request.json())
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid settings payload",
            details={"errors": exc.errors()},
        )
    demo = _get_demo_singletons()
    tenants: TenantService = demo["tenant_service"]  # type: ignore[assignment]
    tenant_id = _tenant_id(request)
    try:
        row = tenants.patch_tenant_settings(
            tenant_id=tenant_id,
            patch=body.settings,
            actor=_actor(request),
            request_id=_request_id(request),
        )
    except ValueError as exc:
        return _error_response(request, status_code=400, code=str(exc), message="Could not update settings")
    return {"ok": True, "tenant": row}


@app.get("/api/users")
async def api_users_list(request: Request):
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    return {"rows": audit.list_users()}


@app.post("/api/users")
async def api_users_create(request: Request):
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    username = str(body.get("username") or "").strip().lower()
    password = str(body.get("password") or "")
    role = str(body.get("role") or "").strip().upper()
    if not username or not password or role not in ROLES:
        return _error_response(request, status_code=400, code="invalid_payload", message="Invalid user payload")
    errs = password_policy_errors(password)
    if errs:
        return _error_response(
            request,
            status_code=400,
            code="weak_password",
            message="Password does not meet policy",
            details={"violations": errs},
        )
    if audit.get_user_by_username(username):
        return _error_response(request, status_code=409, code="username_exists", message="Username already exists")
    uid = f"usr-{uuid.uuid4().hex[:12]}"
    default_tenant_id = str(body.get("default_tenant_id") or _tenant_id(request) or "default").strip().lower() or "default"
    audit.create_user(
        user_id=uid,
        username=username,
        password_hash=hash_password(password),
        full_name=str(body.get("full_name") or "").strip() or None,
        email=str(body.get("email") or "").strip() or None,
        role=role,
        is_active=bool(body.get("is_active", True)),
        default_tenant_id=default_tenant_id,
        must_change_password=bool(body.get("must_change_password", False)),
        actor=_actor(request),
    )
    row = audit.get_user_by_id(uid)
    return {"ok": True, "user": _to_public_user(row)}


@app.patch("/api/users/{user_id}")
async def api_users_patch(user_id: str, request: Request):
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    role = body.get("role")
    if role is not None:
        role = str(role).strip().upper()
        if role not in ROLES:
            return _error_response(request, status_code=400, code="invalid_role", message="Invalid role")
    default_tenant_id = body.get("default_tenant_id")
    if default_tenant_id is not None:
        default_tenant_id = str(default_tenant_id).strip().lower() or "default"
    ok = audit.update_user(
        user_id=user_id,
        role=role,
        is_active=body.get("is_active") if "is_active" in body else None,
        full_name=body.get("full_name") if "full_name" in body else None,
        email=body.get("email") if "email" in body else None,
        default_tenant_id=default_tenant_id if "default_tenant_id" in body else None,
        must_change_password=bool(body.get("must_change_password")) if "must_change_password" in body else None,
    )
    if not ok:
        return _error_response(request, status_code=404, code="user_not_found_or_no_change", message="User not found")
    if default_tenant_id:
        audit.add_user_tenant_membership(
            user_id=user_id,
            tenant_id=default_tenant_id,
            is_default=True,
            actor=_actor(request),
        )
    return {"ok": True, "user": _to_public_user(audit.get_user_by_id(user_id))}


@app.post("/api/users/{user_id}/reset_password")
async def api_users_reset_password(user_id: str, request: Request):
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    provided_password = str(body.get("new_password") or "").strip()
    new_password = provided_password or _generate_temporary_password()
    errs = password_policy_errors(new_password)
    if errs:
        return _error_response(
            request,
            status_code=400,
            code="weak_password",
            message="Password does not meet policy",
            details={"violations": errs},
        )
    ok = audit.set_user_password(user_id=user_id, password_hash=hash_password(new_password))
    if not ok:
        return _error_response(request, status_code=404, code="user_not_found", message="User not found")
    audit.update_user(user_id=user_id, must_change_password=True)
    with contextlib.suppress(Exception):
        audit.revoke_user_sessions(user_id=user_id)
    payload: Dict[str, Any] = {"ok": True, "must_change_password": True}
    if _is_demo():
        payload["temporary_password"] = new_password
    return payload


def _require_admin(request) -> Optional[JSONResponse]:
    return _require_permissions(request, (USERS_MANAGE,))


def _get_demo_singletons() -> Dict[str, object]:
    if _demo_state:
        return _demo_state

    load_dotenv(override=False)
    os.makedirs(DATA_DIR, exist_ok=True)
    customers_path = get_env("CUSTOMERS_XLSX_PATH", os.path.join(DATA_DIR, "customers.xlsx"))
    output_path = get_env("DEMO_OUTPUT_XLSX_PATH", os.path.join(DATA_DIR, "demo_output.xlsx"))
    db_path = get_env("DEMO_DB_PATH", os.path.join(DATA_DIR, "demo.db"))
    knowledge_dir = get_env("KNOWLEDGE_DIR", KNOWLEDGE_DIR_DEFAULT)
    pay_base_url = get_env("DEMO_PAY_BASE_URL", "https://pay.example/demo")
    twilio_account_sid = get_env("TWILIO_ACCOUNT_SID")
    twilio_auth_token = get_env("TWILIO_AUTH_TOKEN")
    twilio_from_number = get_env("TWILIO_PHONE_NUMBER", get_env("TWILIO_FROM_NUMBER"))
    twilio_whatsapp_from = get_env("TWILIO_WHATSAPP_FROM")
    twilio_send_enabled_raw = get_env("TWILIO_SEND_ENABLED")
    if twilio_send_enabled_raw is None or str(twilio_send_enabled_raw).strip() == "":
        twilio_send_enabled: Optional[bool] = None
    else:
        twilio_send_enabled = str(twilio_send_enabled_raw).strip().lower() in {"1", "true", "yes", "y", "on"}
    twilio_timeout_s = float(get_env("TWILIO_TIMEOUT_S", "8") or "8")

    retention_days = int(get_env("PHI_RETENTION_DAYS", "30") or "30")
    metrics_tz = get_env("METRICS_TIMEZONE", "Asia/Kolkata") or "Asia/Kolkata"

    audit = SQLiteAuditStore(db_path, retention_days=retention_days, metrics_tz=metrics_tz)
    _seed_default_users(audit)
    sink = ExcelOutcomeSink(output_path)
    sink.ensure_workbook()
    source = ExcelCustomerSource(customers_path)
    router = ActionRouter(
        pay_base_url=pay_base_url,
        twilio_account_sid=twilio_account_sid,
        twilio_auth_token=twilio_auth_token,
        twilio_from_number=twilio_from_number,
        twilio_whatsapp_from=twilio_whatsapp_from,
        twilio_send_enabled=twilio_send_enabled,
        twilio_timeout_s=twilio_timeout_s,
    )
    knowledge = SQLiteFTSKnowledgeStore(db_path, knowledge_dir)
    strategy = StrategyEngine()
    compliance = ComplianceEngine()
    campaign_service = CampaignService(db_path)
    followups = FollowupService(db_path, tz_name=metrics_tz)
    crm = CRMAdapter(db_path)
    zoho_crm = ZohoCRM(os.path.join(DATA_DIR, "crm_sync.db"))
    portfolio = PortfolioService(db_path, DATA_DIR)
    workbench = WorkbenchService(db_path)
    alerts = AlertsService(db_path)
    reports = ReportsService(db_path, DATA_DIR, tz_name=metrics_tz)
    sync = SyncService(db_path)
    tenants = TenantService(db_path)
    event_bus = EventBusService(db_path)
    customer_360 = Customer360Service(db_path)
    journeys = JourneyOrchestratorService(db_path)
    channel_hub = ChannelHubService(db_path)
    payments = PaymentOrchestrationService(db_path, pay_base_url=pay_base_url)
    settlements = SettlementService(db_path)
    approvals = ApprovalService(db_path)
    experiments = ExperimentService(db_path)
    model_router = ModelRouterService(db_path)
    recovery_brain = RecoveryBrainService(
        db_path,
        model_router=model_router,
        experiments=experiments,
        approvals=approvals,
    )
    # Best-effort reindex on startup for demos (fast for small docs).
    try:
        knowledge.reindex()
    except Exception as exc:
        log_event(logger, "knowledge_reindex_startup_error", error=str(exc))

    _demo_state.update(
        {
            "audit": audit,
            "sink": sink,
            "source": source,
            "router": router,
            "knowledge": knowledge,
            "strategy": strategy,
            "compliance": compliance,
            "campaign_service": campaign_service,
            "followups": followups,
            "crm": crm,
            "zoho_crm": zoho_crm,
            "portfolio": portfolio,
            "workbench": workbench,
            "alerts": alerts,
            "reports": reports,
            "sync": sync,
            "tenant_service": tenants,
            "event_bus": event_bus,
            "customer_360": customer_360,
            "journey_orchestrator": journeys,
            "channel_hub": channel_hub,
            "payment_orchestration": payments,
            "settlement": settlements,
            "approval_service": approvals,
            "experiment_service": experiments,
            "model_router": model_router,
            "recovery_brain": recovery_brain,
            "campaign_tasks": {},
            "sessions": {},  # session_id -> snapshot dict
            "session_objs": {},  # session_id -> WebCallSession
        }
    )
    _start_background_loops()
    return _demo_state


def _start_background_loops() -> None:
    global _report_scheduler_task
    if _report_scheduler_task and not _report_scheduler_task.done():
        return
    try:
        loop = asyncio.get_running_loop()
    except RuntimeError:
        return
    _report_scheduler_task = loop.create_task(_background_tick_loop(), name="ops_background_tick")


async def _background_tick_loop() -> None:
    while True:
        try:
            demo = _get_demo_singletons()
            reports: ReportsService = demo["reports"]  # type: ignore[assignment]
            alerts: AlertsService = demo["alerts"]  # type: ignore[assignment]
            followups: FollowupService = demo["followups"]  # type: ignore[assignment]
            audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]

            due = followups.due_followups(limit=300)
            for row in due:
                try:
                    reminder_type = str(row.get("reminder_type") or "")
                    status = "missed" if reminder_type == "ptp_t_plus_1_miss" else "sent"
                    followups.mark_status(str(row["idempotency_key"]), status=status)
                    audit.record_event(
                        event_type="followup_" + status,
                        session_id=str(row.get("session_id") or ""),
                        payload={
                            "reminder_type": reminder_type,
                            "idempotency_key": row.get("idempotency_key"),
                            "channel": row.get("channel"),
                        },
                    )
                except Exception as exc:
                    log_event(logger, "followup_background_error", error=str(exc))
            with contextlib.suppress(Exception):
                alerts.evaluate()
            with contextlib.suppress(Exception):
                reports.run_scheduler_tick()
        except Exception as exc:
            log_event(logger, "background_tick_error", error=str(exc))
        await asyncio.sleep(30.0)


async def _campaign_run_loop(campaign_id: str) -> None:
    demo = _get_demo_singletons()
    campaign_service: CampaignService = demo["campaign_service"]  # type: ignore[assignment]
    tasks: Dict[str, asyncio.Task] = demo["campaign_tasks"]  # type: ignore[assignment]
    try:
        for _ in range(1200):  # bounded loop to avoid runaway background task
            meta = campaign_service.get_campaign(campaign_id)
            if not meta or meta.get("status") not in {"active", "created"}:
                return
            result = campaign_service.run_pending_batch(campaign_id)
            if int(result.get("processed", 0)) <= 0:
                await asyncio.sleep(0.2)
            else:
                await asyncio.sleep(0.02)
            meta = campaign_service.get_campaign(campaign_id)
            if meta and meta.get("status") == "completed":
                return
    finally:
        tasks.pop(campaign_id, None)


@app.get("/api/metrics")
async def api_metrics(request: Request):
    deny = _require_permissions(request, (VIEW_DASHBOARD,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    crm: CRMAdapter = demo["crm"]  # type: ignore[assignment]
    campaign_id = (request.query_params.get("campaign_id") or "").strip() or None
    data = audit.metrics(campaign_id=campaign_id)
    data["outbound_status"] = crm.status()
    data["followups_due_processed"] = 0
    data["outbound_tick"] = {"processed": 0, "acked": 0, "retry": 0, "dead_letter": 0}
    data["alerts_tick"] = {"evaluated": 0}
    data["reports_tick"] = {"scheduled": 0}
    data["build_info"] = _build_info_payload()
    return data


@app.post("/api/ops/tick")
async def api_ops_tick(request: Request):
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    followups: FollowupService = demo["followups"]  # type: ignore[assignment]
    crm: CRMAdapter = demo["crm"]  # type: ignore[assignment]
    alerts: AlertsService = demo["alerts"]  # type: ignore[assignment]
    reports: ReportsService = demo["reports"]  # type: ignore[assignment]
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    due = followups.due_followups(limit=300)
    processed = 0
    for row in due:
        try:
            reminder_type = str(row.get("reminder_type") or "")
            status = "missed" if reminder_type == "ptp_t_plus_1_miss" else "sent"
            followups.mark_status(str(row["idempotency_key"]), status=status)
            audit.record_event(
                event_type="followup_" + status,
                session_id=str(row.get("session_id") or ""),
                payload={
                    "reminder_type": reminder_type,
                    "idempotency_key": row.get("idempotency_key"),
                    "channel": row.get("channel"),
                },
            )
            processed += 1
        except Exception as exc:
            log_event(logger, "followup_ops_tick_error", error=str(exc))
    outbound = crm.process_queue(limit=120)
    alerts_out = alerts.evaluate()
    reports_out = reports.run_scheduler_tick()
    return {
        "ok": True,
        "request_id": _request_id(request),
        "tick": {
            "followups_due_processed": processed,
            "outbound": outbound,
            "alerts": alerts_out,
            "reports": reports_out,
        },
    }

@app.get("/api/customers")
async def api_customers(request: Request):
    deny = _require_permissions(request, (WORKBENCH_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    source: ExcelCustomerSource = demo["source"]  # type: ignore[assignment]
    try:
        rows = source.list_customers()
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=500)
    return {"rows": rows}


@app.post("/api/portfolio/upload")
async def api_portfolio_upload(request: Request, file: UploadFile = File(...)):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    portfolio: PortfolioService = demo["portfolio"]  # type: ignore[assignment]
    if not file.filename:
        return JSONResponse({"error": "filename_required"}, status_code=400)
    content = await file.read()
    try:
        result = portfolio.upload(filename=file.filename, content=content)
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    return {"ok": True, **result}


@app.post("/api/portfolio/{upload_id}/map")
async def api_portfolio_map(upload_id: str, request: Request):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    portfolio: PortfolioService = demo["portfolio"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    mappings = body.get("mappings") if isinstance(body.get("mappings"), dict) else {}
    portfolio_name = str(body.get("portfolio_name") or f"Portfolio {upload_id}")
    try:
        result = portfolio.map_columns(upload_id=upload_id, mappings=mappings, portfolio_name=portfolio_name)
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    return {"ok": True, **result}


@app.post("/api/portfolio/{portfolio_id}/validate")
async def api_portfolio_validate(portfolio_id: str, request: Request):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    portfolio: PortfolioService = demo["portfolio"]  # type: ignore[assignment]
    try:
        result = portfolio.validate(portfolio_id=portfolio_id)
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    return {"ok": True, **result}


@app.get("/api/portfolio/{portfolio_id}/preview")
async def api_portfolio_preview(portfolio_id: str, request: Request):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    portfolio: PortfolioService = demo["portfolio"]  # type: ignore[assignment]
    try:
        limit = int(request.query_params.get("limit") or 20)
    except Exception:
        limit = 20
    rows = portfolio.preview(portfolio_id=portfolio_id, limit=limit)
    return {"rows": rows, "limit": max(1, limit)}


@app.get("/api/portfolio/{portfolio_id}/errors.csv")
async def api_portfolio_errors_csv(portfolio_id: str, request: Request):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    portfolio: PortfolioService = demo["portfolio"]  # type: ignore[assignment]
    path = portfolio.get_errors_csv_path(portfolio_id)
    if not path or not os.path.exists(path):
        return JSONResponse({"error": "errors_csv_not_found"}, status_code=404)
    return FileResponse(path, filename=f"{portfolio_id}-errors.csv")


@app.post("/api/portfolio/{portfolio_id}/exclusions/upload")
async def api_portfolio_exclusions_upload(portfolio_id: str, request: Request, file: UploadFile = File(...)):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    portfolio: PortfolioService = demo["portfolio"]  # type: ignore[assignment]
    content = await file.read()
    try:
        out = portfolio.add_exclusions(portfolio_id=portfolio_id, filename=file.filename or "exclusions.csv", content=content)
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    return {"ok": True, **out}


@app.post("/api/portfolio/{portfolio_id}/launch")
async def api_portfolio_launch(portfolio_id: str, request: Request):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    prod_deny = _simulation_blocked_in_prod(request)
    if prod_deny:
        return prod_deny
    demo = _get_demo_singletons()
    portfolio: PortfolioService = demo["portfolio"]  # type: ignore[assignment]
    campaign_service: CampaignService = demo["campaign_service"]  # type: ignore[assignment]
    workbench: WorkbenchService = demo["workbench"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}

    campaign_name = str(body.get("campaign_name") or f"Campaign {portfolio_id}")
    max_attempts = int(body.get("retry_policy", {}).get("max_attempts", body.get("max_attempts", 3)) or 3)
    retry_delay_minutes = int(body.get("retry_policy", {}).get("retry_delay_minutes", body.get("retry_delay_minutes", 30)) or 30)
    batch_size = int(body.get("throttle", body.get("batch_size", 250)) or 250)
    exclusion_id = body.get("exclusion_upload_id")
    exclude_predicate = body.get("exclude_predicate")

    try:
        candidates = portfolio.launch_candidates(
            portfolio_id=portfolio_id,
            exclusion_id=str(exclusion_id).strip() if exclusion_id else None,
            exclude_predicate=str(exclude_predicate).strip() if exclude_predicate else None,
        )
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)

    selected_rows = candidates["selected_rows"]
    customer_ids = [str(r.get("customer_id") or "").strip() for r in selected_rows if r.get("customer_id")]
    created = campaign_service.create_campaign(
        name=campaign_name,
        customer_ids=customer_ids,
        max_attempts=max_attempts,
        retry_delay_minutes=retry_delay_minutes,
        batch_size=batch_size,
    )
    campaign_id = str(created["campaign_id"])
    launch_config = {
        "retry_policy": body.get("retry_policy") or {"max_attempts": max_attempts, "retry_delay_minutes": retry_delay_minutes},
        "contact_window": body.get("contact_window") or {"start": "09:00", "end": "20:00"},
        "throttle": batch_size,
        "dpd_bucket_strategies": body.get("dpd_bucket_strategies") or {},
        "exclude_predicate": exclude_predicate,
        "exclusion_upload_id": exclusion_id,
        "excluded_count": candidates.get("excluded_count", 0),
    }
    launch_id = portfolio.record_launch(portfolio_id=portfolio_id, campaign_id=campaign_id, launch_config=launch_config, status="launched")
    tasks_created = workbench.seed_tasks(campaign_id=campaign_id, portfolio_id=portfolio_id, rows=selected_rows, actor=_actor(request))
    metrics = campaign_service.metrics(campaign_id)
    return {
        "ok": True,
        "campaign_id": campaign_id,
        "launch_id": launch_id,
        "seeded_accounts": len(customer_ids),
        "tasks_created": tasks_created,
        "excluded_count": candidates.get("excluded_count", 0),
        "metrics": metrics,
    }


@app.get("/api/portfolios")
async def api_portfolios(request: Request):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    portfolio: PortfolioService = demo["portfolio"]  # type: ignore[assignment]
    page, page_size = _pagination(request)
    out = portfolio.list_portfolios(limit=page_size, offset=(page - 1) * page_size)
    return {"rows": out["rows"], "total": out["total"], "page": page, "page_size": page_size}


@app.get("/api/portfolio/{portfolio_id}")
async def api_portfolio_get(portfolio_id: str, request: Request):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    portfolio: PortfolioService = demo["portfolio"]  # type: ignore[assignment]
    row = portfolio.get_portfolio(portfolio_id)
    if not row:
        return JSONResponse({"error": "portfolio_not_found"}, status_code=404)
    return row


@app.get("/api/sessions")
async def api_sessions(request: Request):
    deny = _require_permissions(request, (WORKBENCH_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    sessions: Dict[str, dict] = demo["sessions"]  # type: ignore[assignment]
    campaign_id = (request.query_params.get("campaign_id") or "").strip()
    dpd_bucket = (request.query_params.get("bucket") or "").strip()
    disposition = (request.query_params.get("disposition") or "").strip()
    rows = list(sessions.values())
    if campaign_id:
        rows = [s for s in rows if str(s.get("campaign_id") or "") == campaign_id]
    if dpd_bucket:
        rows = [s for s in rows if str(s.get("dpd_bucket") or "") == dpd_bucket]
    if disposition:
        rows = [s for s in rows if str(s.get("disposition") or "") == disposition]
    # Return "active" snapshots (updated by the WS session loop).
    return {"sessions": rows}


@app.get("/api/tasks")
async def api_tasks(request: Request):
    deny = _require_permissions(request, (WORKBENCH_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    workbench: WorkbenchService = demo["workbench"]  # type: ignore[assignment]
    page, page_size = _pagination(request)
    role = _get_role(request)
    actor = _actor(request)
    owner = actor if role == "CALLING_AGENT" else None
    out = workbench.list_tasks(
        campaign_id=(request.query_params.get("campaign_id") or "").strip() or None,
        state=(request.query_params.get("state") or "").strip() or None,
        dpd_bucket=(request.query_params.get("dpd_bucket") or "").strip() or None,
        q=(request.query_params.get("q") or "").strip() or None,
        sort=(request.query_params.get("sort") or "updated_desc").strip(),
        page=page,
        page_size=page_size,
        owner=owner,
    )
    return out


@app.get("/api/tasks/summary")
async def api_tasks_summary(request: Request):
    deny = _require_permissions(request, (WORKBENCH_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    workbench: WorkbenchService = demo["workbench"]  # type: ignore[assignment]
    campaign_id = (request.query_params.get("campaign_id") or "").strip() or None
    return workbench.summary_by_state(campaign_id=campaign_id)


@app.get("/api/tasks/{task_id}")
async def api_task_get(task_id: str, request: Request):
    deny = _require_permissions(request, (WORKBENCH_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    workbench: WorkbenchService = demo["workbench"]  # type: ignore[assignment]
    task = workbench.get_task(task_id)
    if not task:
        return JSONResponse({"error": "task_not_found"}, status_code=404)
    role = _get_role(request)
    actor = _actor(request)
    if role == "CALLING_AGENT" and str(task.get("owner") or "") != actor:
        return JSONResponse({"error": "forbidden_not_assigned"}, status_code=403)
    return task


@app.post("/api/tasks/{task_id}/claim")
async def api_task_claim(task_id: str, request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    workbench: WorkbenchService = demo["workbench"]  # type: ignore[assignment]
    role = _get_role(request)
    actor = _actor(request)
    if role == "CALLING_AGENT":
        existing = workbench.get_task(task_id)
        if not existing:
            return JSONResponse({"ok": False, "error": "task_not_found"}, status_code=404)
        owner = str(existing.get("owner") or "")
        if owner and owner != actor:
            return JSONResponse({"ok": False, "error": "forbidden_not_assigned"}, status_code=403)
    result = workbench.claim_task(task_id=task_id, actor=actor)
    status = 200 if result.get("ok") else 409
    return JSONResponse(result, status_code=status)


@app.post("/api/tasks/{task_id}/update")
async def api_task_update(task_id: str, request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    workbench: WorkbenchService = demo["workbench"]  # type: ignore[assignment]
    role = _get_role(request)
    actor = _actor(request)
    if role == "CALLING_AGENT":
        existing = workbench.get_task(task_id)
        if not existing:
            return JSONResponse({"ok": False, "error": "task_not_found"}, status_code=404)
        owner = str(existing.get("owner") or "")
        if owner and owner != actor:
            return JSONResponse({"ok": False, "error": "forbidden_not_assigned"}, status_code=403)
        if not owner:
            # Calling agents are limited to assigned work; auto-claim first update.
            workbench.claim_task(task_id=task_id, actor=actor)
    try:
        body = await request.json()
    except Exception:
        body = {}
    result = workbench.update_task(
        task_id=task_id,
        actor=actor,
        role=role,
        state=body.get("state"),
        disposition=body.get("disposition"),
        notes=body.get("notes"),
        callback_at=body.get("callback_at"),
        compliance_override=bool(body.get("compliance_override", False)),
        escalate_reason=body.get("escalate_reason"),
    )
    status = 200 if result.get("ok") else 400
    return JSONResponse(result, status_code=status)


@app.post("/api/telephony/test_call")
async def api_telephony_test_call(request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        raw_body = await request.json()
    except Exception:
        raw_body = {}
    try:
        body_obj = TelephonyTestCallRequest.model_validate(raw_body if isinstance(raw_body, dict) else {})
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid telephony test call payload",
            details={"errors": exc.errors()},
        )

    demo = _get_demo_singletons()
    router: ActionRouter = demo["router"]  # type: ignore[assignment]
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    sink: ExcelOutcomeSink = demo["sink"]  # type: ignore[assignment]
    actor = _actor(request)
    payload = body_obj.model_dump(exclude_none=True)
    call_result = router.place_test_voice_call(
        customer_phone=body_obj.phone,
        customer_name=body_obj.customer_name,
        amount=body_obj.amount_due,
        timeout_s=int(body_obj.timeout_seconds or 25),
    )
    ok = bool(call_result.get("delivery_ok"))
    session_id = f"tel-{uuid.uuid4().hex[:10]}"

    with contextlib.suppress(Exception):
        sink.log_action(session_id=session_id, action_name="test_voice_call", payload=payload, result=call_result, ok=ok)
    with contextlib.suppress(Exception):
        sink.log_outbox(
            session_id=session_id,
            channel="voice",
            to=str(call_result.get("normalized_to") or call_result.get("to") or body_obj.phone),
            body=f"test_call:{call_result.get('call_status') or call_result.get('delivery_status') or 'unknown'}",
            link=str(call_result.get("call_sid") or ""),
        )
    with contextlib.suppress(Exception):
        audit.record_event(
            event_type="action",
            session_id=session_id,
            payload={
                "name": "test_voice_call",
                "actor": actor,
                "payload": payload,
                "result": {
                    "provider": call_result.get("provider"),
                    "delivery_ok": call_result.get("delivery_ok"),
                    "delivery_status": call_result.get("delivery_status"),
                    "delivery_error": call_result.get("delivery_error"),
                    "call_sid": call_result.get("call_sid"),
                    "normalized_to": call_result.get("normalized_to"),
                },
            },
        )

    status = 200 if ok else 502
    return JSONResponse({"ok": ok, "result": call_result}, status_code=status)


@app.post("/api/telephony/agent_call")
async def api_telephony_agent_call(request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        raw_body = await request.json()
    except Exception:
        raw_body = {}
    try:
        body_obj = TelephonyAgentCallRequest.model_validate(raw_body if isinstance(raw_body, dict) else {})
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid telephony agent call payload",
            details={"errors": exc.errors()},
        )
    stream_ws_url = _telephony_stream_ws_url(request)
    if not stream_ws_url.startswith("wss://"):
        return _error_response(
            request,
            status_code=400,
            code="telephony_stream_url_missing",
            message="Set TELEPHONY_PUBLIC_BASE_URL or TELEPHONY_STREAM_WSS_URL to a public wss:// endpoint",
            details={"resolved_stream_ws_url": stream_ws_url},
        )

    demo = _get_demo_singletons()
    router: ActionRouter = demo["router"]  # type: ignore[assignment]
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    sink: ExcelOutcomeSink = demo["sink"]  # type: ignore[assignment]
    actor = _actor(request)
    payload = body_obj.model_dump(exclude_none=True)
    call_result = router.place_agent_stream_call(
        customer_phone=body_obj.phone,
        stream_ws_url=stream_ws_url,
        customer_name=body_obj.customer_name,
        amount=body_obj.amount_due,
        customer_id=body_obj.customer_id,
        campaign_id=body_obj.campaign_id,
        language=body_obj.language,
        tts_speaker=body_obj.tts_speaker,
        timeout_s=int(body_obj.timeout_seconds or 25),
    )
    ok = bool(call_result.get("delivery_ok"))
    session_id = f"tel-agent-{uuid.uuid4().hex[:10]}"
    with contextlib.suppress(Exception):
        sink.log_action(session_id=session_id, action_name="agent_voice_call", payload=payload, result=call_result, ok=ok)
    with contextlib.suppress(Exception):
        sink.log_outbox(
            session_id=session_id,
            channel="voice",
            to=str(call_result.get("normalized_to") or call_result.get("to") or body_obj.phone),
            body=f"agent_call:{call_result.get('call_status') or call_result.get('delivery_status') or 'unknown'}",
            link=str(call_result.get("call_sid") or ""),
        )
    with contextlib.suppress(Exception):
        audit.record_event(
            event_type="action",
            session_id=session_id,
            payload={
                "name": "agent_voice_call",
                "actor": actor,
                "payload": payload,
                "result": {
                    "provider": call_result.get("provider"),
                    "delivery_ok": call_result.get("delivery_ok"),
                    "delivery_status": call_result.get("delivery_status"),
                    "delivery_error": call_result.get("delivery_error"),
                    "call_sid": call_result.get("call_sid"),
                    "normalized_to": call_result.get("normalized_to"),
                    "stream_ws_url": call_result.get("stream_ws_url"),
                },
            },
        )
    status = 200 if ok else 502
    return JSONResponse({"ok": ok, "result": call_result}, status_code=status)


@app.post("/api/tasks/bulk_update")
async def api_tasks_bulk_update(request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    if _get_role(request) == "CALLING_AGENT":
        return JSONResponse({"ok": False, "error": "bulk_update_not_allowed_for_calling_agent"}, status_code=403)
    demo = _get_demo_singletons()
    workbench: WorkbenchService = demo["workbench"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    ids = body.get("ids") if isinstance(body.get("ids"), list) else []
    action = str(body.get("action") or "").strip()
    payload = body.get("payload") if isinstance(body.get("payload"), dict) else {}
    result = workbench.bulk_update(ids=[str(i) for i in ids], actor=_actor(request), role=_get_role(request), action=action, payload=payload)
    return result


@app.get("/api/sessions/{session_id}/timeline")
async def api_session_timeline(session_id: str, request: Request):
    deny = _require_permissions(request, (WORKBENCH_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    session_objs: Dict[str, WebCallSession] = demo["session_objs"]  # type: ignore[assignment]
    timeline = audit.session_timeline(session_id)
    active = session_objs.get(session_id)
    if active:
        for ev in active.get_timeline():
            if isinstance(ev, dict):
                timeline.append(dict(ev))
        timeline.sort(key=lambda x: float(x.get("ts") or 0))
    return {"session_id": session_id, "timeline": timeline}


@app.get("/api/compliance/violations")
async def api_compliance_violations(request: Request):
    deny = _require_permissions(request, (ALERTS_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    severity = (request.query_params.get("severity") or "").strip() or None
    session_id = (request.query_params.get("session_id") or "").strip() or None
    return {"rows": audit.list_compliance_violations(severity=severity, session_id=session_id, limit=300)}


@app.get("/api/alerts")
async def api_alerts(request: Request):
    deny = _require_permissions(request, (ALERTS_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    alerts: AlertsService = demo["alerts"]  # type: ignore[assignment]
    page, page_size = _pagination(request)
    return alerts.list_alerts(
        status=(request.query_params.get("status") or "").strip() or None,
        severity=(request.query_params.get("severity") or "").strip() or None,
        rtype=(request.query_params.get("type") or "").strip() or None,
        page=page,
        page_size=page_size,
    )


@app.get("/api/alerts/{alert_id}")
async def api_alert_get(alert_id: str, request: Request):
    deny = _require_permissions(request, (ALERTS_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    alerts: AlertsService = demo["alerts"]  # type: ignore[assignment]
    row = alerts.get_alert(alert_id)
    if not row:
        return JSONResponse({"error": "alert_not_found"}, status_code=404)
    return row


@app.post("/api/alerts/{alert_id}/ack")
async def api_alert_ack(alert_id: str, request: Request):
    deny = _require_permissions(request, (ALERTS_MUTATE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    alerts: AlertsService = demo["alerts"]  # type: ignore[assignment]
    ok = alerts.ack_alert(alert_id=alert_id, actor=_actor(request))
    return {"ok": ok}


@app.post("/api/alerts/{alert_id}/assign")
async def api_alert_assign(alert_id: str, request: Request):
    deny = _require_permissions(request, (ALERTS_MUTATE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    alerts: AlertsService = demo["alerts"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    assignee = str(body.get("assignee") or "").strip()
    if not assignee:
        return JSONResponse({"ok": False, "error": "assignee_required"}, status_code=400)
    ok = alerts.assign_alert(alert_id=alert_id, assignee=assignee, actor=_actor(request))
    return {"ok": ok}


@app.post("/api/alerts/{alert_id}/resolve")
async def api_alert_resolve(alert_id: str, request: Request):
    deny = _require_permissions(request, (ALERTS_MUTATE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    alerts: AlertsService = demo["alerts"]  # type: ignore[assignment]
    ok = alerts.resolve_alert(alert_id=alert_id, actor=_actor(request))
    return {"ok": ok}


@app.get("/api/alert_rules")
async def api_alert_rules(request: Request):
    deny = _require_permissions(request, (ALERTS_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    alerts: AlertsService = demo["alerts"]  # type: ignore[assignment]
    return {"rows": alerts.list_rules()}


@app.post("/api/alert_rules")
async def api_alert_rules_upsert(request: Request):
    deny = _require_permissions(request, (ALERT_RULES_EDIT,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    alerts: AlertsService = demo["alerts"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    try:
        row = alerts.upsert_rule(
            rule_id=body.get("id"),
            name=str(body.get("name") or "Rule"),
            rtype=str(body.get("type") or ""),
            enabled=bool(body.get("enabled", True)),
            threshold=body.get("threshold_json") if isinstance(body.get("threshold_json"), dict) else {},
            routing=body.get("routing_json") if isinstance(body.get("routing_json"), dict) else {},
        )
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    return {"ok": True, "rule": row}


@app.post("/api/alert_rules/{rule_id}/toggle")
async def api_alert_rule_toggle(rule_id: str, request: Request):
    deny = _require_permissions(request, (ALERT_RULES_EDIT,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    alerts: AlertsService = demo["alerts"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    ok = alerts.toggle_rule(rule_id=rule_id, enabled=bool(body.get("enabled", False)))
    return {"ok": ok}


@app.post("/api/alerts/evaluate")
async def api_alerts_evaluate(request: Request):
    deny = _require_permissions(request, (ALERTS_MUTATE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    alerts: AlertsService = demo["alerts"]  # type: ignore[assignment]
    return {"ok": True, "result": alerts.evaluate()}


@app.post("/api/campaigns")
async def api_create_campaign(request: Request):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    prod_deny = _simulation_blocked_in_prod(request)
    if prod_deny:
        return prod_deny
    demo = _get_demo_singletons()
    source: ExcelCustomerSource = demo["source"]  # type: ignore[assignment]
    campaign_service: CampaignService = demo["campaign_service"]  # type: ignore[assignment]
    workbench: WorkbenchService = demo["workbench"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    name = (body.get("name") or f"Portfolio Campaign {int(time.time())}").strip()
    max_accounts = int(body.get("max_accounts") or 10000)
    max_attempts = int(body.get("max_attempts") or 3)
    retry_delay_minutes = int(body.get("retry_delay_minutes") or 30)
    batch_size = int(body.get("batch_size") or 250)
    rows = source.list_customers()
    selected_rows = rows[: max(1, max_accounts)]
    customer_ids = [str(r.get("customer_id") or "").strip() for r in selected_rows]
    created = campaign_service.create_campaign(
        name=name,
        customer_ids=customer_ids,
        max_attempts=max_attempts,
        retry_delay_minutes=retry_delay_minutes,
        batch_size=batch_size,
    )
    task_rows = [
        {
            "customer_id": r.get("customer_id"),
            "phone": r.get("phone"),
            "amount_due": r.get("overdue_amount"),
            "dpd": r.get("dpd"),
            "language": r.get("language_preference"),
            "customer_name": r.get("customer_name"),
        }
        for r in selected_rows
    ]
    tasks_created = workbench.seed_tasks(
        campaign_id=str(created["campaign_id"]),
        portfolio_id="legacy_source",
        rows=task_rows,
        actor=_actor(request),
    )
    return {"ok": True, **created, "tasks_created": tasks_created}


@app.post("/api/campaigns/{campaign_id}/start")
async def api_start_campaign(campaign_id: str, request: Request):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    prod_deny = _simulation_blocked_in_prod(request)
    if prod_deny:
        return prod_deny
    demo = _get_demo_singletons()
    campaign_service: CampaignService = demo["campaign_service"]  # type: ignore[assignment]
    tasks: Dict[str, asyncio.Task] = demo["campaign_tasks"]  # type: ignore[assignment]
    campaign_service.set_status(campaign_id, "active")
    if campaign_id not in tasks or tasks[campaign_id].done():
        tasks[campaign_id] = asyncio.create_task(_campaign_run_loop(campaign_id), name=f"campaign_{campaign_id}")
    return {"ok": True, "campaign_id": campaign_id, "status": "active"}


@app.post("/api/campaigns/{campaign_id}/pause")
async def api_pause_campaign(campaign_id: str, request: Request):
    deny = _require_permissions(request, (PORTFOLIO_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    campaign_service: CampaignService = demo["campaign_service"]  # type: ignore[assignment]
    tasks: Dict[str, asyncio.Task] = demo["campaign_tasks"]  # type: ignore[assignment]
    campaign_service.set_status(campaign_id, "paused")
    t = tasks.get(campaign_id)
    if t and not t.done():
        t.cancel()
    tasks.pop(campaign_id, None)
    return {"ok": True, "campaign_id": campaign_id, "status": "paused"}


@app.get("/api/campaigns/{campaign_id}/metrics")
async def api_campaign_metrics(campaign_id: str, request: Request):
    deny = _require_permissions(request, (VIEW_DASHBOARD,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    campaign_service: CampaignService = demo["campaign_service"]  # type: ignore[assignment]
    return campaign_service.metrics(campaign_id)


@app.get("/api/campaigns")
async def api_campaigns(request: Request):
    deny = _require_permissions(request, (VIEW_DASHBOARD,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    campaign_service: CampaignService = demo["campaign_service"]  # type: ignore[assignment]
    return {"rows": campaign_service.list_campaigns(limit=50)}


@app.get("/api/outbound/status")
async def api_outbound_status(request: Request):
    deny = _require_permissions(request, (INTEGRATIONS_VIEW,))
    if deny:
        return deny
    pilot_deny = _require_pilot_mode()
    if pilot_deny:
        return pilot_deny
    demo = _get_demo_singletons()
    crm: CRMAdapter = demo["crm"]  # type: ignore[assignment]
    return crm.status()


@app.get("/api/reports")
async def api_reports(request: Request):
    deny = _require_permissions(request, (REPORTS_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    reports: ReportsService = demo["reports"]  # type: ignore[assignment]
    page, page_size = _pagination(request)
    campaign_id = (request.query_params.get("campaign_id") or "").strip() or None
    out = reports.list_reports(page=page, page_size=page_size, campaign_id=campaign_id)
    out["schedule"] = reports.get_schedule()
    return out


@app.post("/api/reports/generate")
async def api_reports_generate(request: Request):
    deny = _require_permissions(request, (REPORTS_SCHEDULE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    reports: ReportsService = demo["reports"]  # type: ignore[assignment]
    campaign_service: CampaignService = demo["campaign_service"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    campaign_id = (body.get("campaign_id") or "").strip()
    report_date = (body.get("report_date") or "").strip() or None
    rows = []
    if campaign_id:
        rows = [campaign_id]
    else:
        rows = [str(r.get("campaign_id")) for r in campaign_service.list_campaigns(limit=500) if r.get("campaign_id")]
    created = []
    for cid in rows:
        with contextlib.suppress(Exception):
            created.append(reports.generate_report(campaign_id=cid, report_date=report_date))
    return {"ok": True, "count": len(created), "rows": created}


@app.post("/api/reports/schedule")
async def api_reports_schedule(request: Request):
    deny = _require_permissions(request, (REPORTS_SCHEDULE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    reports: ReportsService = demo["reports"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    enabled = bool(body.get("enabled", True))
    daily_time = str(body.get("daily_time") or "09:00")
    return {"ok": True, "schedule": reports.set_schedule(enabled=enabled, daily_time=daily_time)}


@app.get("/api/reports/{report_id}/download")
async def api_reports_download(report_id: str, request: Request):
    deny = _require_permissions(request, (REPORTS_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    reports: ReportsService = demo["reports"]  # type: ignore[assignment]
    row = reports.get_report(report_id)
    if not row:
        return JSONResponse({"error": "report_not_found"}, status_code=404)
    fmt = (request.query_params.get("format") or "json").strip().lower()
    if fmt == "csv":
        return FileResponse(row["path_csv"], filename=f"{report_id}.csv")
    return FileResponse(row["path_json"], filename=f"{report_id}.json")


@app.post("/api/integrations/crm/replay/{session_id}")
async def api_crm_replay(session_id: str, request: Request):
    deny = _require_permissions(request, (INTEGRATIONS_RESOLVE_CONFLICTS,))
    if deny:
        return deny
    pilot_deny = _require_pilot_mode()
    if pilot_deny:
        return pilot_deny
    demo = _get_demo_singletons()
    crm: CRMAdapter = demo["crm"]  # type: ignore[assignment]
    return crm.replay_session(session_id)


@app.post("/api/integrations/crm/mock")
async def api_crm_mock(request: Request):
    deny = _require_permissions(request, (INTEGRATIONS_RESOLVE_CONFLICTS,))
    if deny:
        return deny
    prod_deny = _simulation_blocked_in_prod(request)
    if prod_deny:
        return prod_deny
    pilot_deny = _require_pilot_mode()
    if pilot_deny:
        return pilot_deny
    demo = _get_demo_singletons()
    crm: CRMAdapter = demo["crm"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    should_fail = bool(body.get("should_fail", False))
    payload = body.get("payload") if isinstance(body.get("payload"), dict) else body
    return crm.mock_receive(payload=payload, should_fail=should_fail)


@app.post("/api/integrations/lms/inbound")
async def api_lms_inbound(request: Request):
    deny = _require_permissions(request, (INTEGRATIONS_VIEW,))
    if deny:
        return deny
    pilot_deny = _require_pilot_mode()
    if pilot_deny:
        return pilot_deny
    demo = _get_demo_singletons()
    sync: SyncService = demo["sync"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    try:
        out = sync.inbound_lms(body, actor=_actor(request))
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    return out


@app.get("/api/integrations/sync_events")
async def api_sync_events(request: Request):
    deny = _require_permissions(request, (INTEGRATIONS_VIEW,))
    if deny:
        return deny
    pilot_deny = _require_pilot_mode()
    if pilot_deny:
        return pilot_deny
    demo = _get_demo_singletons()
    sync: SyncService = demo["sync"]  # type: ignore[assignment]
    page, page_size = _pagination(request)
    return sync.list_sync_events(
        direction=(request.query_params.get("direction") or "").strip() or None,
        status=(request.query_params.get("status") or "").strip() or None,
        page=page,
        page_size=page_size,
    )


@app.get("/api/integrations/conflicts")
async def api_sync_conflicts(request: Request):
    deny = _require_permissions(request, (INTEGRATIONS_VIEW,))
    if deny:
        return deny
    pilot_deny = _require_pilot_mode()
    if pilot_deny:
        return pilot_deny
    demo = _get_demo_singletons()
    sync: SyncService = demo["sync"]  # type: ignore[assignment]
    page, page_size = _pagination(request)
    return sync.list_conflicts(
        status=(request.query_params.get("status") or "").strip() or None,
        page=page,
        page_size=page_size,
    )


@app.post("/api/integrations/conflicts/{conflict_id}/resolve")
async def api_sync_conflict_resolve(conflict_id: str, request: Request):
    deny = _require_permissions(request, (INTEGRATIONS_RESOLVE_CONFLICTS,))
    if deny:
        return deny
    pilot_deny = _require_pilot_mode()
    if pilot_deny:
        return pilot_deny
    demo = _get_demo_singletons()
    sync: SyncService = demo["sync"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    action = str(body.get("action") or "").strip().upper()
    try:
        out = sync.resolve_conflict(
            conflict_id=conflict_id,
            action=action,
            resolved_by=_actor(request),
            manual_override=body.get("manual_override") if isinstance(body.get("manual_override"), dict) else None,
        )
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    return out


@app.get("/api/integrations/dead_letters")
async def api_sync_dead_letters(request: Request):
    deny = _require_permissions(request, (INTEGRATIONS_VIEW,))
    if deny:
        return deny
    pilot_deny = _require_pilot_mode()
    if pilot_deny:
        return pilot_deny
    demo = _get_demo_singletons()
    sync: SyncService = demo["sync"]  # type: ignore[assignment]
    page, page_size = _pagination(request)
    return sync.list_dead_letters(page=page, page_size=page_size)


@app.post("/api/integrations/dead_letters/{queue_id}/replay")
async def api_sync_dead_letter_replay(queue_id: str, request: Request):
    deny = _require_permissions(request, (INTEGRATIONS_RESOLVE_CONFLICTS,))
    if deny:
        return deny
    pilot_deny = _require_pilot_mode()
    if pilot_deny:
        return pilot_deny
    demo = _get_demo_singletons()
    sync: SyncService = demo["sync"]  # type: ignore[assignment]
    return sync.replay_dead_letter(queue_id, actor=_actor(request))


@app.post("/api/sessions/{session_id}/disposition")
async def api_set_disposition(session_id: str, request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    objs: Dict[str, WebCallSession] = demo["session_objs"]  # type: ignore[assignment]
    sessions: Dict[str, dict] = demo["sessions"]  # type: ignore[assignment]
    sess = objs.get(session_id)
    if not sess:
        return JSONResponse({"error": "unknown session"}, status_code=404)
    try:
        data = await request.json()
    except Exception:
        data = {}
    disp = (data.get("disposition") or "").strip() or "closed"
    try:
        await sess.admin_set_disposition(disp)
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    sessions[session_id] = sess.get_snapshot()
    return {"ok": True, "session": sessions[session_id]}


@app.get("/api/export/demo.xlsx")
async def api_export_demo(request: Request):
    deny = _require_permissions(request, (REPORTS_VIEW,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    sink: ExcelOutcomeSink = demo["sink"]  # type: ignore[assignment]
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]

    # Build a one-off export file with a Metrics sheet (client-friendly demo artifact).
    try:
        from openpyxl import load_workbook
    except Exception:
        return FileResponse(sink.path, filename="demo.xlsx")

    export_path = os.path.join(DATA_DIR, "demo_export.xlsx")
    wb = load_workbook(sink.path)
    try:
        if "Metrics" in wb.sheetnames:
            ws = wb["Metrics"]
            ws.delete_rows(1, ws.max_row)
        else:
            ws = wb.create_sheet("Metrics")
        m = audit.metrics()
        ws.append(["key", "value"])
        for k, v in m.items():
            ws.append([k, v])
        wb.save(export_path)
    finally:
        wb.close()

    return FileResponse(export_path, filename="demo.xlsx")


@app.post("/api/knowledge/reindex")
async def api_knowledge_reindex(request: Request):
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    knowledge: SQLiteFTSKnowledgeStore = demo["knowledge"]  # type: ignore[assignment]
    n = knowledge.reindex()
    return {"ok": True, "chunks_indexed": n}


def _sync_task_linkages(
    *,
    task_id: Optional[str],
    journey_id: Optional[str] = None,
    loan_account_id: Optional[str] = None,
    nba_decision_id: Optional[str] = None,
    approval_state: Optional[str] = None,
) -> None:
    if not task_id:
        return
    demo = _get_demo_singletons()
    workbench: WorkbenchService = demo["workbench"]  # type: ignore[assignment]
    conn = workbench._connect()  # type: ignore[attr-defined]
    try:
        sets: List[str] = []
        args: List[Any] = []
        if journey_id:
            sets.append("journey_id = ?")
            args.append(journey_id)
        if loan_account_id:
            sets.append("loan_account_id = ?")
            args.append(loan_account_id)
        if nba_decision_id:
            sets.append("current_nba_decision_id = ?")
            args.append(nba_decision_id)
        if approval_state:
            sets.append("approval_state = ?")
            args.append(approval_state)
        if not sets:
            return
        sets.append("updated_at = ?")
        args.append(time.time())
        args.append(task_id)
        conn.execute(f"UPDATE tasks SET {', '.join(sets)} WHERE id = ?", tuple(args))
        conn.commit()
    finally:
        conn.close()


@app.post("/api/v2/tenants")
async def api_v2_create_tenant(request: Request):
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    tenants: TenantService = demo["tenant_service"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    tenant_id = str(body.get("tenant_id") or "").strip().lower()
    name = str(body.get("name") or tenant_id or "").strip()
    metadata = body.get("metadata") if isinstance(body.get("metadata"), dict) else {}
    settings = body.get("settings") if isinstance(body.get("settings"), dict) else {}
    if not tenant_id:
        return JSONResponse({"error": "tenant_id_required"}, status_code=400)
    try:
        row = tenants.create_tenant(
            tenant_id=tenant_id,
            name=name or tenant_id,
            actor=_actor(request),
            request_id=_request_id(request),
            metadata=metadata,
            settings=settings,
        )
    except ValueError as exc:
        code = 409 if str(exc) == "tenant_exists" else 400
        return JSONResponse({"error": str(exc)}, status_code=code)
    ctx = _get_auth_context(request)
    user = ctx.get("user") if isinstance(ctx.get("user"), dict) else {}
    user_id = str((user or {}).get("id") or "").strip()
    if user_id:
        with contextlib.suppress(Exception):
            audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
            audit.add_user_tenant_membership(
                user_id=user_id,
                tenant_id=tenant_id,
                is_default=False,
                actor=_actor(request),
            )
    return {"schema_version": "v2.1", "tenant": row}


@app.get("/api/v2/tenants/{tenant_id}")
async def api_v2_get_tenant(tenant_id: str, request: Request):
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    tenants: TenantService = demo["tenant_service"]  # type: ignore[assignment]
    try:
        row = tenants.get_tenant(tenant_id)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    if not row:
        return JSONResponse({"error": "tenant_not_found"}, status_code=404)
    return {"schema_version": "v2.1", "tenant": row}


@app.patch("/api/v2/tenants/{tenant_id}/settings")
async def api_v2_patch_tenant_settings(tenant_id: str, request: Request):
    deny = _require_permissions(request, (USERS_MANAGE,))
    if deny:
        return deny
    demo = _get_demo_singletons()
    tenants: TenantService = demo["tenant_service"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    patch = body.get("settings") if isinstance(body.get("settings"), dict) else body
    if not isinstance(patch, dict):
        return JSONResponse({"error": "settings_patch_required"}, status_code=400)
    try:
        row = tenants.patch_tenant_settings(
            tenant_id=tenant_id,
            patch=patch,
            actor=_actor(request),
            request_id=_request_id(request),
        )
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    return {"schema_version": "v2.1", "tenant": row}


@app.post("/api/v2/customers/import")
async def api_v2_customers_import(request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    customer_360: Customer360Service = demo["customer_360"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    rows = body.get("rows") if isinstance(body.get("rows"), list) else []
    if not rows:
        return JSONResponse({"error": "rows_required"}, status_code=400)
    out = customer_360.import_customers(
        tenant_id=ctx["tenant_id"],
        rows=[r for r in rows if isinstance(r, dict)],
        actor=ctx["actor"],
        request_id=ctx["request_id"],
    )
    event_bus.publish(
        tenant_id=ctx["tenant_id"],
        event_type="customers.imported",
        payload={"counts": out},
        actor=ctx["actor"],
        request_id=ctx["request_id"],
    )
    return {"schema_version": "v2.1", "result": out}


@app.get("/api/v2/customers/{customer_id}/360")
async def api_v2_customer_360(customer_id: str, request: Request):
    deny = _require_permissions(request, (WORKBENCH_VIEW,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    customer_360: Customer360Service = demo["customer_360"]  # type: ignore[assignment]
    out = customer_360.get_customer_360(tenant_id=ctx["tenant_id"], customer_id=customer_id)
    return {"schema_version": "v2.1", **out}


@app.post("/api/v2/loan_accounts/upsert")
async def api_v2_loan_accounts_upsert(request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    customer_360: Customer360Service = demo["customer_360"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    try:
        out = customer_360.upsert_loan_account(
            tenant_id=ctx["tenant_id"],
            payload=body,
            actor=ctx["actor"],
            request_id=ctx["request_id"],
        )
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    event_bus.publish(
        tenant_id=ctx["tenant_id"],
        event_type="loan_account.upserted",
        payload={"loan_account_id": out.get("loan_account_id"), "customer_id": out.get("customer_id")},
        actor=ctx["actor"],
        request_id=ctx["request_id"],
    )
    return {"schema_version": "v2.1", "loan_account": out}


@app.post("/api/v2/journeys/start")
async def api_v2_journeys_start(request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    journeys: JourneyOrchestratorService = demo["journey_orchestrator"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    try:
        row = journeys.start_journey(
            tenant_id=ctx["tenant_id"],
            payload=body,
            actor=ctx["actor"],
            request_id=ctx["request_id"],
        )
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    _sync_task_linkages(
        task_id=str(body.get("task_id") or "").strip() or None,
        journey_id=row.get("id"),
        loan_account_id=row.get("loan_account_id"),
    )
    event_bus.publish(
        tenant_id=ctx["tenant_id"],
        event_type="journey.started",
        payload={"journey_id": row.get("id"), "customer_id": row.get("customer_id")},
        actor=ctx["actor"],
        request_id=ctx["request_id"],
    )
    return {"schema_version": "v2.1", "journey": row}


@app.post("/api/v2/journeys/{journey_id}/advance")
async def api_v2_journeys_advance(journey_id: str, request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    journeys: JourneyOrchestratorService = demo["journey_orchestrator"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    try:
        row = journeys.advance_journey(
            tenant_id=ctx["tenant_id"],
            journey_id=journey_id,
            payload=body,
            actor=ctx["actor"],
            request_id=ctx["request_id"],
        )
    except ValueError as exc:
        code = 404 if str(exc) == "journey_not_found" else 400
        return JSONResponse({"error": str(exc)}, status_code=code)
    event_bus.publish(
        tenant_id=ctx["tenant_id"],
        event_type="journey.advanced",
        payload={"journey_id": journey_id, "signal": body.get("signal"), "status": row.get("status")},
        actor=ctx["actor"],
        request_id=ctx["request_id"],
    )
    return {"schema_version": "v2.1", "journey": row}


@app.get("/api/v2/journeys/{journey_id}")
async def api_v2_journeys_get(journey_id: str, request: Request):
    deny = _require_permissions(request, (WORKBENCH_VIEW,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    journeys: JourneyOrchestratorService = demo["journey_orchestrator"]  # type: ignore[assignment]
    row = journeys.get_journey(tenant_id=ctx["tenant_id"], journey_id=journey_id)
    if not row:
        return JSONResponse({"error": "journey_not_found"}, status_code=404)
    return {"schema_version": "v2.1", "journey": row}


@app.post("/api/v2/channels/send")
async def api_v2_channels_send(request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    channel_hub: ChannelHubService = demo["channel_hub"]  # type: ignore[assignment]
    journeys: JourneyOrchestratorService = demo["journey_orchestrator"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    try:
        raw_body = await request.json()
    except Exception:
        raw_body = {}
    try:
        body_obj = V2ChannelSendRequest.model_validate(raw_body if isinstance(raw_body, dict) else {})
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid channel send payload",
            details={"errors": exc.errors()},
        )
    body = body_obj.model_dump(exclude_none=True)
    idem_key, body_hash, replay = _idempotency_preflight(
        request=request,
        audit=audit,
        route_key=f"v2:channels:send:{ctx['tenant_id']}",
        body=body,
    )
    if replay is not None:
        return replay
    try:
        msg = channel_hub.send_message(
            tenant_id=ctx["tenant_id"],
            payload=body,
            actor=ctx["actor"],
            request_id=ctx["request_id"],
        )
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    journey_id = str(msg.get("journey_id") or "").strip()
    if journey_id:
        with contextlib.suppress(Exception):
            journeys.advance_journey(
                tenant_id=ctx["tenant_id"],
                journey_id=journey_id,
                payload={
                    "step_type": "channel_send",
                    "direction": "outbound",
                    "channel": msg.get("channel"),
                    "status": "DONE",
                    "signal": "message_sent",
                    "message_id": msg.get("id"),
                },
                actor=ctx["actor"],
                request_id=ctx["request_id"],
            )
    event_bus.publish(
        tenant_id=ctx["tenant_id"],
        event_type="channel.outbound.queued",
        payload={"message_id": msg.get("id"), "channel": msg.get("channel")},
        actor=ctx["actor"],
        request_id=ctx["request_id"],
    )
    payload = {"schema_version": "v2.1", "message": msg}
    _idempotency_store(
        audit=audit,
        route_key=f"v2:channels:send:{ctx['tenant_id']}",
        idempotency_key=idem_key,
        body_hash=body_hash,
        response_payload=payload,
    )
    return payload


@app.post("/api/v2/channels/inbound/webhook/{provider}")
async def api_v2_channels_inbound(provider: str, request: Request):
    raw_body = await request.body()
    deny = _verify_webhook_signature(request, raw_body)
    if deny:
        return deny
    try:
        parsed = json.loads(raw_body.decode("utf-8") or "{}")
    except Exception:
        parsed = {}
    body = parsed if isinstance(parsed, dict) else {}
    raw_tenant_id = (
        request.headers.get("x-tenant-id")
        or request.query_params.get("tenant_id")
        or body.get("tenant_id")
        or "default"
    )
    demo = _get_demo_singletons()
    tenants: TenantService = demo["tenant_service"]  # type: ignore[assignment]
    channel_hub: ChannelHubService = demo["channel_hub"]  # type: ignore[assignment]
    journeys: JourneyOrchestratorService = demo["journey_orchestrator"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    external_event_id = _extract_webhook_event_id(body)
    if _is_prod() and not external_event_id:
        return _error_response(
            request,
            status_code=400,
            code="missing_external_event_id",
            message="Webhook event id is required",
        )
    if external_event_id:
        seen = audit.webhook_seen(
            provider=f"channels:{provider}",
            external_event_id=external_event_id,
            request_hash=hashlib.sha256(raw_body).hexdigest(),
        )
        if seen:
            return {"schema_version": "v2.1", "deduplicated": True, "provider_event_id": external_event_id}
    try:
        tenant_id = tenants.ensure_tenant(str(raw_tenant_id).strip().lower() or "default")
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    req_id = _request_id(request)
    try:
        msg = channel_hub.ingest_inbound(
            provider=provider,
            tenant_id=tenant_id,
            payload=body if isinstance(body, dict) else {},
            actor="webhook",
            request_id=req_id,
        )
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    journey_id = str(msg.get("journey_id") or "").strip()
    if journey_id:
        with contextlib.suppress(Exception):
            journeys.advance_journey(
                tenant_id=tenant_id,
                journey_id=journey_id,
                payload={
                    "step_type": "channel_inbound",
                    "direction": "inbound",
                    "channel": msg.get("channel"),
                    "status": "DONE",
                    "signal": "customer_reply",
                    "message_id": msg.get("id"),
                },
                actor="webhook",
                request_id=req_id,
            )
    event_bus.publish(
        tenant_id=tenant_id,
        event_type="channel.inbound.received",
        payload={"message_id": msg.get("id"), "provider": provider},
        actor="webhook",
        request_id=req_id,
    )
    return {"schema_version": "v2.1", "message": msg}


@app.get("/api/v2/channels/messages")
async def api_v2_channels_list(request: Request):
    deny = _require_permissions(request, (WORKBENCH_VIEW,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    page, page_size = _pagination(request)
    demo = _get_demo_singletons()
    channel_hub: ChannelHubService = demo["channel_hub"]  # type: ignore[assignment]
    out = channel_hub.list_messages(
        tenant_id=ctx["tenant_id"],
        journey_id=(request.query_params.get("journey_id") or "").strip() or None,
        customer_id=(request.query_params.get("customer_id") or "").strip() or None,
        channel=(request.query_params.get("channel") or "").strip() or None,
        direction=(request.query_params.get("direction") or "").strip() or None,
        page=page,
        page_size=page_size,
    )
    out["schema_version"] = "v2.1"
    return out


@app.post("/api/v2/payments/intents")
async def api_v2_payment_intents(request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    payments: PaymentOrchestrationService = demo["payment_orchestration"]  # type: ignore[assignment]
    journeys: JourneyOrchestratorService = demo["journey_orchestrator"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    if _is_prod() and "pay.example" in str(getattr(payments, "pay_base_url", "")):
        return _error_response(
            request,
            status_code=500,
            code="payment_provider_not_configured",
            message="Payment provider configuration is missing",
        )
    try:
        raw_body = await request.json()
    except Exception:
        raw_body = {}
    try:
        body_obj = V2PaymentIntentRequest.model_validate(raw_body if isinstance(raw_body, dict) else {})
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid payment intent payload",
            details={"errors": exc.errors()},
        )
    body = body_obj.model_dump(exclude_none=True)
    idem_key, body_hash, replay = _idempotency_preflight(
        request=request,
        audit=audit,
        route_key=f"v2:payments:intents:{ctx['tenant_id']}",
        body=body,
    )
    if replay is not None:
        return replay
    try:
        intent = payments.create_intent(
            tenant_id=ctx["tenant_id"],
            payload=body,
            actor=ctx["actor"],
            request_id=ctx["request_id"],
        )
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    journey_id = str(intent.get("journey_id") or "").strip()
    if journey_id:
        with contextlib.suppress(Exception):
            journeys.advance_journey(
                tenant_id=ctx["tenant_id"],
                journey_id=journey_id,
                payload={
                    "step_type": "payment_intent_created",
                    "channel": "digital",
                    "status": "DONE",
                    "signal": "payment_link_sent",
                    "payment_intent_id": intent.get("id"),
                },
                actor=ctx["actor"],
                request_id=ctx["request_id"],
            )
    event_bus.publish(
        tenant_id=ctx["tenant_id"],
        event_type="payment.intent.created",
        payload={"payment_intent_id": intent.get("id"), "amount": intent.get("amount")},
        actor=ctx["actor"],
        request_id=ctx["request_id"],
    )
    payload = {"schema_version": "v2.1", "payment_intent": intent}
    _idempotency_store(
        audit=audit,
        route_key=f"v2:payments:intents:{ctx['tenant_id']}",
        idempotency_key=idem_key,
        body_hash=body_hash,
        response_payload=payload,
    )
    return payload


@app.post("/api/v2/payments/webhooks/{rail}")
async def api_v2_payment_webhook(rail: str, request: Request):
    raw_body = await request.body()
    deny = _verify_webhook_signature(request, raw_body)
    if deny:
        return deny
    try:
        parsed = json.loads(raw_body.decode("utf-8") or "{}")
    except Exception:
        parsed = {}
    body = parsed if isinstance(parsed, dict) else {}
    raw_tenant_id = (
        request.headers.get("x-tenant-id")
        or request.query_params.get("tenant_id")
        or body.get("tenant_id")
        or "default"
    )
    demo = _get_demo_singletons()
    tenants: TenantService = demo["tenant_service"]  # type: ignore[assignment]
    payments: PaymentOrchestrationService = demo["payment_orchestration"]  # type: ignore[assignment]
    journeys: JourneyOrchestratorService = demo["journey_orchestrator"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    external_event_id = _extract_webhook_event_id(body)
    if _is_prod() and not external_event_id:
        return _error_response(
            request,
            status_code=400,
            code="missing_external_event_id",
            message="Webhook event id is required",
        )
    if external_event_id:
        seen = audit.webhook_seen(
            provider=f"payments:{rail}",
            external_event_id=external_event_id,
            request_hash=hashlib.sha256(raw_body).hexdigest(),
        )
        if seen:
            return {"schema_version": "v2.1", "deduplicated": True, "provider_event_id": external_event_id}
    try:
        tenant_id = tenants.ensure_tenant(str(raw_tenant_id).strip().lower() or "default")
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    req_id = _request_id(request)
    out = payments.ingest_webhook(
        rail=rail,
        tenant_id=tenant_id,
        payload=body if isinstance(body, dict) else {},
        actor="webhook",
        request_id=req_id,
    )
    intent_id = str(out.get("payment_intent_id") or "").strip()
    if intent_id and str(out.get("status") or "").upper() == "SUCCEEDED":
        intent = payments.get_intent(tenant_id=tenant_id, intent_id=intent_id)
        if intent and intent.get("journey_id"):
            with contextlib.suppress(Exception):
                journeys.advance_journey(
                    tenant_id=tenant_id,
                    journey_id=str(intent.get("journey_id")),
                    payload={
                        "step_type": "payment_webhook",
                        "channel": "digital",
                        "status": "DONE",
                        "signal": "payment_success",
                        "result": "paid",
                        "close": True,
                        "payment_intent_id": intent_id,
                    },
                    actor="webhook",
                    request_id=req_id,
                )
    event_bus.publish(
        tenant_id=tenant_id,
        event_type="payment.webhook.ingested",
        payload={"rail": rail, "payment_intent_id": out.get("payment_intent_id"), "status": out.get("status")},
        actor="webhook",
        request_id=req_id,
    )
    return {"schema_version": "v2.1", "event": out}


@app.post("/api/v2/settlements/offers")
async def api_v2_settlement_offers(request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    settlements: SettlementService = demo["settlement"]  # type: ignore[assignment]
    journeys: JourneyOrchestratorService = demo["journey_orchestrator"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    try:
        raw_body = await request.json()
    except Exception:
        raw_body = {}
    try:
        body_obj = V2SettlementOfferRequest.model_validate(raw_body if isinstance(raw_body, dict) else {})
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid settlement offer payload",
            details={"errors": exc.errors()},
        )
    body = body_obj.model_dump(exclude_none=True)
    idem_key, body_hash, replay = _idempotency_preflight(
        request=request,
        audit=audit,
        route_key=f"v2:settlements:offers:{ctx['tenant_id']}",
        body=body,
    )
    if replay is not None:
        return replay
    try:
        offer = settlements.create_offer(
            tenant_id=ctx["tenant_id"],
            payload=body,
            actor=ctx["actor"],
            request_id=ctx["request_id"],
        )
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    journey_id = str(offer.get("journey_id") or "").strip()
    if journey_id:
        with contextlib.suppress(Exception):
            journeys.advance_journey(
                tenant_id=ctx["tenant_id"],
                journey_id=journey_id,
                payload={
                    "step_type": "settlement_offer_created",
                    "status": "DONE",
                    "signal": "settlement_offer",
                    "offer_id": offer.get("id"),
                },
                actor=ctx["actor"],
                request_id=ctx["request_id"],
            )
    event_bus.publish(
        tenant_id=ctx["tenant_id"],
        event_type="settlement.offer.created",
        payload={"offer_id": offer.get("id"), "customer_id": offer.get("customer_id")},
        actor=ctx["actor"],
        request_id=ctx["request_id"],
    )
    payload = {"schema_version": "v2.1", "offer": offer}
    _idempotency_store(
        audit=audit,
        route_key=f"v2:settlements:offers:{ctx['tenant_id']}",
        idempotency_key=idem_key,
        body_hash=body_hash,
        response_payload=payload,
    )
    return payload


@app.post("/api/v2/settlements/{offer_id}/accept")
async def api_v2_settlement_accept(offer_id: str, request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    settlements: SettlementService = demo["settlement"]  # type: ignore[assignment]
    journeys: JourneyOrchestratorService = demo["journey_orchestrator"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    try:
        raw_body = await request.json()
    except Exception:
        raw_body = {}
    try:
        body_obj = V2SettlementAcceptRequest.model_validate(raw_body if isinstance(raw_body, dict) else {})
    except ValidationError as exc:
        return _error_response(
            request,
            status_code=400,
            code="invalid_payload",
            message="Invalid settlement accept payload",
            details={"errors": exc.errors()},
        )
    body = body_obj.model_dump(exclude_none=True)
    idem_key, body_hash, replay = _idempotency_preflight(
        request=request,
        audit=audit,
        route_key=f"v2:settlements:accept:{ctx['tenant_id']}:{offer_id}",
        body=body,
    )
    if replay is not None:
        return replay
    try:
        out = settlements.accept_offer(
            tenant_id=ctx["tenant_id"],
            offer_id=offer_id,
            payload=body if isinstance(body, dict) else {},
            actor=ctx["actor"],
            request_id=ctx["request_id"],
        )
    except ValueError as exc:
        code = 404 if str(exc) == "offer_not_found" else 400
        return JSONResponse({"error": str(exc)}, status_code=code)
    journey_id = str(out.get("journey_id") or "").strip()
    if journey_id:
        with contextlib.suppress(Exception):
            journeys.advance_journey(
                tenant_id=ctx["tenant_id"],
                journey_id=journey_id,
                payload={
                    "step_type": "settlement_accepted",
                    "status": "DONE",
                    "signal": "settlement_accept",
                    "result": "settled",
                    "close": True,
                    "offer_id": offer_id,
                },
                actor=ctx["actor"],
                request_id=ctx["request_id"],
            )
    event_bus.publish(
        tenant_id=ctx["tenant_id"],
        event_type="settlement.offer.accepted",
        payload={"offer_id": offer_id, "status": out.get("status")},
        actor=ctx["actor"],
        request_id=ctx["request_id"],
    )
    payload = {"schema_version": "v2.1", "offer": out}
    _idempotency_store(
        audit=audit,
        route_key=f"v2:settlements:accept:{ctx['tenant_id']}:{offer_id}",
        idempotency_key=idem_key,
        body_hash=body_hash,
        response_payload=payload,
    )
    return payload


@app.post("/api/v2/recovery/nba/decide")
async def api_v2_recovery_nba_decide(request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    brain: RecoveryBrainService = demo["recovery_brain"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    try:
        decision = brain.decide_nba(
            tenant_id=ctx["tenant_id"],
            payload=body if isinstance(body, dict) else {},
            actor=ctx["actor"],
            request_id=ctx["request_id"],
        )
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    _sync_task_linkages(
        task_id=str(body.get("task_id") or "").strip() or None,
        journey_id=str(decision.get("journey_id") or "").strip() or None,
        loan_account_id=str(decision.get("loan_account_id") or "").strip() or None,
        nba_decision_id=str(decision.get("id") or "").strip() or None,
        approval_state="PENDING" if decision.get("approval_required") else "NOT_REQUIRED",
    )
    event_bus.publish(
        tenant_id=ctx["tenant_id"],
        event_type="recovery.nba.decided",
        payload={"decision_id": decision.get("id"), "action_type": decision.get("action_type")},
        actor=ctx["actor"],
        request_id=ctx["request_id"],
    )
    return {"schema_version": "v2.1", "decision": decision}


@app.post("/api/v2/recovery/experiments")
async def api_v2_recovery_experiments(request: Request):
    deny = _require_permissions(request, (WORKBENCH_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    experiments: ExperimentService = demo["experiment_service"]  # type: ignore[assignment]
    event_bus: EventBusService = demo["event_bus"]  # type: ignore[assignment]
    try:
        body = await request.json()
    except Exception:
        body = {}
    try:
        experiment = experiments.create_experiment(
            tenant_id=ctx["tenant_id"],
            payload=body if isinstance(body, dict) else {},
            actor=ctx["actor"],
            request_id=ctx["request_id"],
        )
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    assignment = None
    unit_id = str(body.get("unit_id") or "").strip()
    if unit_id:
        with contextlib.suppress(Exception):
            assignment = experiments.assign(
                tenant_id=ctx["tenant_id"],
                experiment_id=str(experiment.get("id")),
                unit_id=unit_id,
                actor=ctx["actor"],
                request_id=ctx["request_id"],
                source="api",
            )
    event_bus.publish(
        tenant_id=ctx["tenant_id"],
        event_type="recovery.experiment.created",
        payload={"experiment_id": experiment.get("id"), "name": experiment.get("name")},
        actor=ctx["actor"],
        request_id=ctx["request_id"],
    )
    return {"schema_version": "v2.1", "experiment": experiment, "assignment": assignment}


@app.get("/api/v2/recovery/uplift")
async def api_v2_recovery_uplift(request: Request):
    deny = _require_permissions(request, (VIEW_DASHBOARD,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    demo = _get_demo_singletons()
    brain: RecoveryBrainService = demo["recovery_brain"]  # type: ignore[assignment]
    experiment_id = (request.query_params.get("experiment_id") or "").strip() or None
    out = brain.get_uplift_snapshot(tenant_id=ctx["tenant_id"], experiment_id=experiment_id)
    out["schema_version"] = "v2.1"
    return out


@app.get("/api/v2/approvals/queue")
async def api_v2_approvals_queue(request: Request):
    deny = _require_permissions(request, (ALERTS_VIEW,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    page, page_size = _pagination(request)
    status = (request.query_params.get("status") or "").strip() or None
    demo = _get_demo_singletons()
    approvals: ApprovalService = demo["approval_service"]  # type: ignore[assignment]
    out = approvals.list_queue(
        tenant_id=ctx["tenant_id"],
        status=status,
        page=page,
        page_size=page_size,
    )
    out["schema_version"] = "v2.1"
    return out


@app.post("/api/v2/approvals/{approval_id}/approve")
async def api_v2_approvals_approve(approval_id: str, request: Request):
    deny = _require_permissions(request, (ALERTS_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    try:
        body = await request.json()
    except Exception:
        body = {}
    demo = _get_demo_singletons()
    approvals: ApprovalService = demo["approval_service"]  # type: ignore[assignment]
    try:
        out = approvals.approve(
            tenant_id=ctx["tenant_id"],
            approval_id=approval_id,
            actor=ctx["actor"],
            request_id=ctx["request_id"],
            note=str(body.get("note") or "").strip() or None,
            payload=body if isinstance(body, dict) else None,
        )
    except ValueError as exc:
        code = 404 if str(exc) == "approval_not_found" else 400
        return JSONResponse({"error": str(exc)}, status_code=code)
    return {"schema_version": "v2.1", "approval": out}


@app.post("/api/v2/approvals/{approval_id}/reject")
async def api_v2_approvals_reject(approval_id: str, request: Request):
    deny = _require_permissions(request, (ALERTS_MUTATE,))
    if deny:
        return deny
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    try:
        body = await request.json()
    except Exception:
        body = {}
    demo = _get_demo_singletons()
    approvals: ApprovalService = demo["approval_service"]  # type: ignore[assignment]
    try:
        out = approvals.reject(
            tenant_id=ctx["tenant_id"],
            approval_id=approval_id,
            actor=ctx["actor"],
            request_id=ctx["request_id"],
            note=str(body.get("note") or "").strip() or None,
            payload=body if isinstance(body, dict) else None,
        )
    except ValueError as exc:
        code = 404 if str(exc) == "approval_not_found" else 400
        return JSONResponse({"error": str(exc)}, status_code=code)
    return {"schema_version": "v2.1", "approval": out}


class WebSocketSender:
    _CONTROL_MAX = 500
    _AUDIO_MAX = 300
    _CONTROL_BURST = 3

    def __init__(self, websocket: WebSocket) -> None:
        self.websocket = websocket
        self._control_queue: asyncio.Queue[str] = asyncio.Queue(maxsize=self._CONTROL_MAX)
        self._audio_queue: asyncio.Queue[bytes] = asyncio.Queue(maxsize=self._AUDIO_MAX)
        self._control_burst = 0
        self._task: Optional[asyncio.Task] = None
        self._closed = asyncio.Event()

    def start(self) -> None:
        if not self._task:
            self._task = asyncio.create_task(self._send_loop(), name="ws_send_loop")

    async def close(self) -> None:
        self._closed.set()
        if self._task and not self._task.done():
            self._task.cancel()
            with contextlib.suppress(asyncio.CancelledError):
                await self._task

    async def send_json(self, payload: dict) -> None:
        if self.websocket.application_state != WebSocketState.CONNECTED:
            return
        msg = json.dumps(payload)
        if self._control_queue.full():
            with contextlib.suppress(asyncio.QueueEmpty):
                _ = self._control_queue.get_nowait()
        with contextlib.suppress(asyncio.QueueFull):
            self._control_queue.put_nowait(msg)

    async def send_bytes(self, data: bytes) -> None:
        if self.websocket.application_state != WebSocketState.CONNECTED:
            return
        if self._audio_queue.full():
            with contextlib.suppress(asyncio.QueueEmpty):
                _ = self._audio_queue.get_nowait()
        with contextlib.suppress(asyncio.QueueFull):
            self._audio_queue.put_nowait(data)

    def _pick_ready_item(self) -> Optional[tuple[str, object]]:
        if self._control_burst < self._CONTROL_BURST:
            with contextlib.suppress(asyncio.QueueEmpty):
                self._control_burst += 1
                return ("text", self._control_queue.get_nowait())
        with contextlib.suppress(asyncio.QueueEmpty):
            self._control_burst = 0
            return ("bytes", self._audio_queue.get_nowait())
        with contextlib.suppress(asyncio.QueueEmpty):
            self._control_burst += 1
            return ("text", self._control_queue.get_nowait())
        with contextlib.suppress(asyncio.QueueEmpty):
            self._control_burst = 0
            return ("bytes", self._audio_queue.get_nowait())
        return None

    async def _send_loop(self) -> None:
        try:
            while not self._closed.is_set():
                if self.websocket.application_state != WebSocketState.CONNECTED:
                    break
                item = self._pick_ready_item()
                if item is None:
                    await asyncio.sleep(0.01)
                    continue
                kind, payload = item
                try:
                    if kind == "text":
                        await self.websocket.send_text(str(payload))
                    else:
                        await self.websocket.send_bytes(bytes(payload))
                except (RuntimeError, WebSocketDisconnect):
                    break
        finally:
            self._closed.set()


_MULAW_BIAS = 0x84
_MULAW_CLIP = 32635


def _mulaw_to_linear16(value: int) -> int:
    u = (~value) & 0xFF
    sign = u & 0x80
    exponent = (u >> 4) & 0x07
    mantissa = u & 0x0F
    sample = ((mantissa << 3) + _MULAW_BIAS) << exponent
    sample -= _MULAW_BIAS
    if sign:
        sample = -sample
    return int(max(-32768, min(32767, sample)))


def _linear16_to_mulaw(sample: int) -> int:
    v = int(sample)
    sign = 0x80 if v < 0 else 0x00
    if v < 0:
        v = -v
    if v > _MULAW_CLIP:
        v = _MULAW_CLIP
    v += _MULAW_BIAS
    exponent = 7
    exp_mask = 0x4000
    while exponent > 0 and (v & exp_mask) == 0:
        exponent -= 1
        exp_mask >>= 1
    mantissa = (v >> (exponent + 3)) & 0x0F
    return (~(sign | (exponent << 4) | mantissa)) & 0xFF


def _pcm16le_to_samples(pcm: bytes) -> List[int]:
    if not pcm:
        return []
    usable = len(pcm) - (len(pcm) % 2)
    if usable <= 0:
        return []
    return list(struct.unpack(f"<{usable // 2}h", pcm[:usable]))


def _samples_to_pcm16le(samples: List[int]) -> bytes:
    if not samples:
        return b""
    clipped = [max(-32768, min(32767, int(s))) for s in samples]
    return struct.pack(f"<{len(clipped)}h", *clipped)


def _mulaw_bytes_to_pcm16le(data: bytes) -> bytes:
    return _samples_to_pcm16le([_mulaw_to_linear16(b) for b in data])


def _pcm16le_to_mulaw_bytes(pcm: bytes) -> bytes:
    return bytes(_linear16_to_mulaw(s) for s in _pcm16le_to_samples(pcm))


def _upsample_pcm16_mono_8k_to_16k(pcm8: bytes) -> bytes:
    src = _pcm16le_to_samples(pcm8)
    if not src:
        return b""
    out: List[int] = []
    for i, cur in enumerate(src):
        out.append(cur)
        nxt = src[i + 1] if i + 1 < len(src) else cur
        out.append((cur + nxt) // 2)
    return _samples_to_pcm16le(out)


def _downsample_pcm16_mono_16k_to_8k(pcm16: bytes) -> bytes:
    src = _pcm16le_to_samples(pcm16)
    if not src:
        return b""
    out = [src[i] for i in range(0, len(src), 2)]
    return _samples_to_pcm16le(out)


def _to_ws_url(base_url: str, path: str) -> str:
    raw = (base_url or "").strip().rstrip("/")
    if raw.startswith("https://"):
        return f"wss://{raw[len('https://'):]}/{path.lstrip('/')}"
    if raw.startswith("http://"):
        return f"ws://{raw[len('http://'):]}/{path.lstrip('/')}"
    if raw.startswith("wss://") or raw.startswith("ws://"):
        return f"{raw}/{path.lstrip('/')}"
    return ""


def _telephony_stream_ws_url(request: Optional[Request] = None) -> str:
    explicit = (get_env("TELEPHONY_STREAM_WSS_URL", "") or "").strip()
    if explicit:
        return explicit
    base = (get_env("TELEPHONY_PUBLIC_BASE_URL", "") or "").strip()
    if not base and request is not None:
        scheme = (request.headers.get("x-forwarded-proto") or request.url.scheme or "").strip()
        host = (request.headers.get("x-forwarded-host") or request.headers.get("host") or "").strip()
        if scheme and host:
            base = f"{scheme}://{host}"
    return _to_ws_url(base, "/ws/twilio-media")


class TwilioMediaSender:
    def __init__(self, websocket: WebSocket) -> None:
        self.websocket = websocket
        self.stream_sid: Optional[str] = None
        self._lock = asyncio.Lock()
        self._pcm16_out_buffer = bytearray()

    async def set_stream_sid(self, stream_sid: str) -> None:
        self.stream_sid = (stream_sid or "").strip() or None

    async def send_pcm16(self, pcm16: bytes) -> None:
        if not pcm16:
            return
        if not self.stream_sid:
            return
        self._pcm16_out_buffer.extend(pcm16)
        frames: List[bytes] = []
        frame_bytes_16k = 640  # 20ms @ 16kHz mono PCM16
        while len(self._pcm16_out_buffer) >= frame_bytes_16k:
            frame16 = bytes(self._pcm16_out_buffer[:frame_bytes_16k])
            del self._pcm16_out_buffer[:frame_bytes_16k]
            frames.append(frame16)
        for frame16 in frames:
            frame8 = _downsample_pcm16_mono_16k_to_8k(frame16)
            ulaw = _pcm16le_to_mulaw_bytes(frame8)
            payload = base64.b64encode(ulaw).decode("ascii")
            await self._send_media_payload(payload)

    async def flush(self) -> None:
        if not self.stream_sid:
            self._pcm16_out_buffer.clear()
            return
        if not self._pcm16_out_buffer:
            return
        frame16 = bytes(self._pcm16_out_buffer)
        self._pcm16_out_buffer.clear()
        frame8 = _downsample_pcm16_mono_16k_to_8k(frame16)
        ulaw = _pcm16le_to_mulaw_bytes(frame8)
        payload = base64.b64encode(ulaw).decode("ascii")
        await self._send_media_payload(payload)

    async def _send_media_payload(self, payload: str) -> None:
        if self.websocket.application_state != WebSocketState.CONNECTED:
            return
        if not self.stream_sid:
            return
        msg = {
            "event": "media",
            "streamSid": self.stream_sid,
            "media": {"payload": payload},
        }
        async with self._lock:
            await self.websocket.send_text(json.dumps(msg))


@app.websocket("/ws/twilio-media")
async def twilio_media_socket(websocket: WebSocket) -> None:
    await websocket.accept()

    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    sink: ExcelOutcomeSink = demo["sink"]  # type: ignore[assignment]
    router: ActionRouter = demo["router"]  # type: ignore[assignment]
    knowledge: SQLiteFTSKnowledgeStore = demo["knowledge"]  # type: ignore[assignment]
    strategy: StrategyEngine = demo["strategy"]  # type: ignore[assignment]
    compliance: ComplianceEngine = demo["compliance"]  # type: ignore[assignment]
    followups: FollowupService = demo["followups"]  # type: ignore[assignment]
    crm: CRMAdapter = demo["crm"]  # type: ignore[assignment]
    sessions: Dict[str, dict] = demo["sessions"]  # type: ignore[assignment]
    session_objs: Dict[str, WebCallSession] = demo["session_objs"]  # type: ignore[assignment]

    load_dotenv(override=True)
    stt_ws_url = get_env(
        "SAARIKA_WS_URL",
        get_env("SARVAM_STT_WS_URL", "wss://api.sarvam.ai/speech-to-text/ws"),
    )
    llm_base_url = get_env("SARVAM_BASE_URL", "https://api.sarvam.ai")
    tts_ws_url = get_env(
        "BULBUL_WS_URL",
        get_env("SARVAM_TTS_WS_URL", "wss://api.sarvam.ai/text-to-speech/ws"),
    )

    api_key = get_env("SARVAM_API_KEY")
    stt_api_key = get_env("SAARIKA_API_KEY", api_key)
    llm_api_key = get_env("SARVAM_LLM_API_KEY", api_key)
    tts_api_key = get_env("BULBUL_API_KEY", api_key)
    if not stt_api_key:
        await websocket.close(code=4500)
        return

    llm_model = get_env("SARVAM_CHAT_MODEL", get_env("SARVAM_LLM_MODEL", "sarvam-m"))
    tts_voice = get_env("BULBUL_VOICE")
    tts_speaker = get_env("SARVAM_TTS_SPEAKER", "shubh")
    stt_model = get_env("SARVAM_STT_MODEL", "saarika:v2.5")
    tts_model = get_env("SARVAM_TTS_MODEL", "bulbul:v3")
    stt_language = get_env("SARVAM_STT_LANGUAGE", "en-IN")
    stt_language_key = get_env("SARVAM_STT_LANGUAGE_PARAM", "language_code")
    vad_threshold = float(get_env("VAD_RMS_THRESHOLD", "500"))
    vad_silence_ms = int(get_env("VAD_SILENCE_MS", "600"))
    stt_flush_interval_ms = int(get_env("STT_FLUSH_INTERVAL_MS", "250"))
    stt_vad_signals = get_env_bool("STT_VAD_SIGNALS", True)
    stt_high_vad = get_env_bool("STT_HIGH_VAD_SENSITIVITY", True)
    use_sdk = get_env_bool("USE_SARVAMAI_SDK", True)
    use_sdk_stt = get_env_bool("USE_SARVAMAI_SDK_STT", use_sdk)
    use_sdk_tts = get_env_bool("USE_SARVAMAI_SDK_TTS", False)
    use_sdk_llm = get_env_bool("USE_SARVAMAI_SDK_LLM", False)
    stt_flush_signal_env = get_env("STT_FLUSH_SIGNAL")
    if stt_flush_signal_env is None or stt_flush_signal_env == "":
        stt_flush_signal = use_sdk_stt
    else:
        stt_flush_signal = get_env_bool("STT_FLUSH_SIGNAL", False)
    stt_audio_payload_format = get_env("STT_AUDIO_PAYLOAD_FORMAT", "flat")
    stt_input_audio_codec = get_env("STT_INPUT_AUDIO_CODEC", "pcm_s16le")
    stt_audio_encoding = get_env("STT_AUDIO_ENCODING")
    stt_log_raw = get_env_bool("LOG_STT_RAW", False)
    log_partials = get_env_bool("LOG_PARTIAL_TRANSCRIPTS", True)
    log_tokens = get_env_bool("LOG_LLM_TOKENS", False)
    log_audio_chunks = get_env_bool("LOG_TTS_CHUNKS", False)
    greeting_text = get_env("GREETING_TEXT")
    enable_greeting = get_env_bool("ENABLE_GREETING", True)
    max_history_turns = int(get_env("MAX_HISTORY_TURNS", "10"))
    session_store_path = get_env("SESSION_STORE_PATH")
    dynamic_stt_language = get_env_bool("STT_DYNAMIC_LANGUAGE", True)
    preview_partials = get_env_bool("PREVIEW_PARTIALS", True)
    preview_after_ms = int(get_env("PREVIEW_AFTER_MS", "300"))
    llm_timeout_s = float(get_env("LLM_STREAM_TIMEOUT_S", "20"))
    tts_timeout_s = float(get_env("TTS_STREAM_TIMEOUT_S", "20"))
    stt_connect_timeout_s = float(get_env("STT_CONNECT_TIMEOUT_S", "10"))
    tts_stream_chunk_chars = int(get_env("TTS_STREAM_CHUNK_CHARS", "50"))
    tts_stream_flush_punct = get_env_bool("TTS_STREAM_FLUSH_PUNCT", True)
    tts_min_buffer_size = int(get_env("TTS_MIN_BUFFER_SIZE", "30"))
    tts_max_chunk_length = int(get_env("TTS_MAX_CHUNK_LENGTH", "120"))
    tts_output_audio_codec = get_env("TTS_OUTPUT_AUDIO_CODEC", "linear16")
    tts_output_audio_bitrate = get_env("TTS_OUTPUT_AUDIO_BITRATE")
    post_speech_pause_ms = int(get_env("POST_SPEECH_PAUSE_MS", "800"))  # 800ms suits Indian speech cadence

    stt = SaarikaSTTService(
        ws_url=stt_ws_url,
        api_key=stt_api_key,
        model=stt_model,
        language=stt_language,
        language_param_key=stt_language_key,
        vad_signals=stt_vad_signals,
        high_vad_sensitivity=stt_high_vad,
        flush_signal=stt_flush_signal,
        audio_payload_format=stt_audio_payload_format,
        input_audio_codec=stt_input_audio_codec,
        audio_encoding=stt_audio_encoding,
        log_raw_messages=stt_log_raw,
        use_sdk=use_sdk_stt,
    )
    llm = SarvamLLMService(base_url=llm_base_url, api_key=llm_api_key, model=llm_model, use_sdk=use_sdk_llm)

    async def _noop_send_event(_: dict) -> None:
        return

    twilio_sender = TwilioMediaSender(websocket)
    session = WebCallSession(
        stt_service=stt,
        llm_service=llm,
        tts_ws_url=tts_ws_url,
        send_event=_noop_send_event,
        send_audio=twilio_sender.send_pcm16,
        tts_api_key=tts_api_key,
        tts_voice=tts_voice,
        tts_model=tts_model,
        tts_speaker=tts_speaker,
        greeting_text=greeting_text if enable_greeting else None,
        max_history_turns=max_history_turns,
        session_store_path=session_store_path,
        dynamic_stt_language=dynamic_stt_language,
        preview_partials=preview_partials,
        preview_after_ms=preview_after_ms,
        llm_timeout_s=llm_timeout_s,
        tts_timeout_s=tts_timeout_s,
        stt_connect_timeout_s=stt_connect_timeout_s,
        vad_rms_threshold=vad_threshold,
        vad_silence_ms=vad_silence_ms,
        stt_flush_interval_ms=stt_flush_interval_ms,
        log_partial_transcripts=log_partials,
        log_tokens=log_tokens,
        log_audio_chunks=log_audio_chunks,
        use_sdk=use_sdk_tts,
        tts_stream_chunk_chars=tts_stream_chunk_chars,
        tts_stream_flush_punct=tts_stream_flush_punct,
        tts_min_buffer_size=tts_min_buffer_size,
        tts_max_chunk_length=tts_max_chunk_length,
        tts_output_audio_codec=tts_output_audio_codec,
        tts_output_audio_bitrate=tts_output_audio_bitrate,
        post_speech_pause_ms=post_speech_pause_ms,
        audit_store=audit,
        excel_sink=sink,
        action_router=router,
        knowledge_store=knowledge,
        session_registry=sessions,
        strategy_engine=strategy,
        compliance_engine=compliance,
        followup_service=followups,
        crm_adapter=crm,
    )
    session_objs[session._session_id] = session

    inbound_pcm16_buffer = bytearray()
    stream_sid = ""
    session_started = False
    try:
        while True:
            packet = await websocket.receive()
            if packet.get("type") == "websocket.disconnect":
                break
            text = packet.get("text")
            if not text:
                continue
            try:
                data = json.loads(text)
            except Exception:
                continue
            event_type = str(data.get("event") or "").strip().lower()
            if event_type == "start":
                start = data.get("start") if isinstance(data.get("start"), dict) else {}
                stream_sid = str(start.get("streamSid") or "").strip()
                await twilio_sender.set_stream_sid(stream_sid)
                params = start.get("customParameters") if isinstance(start.get("customParameters"), dict) else {}
                context: Dict[str, Any] = {
                    "customer_id": params.get("customer_id"),
                    "customer_name": params.get("customer_name"),
                    "phone": params.get("phone"),
                    "overdue_amount": params.get("amount_due"),
                    "campaign_id": params.get("campaign_id"),
                    "language_preference": params.get("language") or "hi-IN",
                    "tts_speaker": params.get("tts_speaker"),
                }
                context = {k: v for k, v in context.items() if v not in (None, "")}
                if session_started:
                    continue
                preferred_lang = context.get("language_preference")
                if preferred_lang:
                    session.stt.language = session._resolve_stt_connect_language(str(preferred_lang)) or str(preferred_lang)
                if context:
                    with contextlib.suppress(Exception):
                        await session.set_context(context)
                try:
                    await session.start()
                    session_started = True
                except Exception as exc:
                    log_event(logger, "session_start_failed", error=str(exc), stream_sid=stream_sid or None)
                    await websocket.close(code=4500)
                    return
                if hasattr(session, "start_greeting"):
                    with contextlib.suppress(Exception):
                        await session.start_greeting()
                continue
            if event_type == "media":
                if not session_started:
                    continue
                media = data.get("media") if isinstance(data.get("media"), dict) else {}
                payload_b64 = str(media.get("payload") or "").strip()
                if not payload_b64:
                    continue
                try:
                    ulaw = base64.b64decode(payload_b64)
                except Exception:
                    continue
                pcm8 = _mulaw_bytes_to_pcm16le(ulaw)
                pcm16 = _upsample_pcm16_mono_8k_to_16k(pcm8)
                if not pcm16:
                    continue
                inbound_pcm16_buffer.extend(pcm16)
                frame_bytes_16k = 640
                while len(inbound_pcm16_buffer) >= frame_bytes_16k:
                    frame = bytes(inbound_pcm16_buffer[:frame_bytes_16k])
                    del inbound_pcm16_buffer[:frame_bytes_16k]
                    await session.handle_audio(frame)
                continue
            if event_type in {"stop", "closed"}:
                break
    except WebSocketDisconnect:
        pass
    except Exception as exc:
        log_event(logger, "twilio_media_socket_error", error=str(exc), stream_sid=stream_sid or None)
    finally:
        with contextlib.suppress(Exception):
            await twilio_sender.flush()
        with contextlib.suppress(Exception):
            await session.stop()
        with contextlib.suppress(Exception):
            session_objs.pop(session._session_id, None)
        with contextlib.suppress(Exception):
            await websocket.close()


@app.websocket("/ws/voice")
async def voice_socket(websocket: WebSocket) -> None:
    ws_user_id = ""
    ws_username = ""
    ws_role = ""
    ws_tenant_id = ""
    ws_session_id = ""
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    ws_token = (websocket.query_params.get("token") or websocket.query_params.get("ws_token") or "").strip()
    if not ws_token:
        await websocket.close(code=4401)
        return
    payload = decode_token(ws_token, _jwt_secret(), expected_token_type="ws", expected_aud="voice_ws")
    if not payload:
        await websocket.close(code=4401)
        return
    ws_user_id = str(payload.get("sub") or "").strip()
    ws_username = str(payload.get("username") or "").strip()
    ws_role = str(payload.get("role") or "").strip().upper()
    ws_tenant_id = str(payload.get("tenant_id") or "").strip().lower()
    ws_session_id = str(payload.get("sid") or "").strip()
    if not ws_user_id or not ws_tenant_id or not ws_role:
        await websocket.close(code=4401)
        return
    user_row = audit.get_user_by_id(ws_user_id)
    if not user_row or int(user_row.get("is_active") or 0) != 1:
        await websocket.close(code=4403)
        return
    if not audit.user_has_tenant(user_id=ws_user_id, tenant_id=ws_tenant_id):
        await websocket.close(code=4403)
        return
    if ws_session_id:
        sess = audit.get_user_session(session_id=ws_session_id)
        if not sess or sess.get("revoked_at") or float(sess.get("expires_at") or 0) < time.time():
            await websocket.close(code=4401)
            return
        if str(sess.get("user_id") or "") != ws_user_id:
            await websocket.close(code=4401)
            return
    live = int(_ws_user_connections.get(ws_user_id, 0))
    if live >= max(1, WS_CONNECTION_LIMIT_PER_USER):
        await websocket.close(code=4429)
        return
    _ws_user_connections[ws_user_id] = live + 1
    await websocket.accept()
    sender = WebSocketSender(websocket)
    sender.start()
    sink: ExcelOutcomeSink = demo["sink"]  # type: ignore[assignment]
    source: ExcelCustomerSource = demo["source"]  # type: ignore[assignment]
    router: ActionRouter = demo["router"]  # type: ignore[assignment]
    knowledge: SQLiteFTSKnowledgeStore = demo["knowledge"]  # type: ignore[assignment]
    strategy: StrategyEngine = demo["strategy"]  # type: ignore[assignment]
    compliance: ComplianceEngine = demo["compliance"]  # type: ignore[assignment]
    followups: FollowupService = demo["followups"]  # type: ignore[assignment]
    crm: CRMAdapter = demo["crm"]  # type: ignore[assignment]
    workbench: WorkbenchService = demo["workbench"]  # type: ignore[assignment]
    sessions: Dict[str, dict] = demo["sessions"]  # type: ignore[assignment]
    session_objs: Dict[str, WebCallSession] = demo["session_objs"]  # type: ignore[assignment]

    client = websocket.client
    log_event(
        logger,
        "ws_open",
        client=f"{client.host}:{client.port}" if client else None,
    )

    def _extract_context_from_payload(data: dict) -> dict:
        # Accept either {type: 'context', context: {...}} OR flat keys on the payload.
        ctx = data.get("context") if isinstance(data.get("context"), dict) else {}
        if not isinstance(ctx, dict):
            ctx = {}

        # Merge flat keys (excluding type/context) if caller sends them at top-level.
        for k, v in (data or {}).items():
            if k in {"type", "context"}:
                continue
            if v is None:
                continue
            ctx.setdefault(k, v)

        # Normalize common aliases
        # amount / overdue
        if "overdue_amount" not in ctx:
            if "amount" in ctx:
                ctx["overdue_amount"] = ctx.get("amount")
            elif "overdue" in ctx:
                ctx["overdue_amount"] = ctx.get("overdue")

        # due date
        if "due_date" not in ctx:
            if "due" in ctx:
                ctx["due_date"] = ctx.get("due")
            elif "overdue_date" in ctx:
                ctx["due_date"] = ctx.get("overdue_date")

        # customer name
        if "customer_name" not in ctx and "name" in ctx:
            ctx["customer_name"] = ctx.get("name")

        # Normalize empty strings
        cleaned = {}
        for k, v in ctx.items():
            if isinstance(v, str):
                vv = v.strip()
                if vv:
                    cleaned[k] = vv
            else:
                cleaned[k] = v
        return cleaned

    load_dotenv(override=True)
    stt_ws_url = get_env(
        "SAARIKA_WS_URL",
        get_env("SARVAM_STT_WS_URL", "wss://api.sarvam.ai/speech-to-text/ws"),
    )
    llm_base_url = get_env("SARVAM_BASE_URL", "https://api.sarvam.ai")
    tts_ws_url = get_env(
        "BULBUL_WS_URL",
        get_env("SARVAM_TTS_WS_URL", "wss://api.sarvam.ai/text-to-speech/ws"),
    )

    api_key = get_env("SARVAM_API_KEY")
    stt_api_key = get_env("SAARIKA_API_KEY", api_key)
    llm_api_key = get_env("SARVAM_LLM_API_KEY", api_key)
    tts_api_key = get_env("BULBUL_API_KEY", api_key)
    if not stt_api_key:
        log_event(logger, "config_error", error="Missing SAARIKA_API_KEY for STT.")
        await sender.send_json({"type": "error", "message": "Missing SAARIKA_API_KEY for STT."})
        await websocket.close()
        return

    llm_model = get_env("SARVAM_CHAT_MODEL", get_env("SARVAM_LLM_MODEL", "sarvam-m"))
    tts_voice = get_env("BULBUL_VOICE")
    tts_speaker = get_env("SARVAM_TTS_SPEAKER", "shubh")
    stt_model = get_env("SARVAM_STT_MODEL", "saarika:v2.5")
    tts_model = get_env("SARVAM_TTS_MODEL", "bulbul:v3")
    stt_language = get_env("SARVAM_STT_LANGUAGE", "en-IN")
    stt_language_key = get_env("SARVAM_STT_LANGUAGE_PARAM", "language_code")

    vad_threshold = float(get_env("VAD_RMS_THRESHOLD", "500"))
    vad_silence_ms = int(get_env("VAD_SILENCE_MS", "600"))
    stt_flush_interval_ms = int(get_env("STT_FLUSH_INTERVAL_MS", "250"))
    stt_vad_signals = get_env_bool("STT_VAD_SIGNALS", True)
    stt_high_vad = get_env_bool("STT_HIGH_VAD_SENSITIVITY", True)
    use_sdk = get_env_bool("USE_SARVAMAI_SDK", True)
    use_sdk_stt = get_env_bool("USE_SARVAMAI_SDK_STT", use_sdk)
    # SDK TTS streams MP3; browser/player expects PCM16. Default to raw WS unless explicitly enabled.
    use_sdk_tts = get_env_bool("USE_SARVAMAI_SDK_TTS", False)
    # LLM streaming via SDK doesn't handle event-stream correctly; default to HTTP SSE unless overridden.
    use_sdk_llm = get_env_bool("USE_SARVAMAI_SDK_LLM", False)
    stt_flush_signal_env = get_env("STT_FLUSH_SIGNAL")
    if stt_flush_signal_env is None or stt_flush_signal_env == "":
        stt_flush_signal = use_sdk_stt
    else:
        stt_flush_signal = get_env_bool("STT_FLUSH_SIGNAL", False)
    stt_audio_payload_format = get_env("STT_AUDIO_PAYLOAD_FORMAT", "flat")
    stt_input_audio_codec = get_env("STT_INPUT_AUDIO_CODEC", "pcm_s16le")
    stt_audio_encoding = get_env("STT_AUDIO_ENCODING")
    stt_log_raw = get_env_bool("LOG_STT_RAW", False)
    log_partials = get_env_bool("LOG_PARTIAL_TRANSCRIPTS", True)
    log_tokens = get_env_bool("LOG_LLM_TOKENS", False)
    log_audio_chunks = get_env_bool("LOG_TTS_CHUNKS", False)
    greeting_text = get_env("GREETING_TEXT")
    enable_greeting = get_env_bool("ENABLE_GREETING", True)
    max_history_turns = int(get_env("MAX_HISTORY_TURNS", "10"))
    session_store_path = get_env("SESSION_STORE_PATH")
    dynamic_stt_language = get_env_bool("STT_DYNAMIC_LANGUAGE", True)
    preview_partials = get_env_bool("PREVIEW_PARTIALS", True)
    preview_after_ms = int(get_env("PREVIEW_AFTER_MS", "300"))
    llm_timeout_s = float(get_env("LLM_STREAM_TIMEOUT_S", "20"))
    tts_timeout_s = float(get_env("TTS_STREAM_TIMEOUT_S", "20"))
    stt_connect_timeout_s = float(get_env("STT_CONNECT_TIMEOUT_S", "10"))
    tts_stream_chunk_chars = int(get_env("TTS_STREAM_CHUNK_CHARS", "50"))
    tts_stream_flush_punct = get_env_bool("TTS_STREAM_FLUSH_PUNCT", True)
    tts_min_buffer_size = int(get_env("TTS_MIN_BUFFER_SIZE", "30"))
    tts_max_chunk_length = int(get_env("TTS_MAX_CHUNK_LENGTH", "120"))
    tts_output_audio_codec = get_env("TTS_OUTPUT_AUDIO_CODEC", "linear16")
    tts_output_audio_bitrate = get_env("TTS_OUTPUT_AUDIO_BITRATE")
    post_speech_pause_ms = int(get_env("POST_SPEECH_PAUSE_MS", "800"))  # 800ms suits Indian speech cadence

    stt = SaarikaSTTService(
        ws_url=stt_ws_url,
        api_key=stt_api_key,
        model=stt_model,
        language=stt_language,
        language_param_key=stt_language_key,
        vad_signals=stt_vad_signals,
        high_vad_sensitivity=stt_high_vad,
        flush_signal=stt_flush_signal,
        audio_payload_format=stt_audio_payload_format,
        input_audio_codec=stt_input_audio_codec,
        audio_encoding=stt_audio_encoding,
        log_raw_messages=stt_log_raw,
        use_sdk=use_sdk_stt,
    )
    llm = SarvamLLMService(base_url=llm_base_url, api_key=llm_api_key, model=llm_model, use_sdk=use_sdk_llm)

    log_event(
        logger,
        "config_loaded",
        stt_ws_url=stt_ws_url,
        llm_base_url=llm_base_url,
        tts_ws_url=tts_ws_url,
        stt_model=stt_model,
        llm_model=llm_model,
        tts_model=tts_model,
        api_key_prefix=mask_secret(api_key),
        stt_key_prefix=mask_secret(stt_api_key),
        llm_key_prefix=mask_secret(llm_api_key),
        tts_key_prefix=mask_secret(tts_api_key),
        use_sdk=use_sdk,
        use_sdk_stt=use_sdk_stt,
        use_sdk_llm=use_sdk_llm,
        use_sdk_tts=use_sdk_tts,
        tts_output_audio_codec=tts_output_audio_codec,
        tts_min_buffer_size=tts_min_buffer_size,
        tts_max_chunk_length=tts_max_chunk_length,
    )

    def _on_snapshot_update(snapshot: Dict[str, Any]) -> None:
        # Single propagation call: session snapshot updates -> task compliance block sync.
        snapshot["tenant_id"] = ws_tenant_id
        snapshot["actor_user_id"] = ws_user_id
        snapshot["actor_username"] = ws_username or str(user_row.get("username") or "")
        workbench.apply_session_gate_status(session_snapshot=snapshot, actor="system")

    session = WebCallSession(
        stt_service=stt,
        llm_service=llm,
        tts_ws_url=tts_ws_url,
        send_event=sender.send_json,
        send_audio=sender.send_bytes,
        tts_api_key=tts_api_key,
        tts_voice=tts_voice,
        tts_model=tts_model,
        tts_speaker=tts_speaker,
        greeting_text=greeting_text if enable_greeting else None,
        max_history_turns=max_history_turns,
        session_store_path=session_store_path,
        dynamic_stt_language=dynamic_stt_language,
        preview_partials=preview_partials,
        preview_after_ms=preview_after_ms,
        llm_timeout_s=llm_timeout_s,
        tts_timeout_s=tts_timeout_s,
        stt_connect_timeout_s=stt_connect_timeout_s,
        vad_rms_threshold=vad_threshold,
        vad_silence_ms=vad_silence_ms,
        stt_flush_interval_ms=stt_flush_interval_ms,
        log_partial_transcripts=log_partials,
        log_tokens=log_tokens,
        log_audio_chunks=log_audio_chunks,
        use_sdk=use_sdk_tts,
        tts_stream_chunk_chars=tts_stream_chunk_chars,
        tts_stream_flush_punct=tts_stream_flush_punct,
        tts_min_buffer_size=tts_min_buffer_size,
        tts_max_chunk_length=tts_max_chunk_length,
        tts_output_audio_codec=tts_output_audio_codec,
        tts_output_audio_bitrate=tts_output_audio_bitrate,
        post_speech_pause_ms=post_speech_pause_ms,
        audit_store=audit,
        excel_sink=sink,
        action_router=router,
        knowledge_store=knowledge,
        session_registry=sessions,
        snapshot_update_hook=_on_snapshot_update,
        strategy_engine=strategy,
        compliance_engine=compliance,
        followup_service=followups,
        crm_adapter=crm,
    )

    def _sync_session_snapshot() -> Dict[str, Any]:
        snap_local = session.get_snapshot()
        snap_local["tenant_id"] = ws_tenant_id
        snap_local["actor_user_id"] = ws_user_id
        snap_local["actor_username"] = ws_username or str(user_row.get("username") or "")
        sessions[session._session_id] = snap_local
        _on_snapshot_update(snap_local)
        return snap_local

    audit.record_event(
        event_type="ws_open",
        session_id=session._session_id,
        payload={
            "client": f"{client.host}:{client.port}" if client else None,
            "user_id": ws_user_id,
            "role": ws_role,
            "tenant_id": ws_tenant_id,
            "session_id": ws_session_id or None,
        },
    )
    sink.upsert_call(session_id=session._session_id, customer_id=None, start_ts=round(time.time(), 3))
    snap = _sync_session_snapshot()
    audit.upsert_outcome(
        session_id=session._session_id,
        start_ts=round(time.time(), 3),
        customer_id=snap.get("customer_id"),
        campaign_id=snap.get("campaign_id"),
        dpd_bucket=snap.get("dpd_bucket"),
    )
    session_objs[session._session_id] = session

    session_started = False
    text_message_window: List[float] = []
    try:
        while True:
            message = await websocket.receive()
            if message.get("type") == "websocket.disconnect":
                break
            if message.get("bytes") is not None:
                if not session_started:
                    continue
                await session.handle_audio(message["bytes"])
                continue
            if message.get("text"):
                now_ts = time.time()
                text_message_window = [ts for ts in text_message_window if now_ts - ts < 1.0]
                text_message_window.append(now_ts)
                if len(text_message_window) > 40:
                    await sender.send_json({"type": "error", "message": "Message rate limit exceeded"})
                    continue
                try:
                    data = json.loads(message["text"])
                except json.JSONDecodeError:
                    continue
                msg_type = data.get("type")

                if msg_type == "start":
                    if session_started:
                        await sender.send_json({"type": "started"})
                        continue
                    # UI may send session context together with start.
                    ctx = _extract_context_from_payload(data)
                    if not ctx.get("customer_id"):
                        await sender.send_json({"type": "error", "message": "Please select a customer before starting."})
                        continue
                    preferred_lang = ctx.get("language_preference")
                    if preferred_lang:
                        session.stt.language = session._resolve_stt_connect_language(str(preferred_lang)) or str(preferred_lang)
                    if ctx and hasattr(session, "set_context"):
                        await session.set_context(ctx)
                    try:
                        await session.start()
                    except Exception as exc:
                        log_event(logger, "session_start_failed", error=str(exc))
                        await sender.send_json({"type": "error", "message": "STT connection failed"})
                        await websocket.close()
                        return
                    await sender.send_json({"type": "context_received", "context": ctx})
                    await sender.send_json({"type": "started"})
                    if hasattr(session, "start_greeting"):
                        await session.start_greeting()
                    _sync_session_snapshot()
                    session_started = True
                    continue

                if msg_type == "stop":
                    break

                if msg_type == "ping":
                    await sender.send_json({"type": "pong"})
                    continue

                if msg_type == "audio_started":
                    if not session_started:
                        continue
                    log_event(logger, "audio_started", session_id=session._session_id)
                    audit.record_event(event_type="audio_started", session_id=session._session_id, payload={})
                    continue

                # Allow typed/text input from the UI (e.g., the yellow input block)
                if msg_type in {"text", "user_text", "chat"}:
                    if not session_started:
                        await sender.send_json({"type": "error", "message": "Start the session before sending text."})
                        continue
                    text = (data.get("text") or "").strip()
                    if text:
                        if hasattr(session, "handle_text"):
                            await session.handle_text(text)
                        else:
                            await sender.send_json(
                                {"type": "error", "message": "Text input is not supported by this server build."}
                            )
                        _sync_session_snapshot()
                    continue

                # Allow the UI to pass known context (customer name / amount / due date etc.)
                if msg_type in {"set_context", "context"}:
                    ctx = _extract_context_from_payload(data)
                    if ctx and hasattr(session, "set_context"):
                        await session.set_context(ctx)
                    await sender.send_json({"type": "context_received", "context": ctx})
                    _sync_session_snapshot()
                    continue


                if msg_type == "set_disposition":
                    disp = str(data.get("disposition") or "").strip()
                    if disp and hasattr(session, "_wf_state"):
                        session._wf_state.disposition = disp
                        log_event(logger, "manual_disposition", session_id=session._session_id, disposition=disp)
                        await sender.send_json({"type": "disposition_set", "disposition": disp})
                        _sync_session_snapshot()
                    continue

                if msg_type == "action":
                    if not session_started:
                        await sender.send_json(
                            {"type": "action_error", "name": data.get("name"), "ok": False, "error": "Start the session before taking actions."}
                        )
                        continue
                    if not has_permissions(ws_role, (WORKBENCH_MUTATE,)):
                        await sender.send_json(
                            {
                                "type": "action_error",
                                "name": data.get("name"),
                                "ok": False,
                                "error": "Role is not allowed to perform call actions.",
                            }
                        )
                        continue
                    name = str(data.get("name") or "").strip()
                    payload = data.get("payload")
                    if not isinstance(payload, dict):
                        payload = {}
                    if not name:
                        await sender.send_json(
                            {"type": "action_error", "name": name, "ok": False, "error": "Missing action name."}
                        )
                        continue
                    try:
                        result = await session.handle_action(name=name, payload=payload)
                    except Exception as exc:
                        await sender.send_json(
                            {"type": "action_error", "name": name, "ok": False, "error": str(exc)}
                        )
                        continue
                    await sender.send_json({"type": "action_result", "name": name, "ok": True, "result": result})
                    _sync_session_snapshot()
                    continue
    except WebSocketDisconnect:
        pass
    except Exception as exc:
        logger.exception("Websocket error: %s", exc)
    finally:
        if ws_user_id:
            live_count = int(_ws_user_connections.get(ws_user_id, 0)) - 1
            if live_count <= 0:
                _ws_user_connections.pop(ws_user_id, None)
            else:
                _ws_user_connections[ws_user_id] = live_count
        # Final snapshot + outcome.
        try:
            _sync_session_snapshot()
        except Exception as exc:
            log_event(logger, "session_snapshot_error", session_id=session._session_id, error=str(exc))
        await session.stop()
        await sender.close()
        try:
            session_objs.pop(session._session_id, None)
        except Exception as exc:
            log_event(logger, "session_registry_remove_error", session_id=session._session_id, error=str(exc))
        try:
            audit.upsert_outcome(
                session_id=session._session_id,
                end_ts=round(time.time(), 3),
                customer_id=session.get_snapshot().get("customer_id"),
                campaign_id=session.get_snapshot().get("campaign_id"),
                dpd_bucket=session.get_snapshot().get("dpd_bucket"),
                disposition=session.get_snapshot().get("disposition"),
                ptp_date=session.get_snapshot().get("ptp_date"),
                callback_time=session.get_snapshot().get("callback_time"),
            )
            sink.upsert_call(
                session_id=session._session_id,
                customer_id=session.get_snapshot().get("customer_id"),
                end_ts=round(time.time(), 3),
                disposition=session.get_snapshot().get("disposition"),
                ptp_date=session.get_snapshot().get("ptp_date"),
                callback_time=session.get_snapshot().get("callback_time"),
            )
        except Exception as exc:
            log_event(logger, "session_finalize_persist_error", session_id=session._session_id, error=str(exc))
        log_event(logger, "ws_close")


# ---------------------------------------------------------------------------
# Analytics endpoints
# ---------------------------------------------------------------------------

@app.get("/api/metrics/roll-forward")
async def api_metrics_roll_forward(request: Request):
    deny = _require_admin(request)
    if deny:
        return deny
    try:
        days = int(request.query_params.get("days") or 30)
    except Exception:
        days = 30
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    return audit.roll_forward_matrix(days=days)


@app.get("/api/metrics/recovery")
async def api_metrics_recovery(request: Request):
    deny = _require_admin(request)
    if deny:
        return deny
    try:
        days = int(request.query_params.get("days") or 30)
    except Exception:
        days = 30
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    return audit.realized_recovery_trend(days=days)


@app.get("/api/metrics/agents")
async def api_metrics_agents(request: Request):
    deny = _require_admin(request)
    if deny:
        return deny
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    return {"agents": audit.agent_metrics()}


# ---------------------------------------------------------------------------
# Borrower 360
# ---------------------------------------------------------------------------

@app.get("/api/customers/{customer_id}/360")
async def api_customer_360(customer_id: str, request: Request):
    deny = _require_admin(request)
    if deny:
        return deny
    demo = _get_demo_singletons()
    c360: Customer360Service = demo["customer_360"]  # type: ignore[assignment]
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    data = c360.get_customer_360(tenant_id=ctx["tenant_id"], customer_id=customer_id)
    timeline = c360.get_timeline(customer_id=customer_id)
    data["timeline"] = timeline
    return data


@app.get("/api/customers/{customer_id}/nba")
async def api_customer_nba(customer_id: str, request: Request):
    deny = _require_admin(request)
    if deny:
        return deny
    demo = _get_demo_singletons()
    brain: RecoveryBrainService = demo["recovery_brain"]  # type: ignore[assignment]
    c360: Customer360Service = demo["customer_360"]  # type: ignore[assignment]
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)

    snapshot = c360.get_customer_360(tenant_id=ctx["tenant_id"], customer_id=customer_id)
    profile = snapshot.get("profile") if isinstance(snapshot.get("profile"), dict) else {}
    loans = snapshot.get("loan_accounts") if isinstance(snapshot.get("loan_accounts"), list) else []
    timeline = c360.get_timeline(customer_id=customer_id)

    def _to_int(v: Any) -> Optional[int]:
        if v in (None, ""):
            return None
        try:
            return int(float(v))
        except Exception:
            return None

    def _to_float(v: Any) -> Optional[float]:
        if v in (None, ""):
            return None
        try:
            return float(v)
        except Exception:
            return None

    dpd_candidates: List[int] = []
    for v in [profile.get("dpd")] + [x.get("dpd") for x in loans if isinstance(x, dict)]:
        parsed = _to_int(v)
        if parsed is not None:
            dpd_candidates.append(parsed)
    dpd = max(dpd_candidates) if dpd_candidates else 0

    amount_due = 0.0
    for loan in loans:
        if not isinstance(loan, dict):
            continue
        parsed_amount = _to_float(loan.get("principal_outstanding"))
        if parsed_amount is not None and parsed_amount > amount_due:
            amount_due = parsed_amount

    last_outcome = ""
    for ev in timeline:
        if str(ev.get("type") or "") != "call":
            continue
        disposition = str(ev.get("disposition") or "").strip().lower()
        if disposition:
            last_outcome = disposition
            break

    tags = profile.get("tags") if isinstance(profile.get("tags"), list) else []
    norm_tags = [str(t).strip().lower() for t in tags if str(t).strip()]
    risk_band = str(profile.get("risk_band") or "").strip().lower()
    hardship_flag = (
        "hardship" in risk_band
        or "vulnerable" in risk_band
        or any(("hardship" in t or "job_loss" in t or "medical" in t) for t in norm_tags)
    )
    dispute_flag = ("dispute" in last_outcome) or any("dispute" in t for t in norm_tags)

    features = {
        "dpd": dpd,
        "amount_due": amount_due,
        "risk_band": risk_band,
        "last_outcome": last_outcome,
        "hardship": hardship_flag,
        "dispute": dispute_flag,
    }

    try:
        result = brain.decide_nba(
            tenant_id=ctx["tenant_id"],
            payload={
                "customer_id": customer_id,
                "dpd": dpd,
                "amount_due": amount_due,
                "features": features,
            },
            actor=ctx["actor"],
            request_id=ctx["request_id"],
        )
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=500)

    nba = dict(result)
    message = nba.get("message") if isinstance(nba.get("message"), dict) else {}
    action_type = str(nba.get("action_type") or "")
    timing = "immediate"
    if action_type == "reschedule_contact":
        timing = "next_working_day"
    elif action_type in {"offer_settlement", "escalate_dispute"}:
        timing = "supervisor_gate"
    action = {
        "action_type": action_type,
        "channel": nba.get("channel"),
        "timing": timing,
        "message_template": message.get("text"),
        "rationale": f"DPD {dpd}, risk {risk_band or 'unknown'}, last_outcome {last_outcome or 'none'}",
    }
    nba["action"] = action
    nba["recommended_action"] = action
    nba["features_used"] = features
    return {"schema_version": "v2.1", "nba": nba}


@app.get("/api/customers/{customer_id}/settlements")
async def api_customer_settlements(customer_id: str, request: Request):
    deny = _require_admin(request)
    if deny:
        return deny
    demo = _get_demo_singletons()
    settlements: SettlementService = demo["settlement"]  # type: ignore[assignment]
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    offers = settlements.list_offers(tenant_id=ctx["tenant_id"], customer_id=customer_id)
    return {"offers": offers, "total": len(offers)}


@app.post("/api/sessions/{session_id}/settlement")
async def api_session_settlement(session_id: str, request: Request):
    deny = _require_admin(request)
    if deny:
        return deny
    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"error": "invalid_json"}, status_code=400)
    demo = _get_demo_singletons()
    settlements: SettlementService = demo["settlement"]  # type: ignore[assignment]
    approvals: ApprovalService = demo["approval_service"]  # type: ignore[assignment]
    try:
        ctx = _v2_context(request)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)

    try:
        offered_amount = float(body.get("offered_amount") or 0)
    except (TypeError, ValueError):
        offered_amount = 0.0
    if offered_amount <= 0:
        return JSONResponse({"error": "offered_amount_required"}, status_code=400)

    original_amount = None
    try:
        original_amount = float(body.get("original_amount") or 0) or None
    except (TypeError, ValueError):
        pass
    counter_offer_of = str(body.get("counter_offer_of") or "").strip() or None

    try:
        payload = {
            "customer_id": str(body.get("customer_id") or "").strip() or None,
            "loan_account_id": str(body.get("loan_account_id") or "").strip() or None,
            "offered_amount": offered_amount,
            "original_due_amount": original_amount,
            "terms": body.get("terms") or {"channel": "upi", "one_time": True},
            "journey_id": session_id,
        }
        if counter_offer_of:
            offer = settlements.counter_offer(
                tenant_id=ctx["tenant_id"],
                base_offer_id=counter_offer_of,
                payload=payload,
                actor=ctx["actor"],
                request_id=ctx["request_id"],
            )
        else:
            offer = settlements.create_offer(
                tenant_id=ctx["tenant_id"],
                payload=payload,
                actor=ctx["actor"],
                request_id=ctx["request_id"],
            )
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)

    discount_pct = 0.0
    terms = offer.get("terms") if isinstance(offer.get("terms"), dict) else {}
    try:
        discount_pct = float(terms.get("discount_pct") or 0.0)
    except Exception:
        discount_pct = 0.0
    if discount_pct == 0.0 and original_amount and original_amount > 0:
        discount_pct = round((1.0 - offered_amount / original_amount) * 100.0, 2)

    requires_approval = discount_pct > 10
    if requires_approval:
        try:
            approvals.enqueue(
                tenant_id=ctx["tenant_id"],
                action_type="settlement_offer",
                actor=ctx["actor"],
                request_id=ctx["request_id"],
                reference_type="settlement_offer",
                reference_id=offer.get("id"),
                risk_score=discount_pct / 100.0,
                payload={"offer": offer, "discount_pct": discount_pct},
            )
        except Exception as exc:
            log_event(logger, "settlement_approval_enqueue_error", session_id=session_id, error=str(exc))

    return {
        "offer": offer,
        "discount_pct": discount_pct,
        "requires_approval": requires_approval,
        "round": int(terms.get("round") or 1),
        "counter_offer_of": terms.get("counter_of"),
    }


# ---------------------------------------------------------------------------
# Post-call summaries
# ---------------------------------------------------------------------------

@app.get("/api/sessions/{session_id}/summary")
async def api_session_summary(session_id: str, request: Request):
    deny = _require_admin(request)
    if deny:
        return deny
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    summary = audit.get_call_summary(session_id)
    if summary is None:
        return JSONResponse({"error": "summary_not_found"}, status_code=404)
    return {"summary": summary}


@app.get("/api/summaries/search")
async def api_summaries_search(request: Request):
    deny = _require_admin(request)
    if deny:
        return deny
    q = str(request.query_params.get("q") or "").strip()
    try:
        limit = int(request.query_params.get("limit") or 20)
    except Exception:
        limit = 20
    demo = _get_demo_singletons()
    audit: SQLiteAuditStore = demo["audit"]  # type: ignore[assignment]
    rows = audit.search_summaries(query=q, limit=limit)
    return {"summaries": rows, "total": len(rows), "query": q}


# ---------------------------------------------------------------------------
# CRM (Zoho) sync
# ---------------------------------------------------------------------------

@app.get("/api/crm/status")
async def api_crm_status(request: Request):
    deny = _require_admin(request)
    if deny:
        return deny
    demo = _get_demo_singletons()
    zoho: ZohoCRM = demo["zoho_crm"]  # type: ignore[assignment]
    try:
        limit = int(request.query_params.get("limit") or 20)
    except Exception:
        limit = 20
    status = zoho.status()
    log_entries = zoho.sync_log(limit=limit)
    return {"provider": "zoho", "status": status, "log": log_entries}


@app.post("/api/crm/sync")
async def api_crm_sync(request: Request):
    deny = _require_admin(request)
    if deny:
        return deny
    demo = _get_demo_singletons()
    zoho: ZohoCRM = demo["zoho_crm"]  # type: ignore[assignment]
    db_path = get_env("DEMO_DB_PATH", os.path.join(DATA_DIR, "demo.db"))
    stats = zoho.process_queue(main_db_path=db_path)
    return {"ok": True, "stats": stats}


def main() -> None:
    import uvicorn

    load_dotenv(override=True)
    configure_logging()
    host = get_env("API_HOST", "0.0.0.0")
    port = int(get_env("API_PORT", get_env("PORT", "8000")))
    log_level = get_env("LOG_LEVEL", "info").lower()

    uvicorn.run("web_app:app", host=host, port=port, log_level=log_level, reload=False)


if __name__ == "__main__":
    configure_logging()
    main()
